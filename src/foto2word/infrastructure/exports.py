"""Extra output formats: a clean 'scanned' PDF of the photos and the tables as Excel."""

from collections.abc import Sequence
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from PIL import Image, ImageFilter, ImageOps

from foto2word.domain.document import Document, TableBlock, plain_text
from foto2word.domain.images import PageImage

_A4_PORTRAIT = (1240, 1754)  # 150 dpi
_SCAN_DPI = 150


def scan_pdf(pages: Sequence[PageImage]) -> bytes:
    """Photos as a printable 'scan': greyscale, levelled contrast, sharpened, fitted to A4."""
    sheets = []
    for page in pages:
        with Image.open(BytesIO(page.data)) as image:
            gray = ImageOps.autocontrast(image.convert("L"), cutoff=2)
        gray = gray.filter(ImageFilter.UnsharpMask(radius=2, percent=80))
        size = _A4_PORTRAIT if gray.height >= gray.width else _A4_PORTRAIT[::-1]
        sheet = Image.new("L", size, 255)
        fitted = ImageOps.contain(gray, size)
        sheet.paste(fitted, ((size[0] - fitted.width) // 2, (size[1] - fitted.height) // 2))
        sheets.append(sheet)
    if not sheets:
        raise ValueError("No pages to export")
    buffer = BytesIO()
    sheets[0].save(
        buffer, format="PDF", save_all=True, append_images=sheets[1:], resolution=_SCAN_DPI
    )
    return buffer.getvalue()


def tables_xlsx(document: Document) -> bytes | None:
    """Every table of the document on its own sheet; None when there are no tables."""
    tables = [b for b in document.blocks if isinstance(b, TableBlock)]
    if not tables:
        return None
    workbook = Workbook()
    workbook.remove(workbook.active)  # type: ignore[arg-type]
    for number, table in enumerate(tables, start=1):
        sheet = workbook.create_sheet(f"Таблица {number}")
        widths: dict[int, int] = {}
        for placement in table.placements():
            row, col = placement.row + 1, placement.col + 1
            text = plain_text(placement.cell.runs)
            cell = sheet.cell(row=row, column=col, value=_number_or_text(text))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if placement.row < table.header_rows:
                cell.font = Font(bold=True)
            if placement.cell.col_span > 1 or placement.cell.row_span > 1:
                sheet.merge_cells(
                    start_row=row,
                    start_column=col,
                    end_row=row + placement.cell.row_span - 1,
                    end_column=col + placement.cell.col_span - 1,
                )
            widths[col] = max(widths.get(col, 8), min(len(text) + 2, 60))
        for col, width in widths.items():
            sheet.column_dimensions[get_column_letter(col)].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _number_or_text(text: str) -> float | str:
    """'14,3' -> 14.3 so Excel can sum it; everything else stays text exactly as printed."""
    candidate = text.strip().replace(" ", "").replace(" ", "")
    if candidate.count(",") == 1 and "." not in candidate:
        candidate = candidate.replace(",", ".")
    try:
        value = float(candidate)
    except ValueError:
        return text
    is_code = candidate.startswith("0") and len(candidate) > 1 and "." not in candidate
    return text if is_code else value  # '007' is a code, not seven
