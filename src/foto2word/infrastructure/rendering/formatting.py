"""Low-level python-docx helpers: paragraph formatting and run writing."""

from dataclasses import dataclass
from typing import Protocol

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Mm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run as DocxRun

from foto2word.domain.document import Alignment, BorderLine, CellBorders, ParagraphLayout, Run
from foto2word.domain.extras import NO_EXTRAS, RenderExtras
from foto2word.domain.profile import PageSetup, Profile, RenderMode, TextStyle
from foto2word.infrastructure.rendering.pictures import add_picture, set_crop

WORD_ALIGNMENT = {  # DocModel alignment -> python-docx
    Alignment.LEFT: WD_ALIGN_PARAGRAPH.LEFT,
    Alignment.CENTER: WD_ALIGN_PARAGRAPH.CENTER,
    Alignment.RIGHT: WD_ALIGN_PARAGRAPH.RIGHT,
    Alignment.JUSTIFY: WD_ALIGN_PARAGRAPH.JUSTIFY,
}


_TAB_ALIGN = {
    "left": WD_TAB_ALIGNMENT.LEFT,
    "center": WD_TAB_ALIGNMENT.CENTER,
    "right": WD_TAB_ALIGNMENT.RIGHT,
    "decimal": WD_TAB_ALIGNMENT.DECIMAL,
}
_TAB_LEADER = {
    "none": WD_TAB_LEADER.SPACES,
    "dot": WD_TAB_LEADER.DOTS,
    "underscore": WD_TAB_LEADER.LINES,
    "hyphen": WD_TAB_LEADER.DASHES,
}


@dataclass(frozen=True, slots=True)
class Tab:
    position_mm: float
    right_aligned: bool = False


class Container(Protocol):
    """Where blocks are written: the document body, a header or a footer."""

    def add_paragraph(self) -> Paragraph: ...

    def add_table(self, rows: int, cols: int) -> Table: ...


@dataclass(frozen=True, slots=True)
class RenderContext:
    docx: Container
    profile: Profile
    mode: RenderMode
    extras: RenderExtras = NO_EXTRAS
    page: PageSetup | None = None
    """Page of the original document; overrides the profile page in faithful mode."""

    @property
    def faithful(self) -> bool:
        return self.mode is RenderMode.FAITHFUL

    @property
    def content_width_mm(self) -> float:
        return (self.page or self.profile.page).content_width_mm

    def styled(
        self, style: TextStyle, layout: ParagraphLayout | None, *, own_tabs: bool = True
    ) -> TextStyle:
        """Faithful mode applies the spacing, indents and tab stops measured on the original.

        own_tabs=False is for blocks that place their columns with tab stops of their own.
        """
        if not self.faithful or layout is None:
            return style
        update = {
            name: value
            for name in type(layout).model_fields
            if (value := getattr(layout, name)) not in (None, ()) and (own_tabs or name != "tabs")
        }
        if self.extras.editable:
            update["line_at_least"] = True  # an exact line clips letters of a larger size
        return style.model_copy(update=update)

    def alignment(self, style: TextStyle, seen_on_photo: Alignment | None) -> Alignment:
        """Faithful mode keeps the photo's alignment when the model reported it."""
        if self.mode is RenderMode.FAITHFUL and seen_on_photo is not None:
            return seen_on_photo
        return style.align

    def add_paragraph(
        self,
        style: TextStyle,
        align: Alignment | None = None,
        *,
        left_indent_mm: float | None = None,
        first_line_indent_mm: float | None = None,
        tabs: tuple[Tab, ...] = (),
    ) -> Paragraph:
        paragraph = self.docx.add_paragraph()
        self.format_paragraph(
            paragraph,
            style,
            align,
            left_indent_mm=left_indent_mm,
            first_line_indent_mm=first_line_indent_mm,
            tabs=tabs,
        )
        return paragraph

    @staticmethod
    def format_paragraph(
        paragraph: Paragraph,
        style: TextStyle,
        align: Alignment | None = None,
        *,
        left_indent_mm: float | None = None,
        first_line_indent_mm: float | None = None,
        tabs: tuple[Tab, ...] = (),
    ) -> None:
        fmt = paragraph.paragraph_format
        fmt.alignment = WORD_ALIGNMENT[align or style.align]
        if style.line_pt is None:
            fmt.line_spacing = style.line_spacing
        else:
            fmt.line_spacing = Pt(style.line_pt)
            fmt.line_spacing_rule = (
                WD_LINE_SPACING.AT_LEAST if style.line_at_least else WD_LINE_SPACING.EXACTLY
            )
        fmt.space_before = Pt(style.space_before_pt)
        fmt.space_after = Pt(style.space_after_pt)
        fmt.left_indent = Mm(style.left_indent_mm if left_indent_mm is None else left_indent_mm)
        fmt.first_line_indent = Mm(
            style.first_line_indent_mm if first_line_indent_mm is None else first_line_indent_mm
        )
        fmt.right_indent = Mm(style.right_indent_mm)
        if style.mark_font:  # an empty line is as tall as its mark
            set_mark_font(paragraph, style.mark_font)
        if style.mark_size_pt:
            set_mark_size(paragraph, style.mark_size_pt)
        if style.keep_next:
            fmt.keep_with_next = True
        if style.keep_lines:
            fmt.keep_together = True
        _set_place(paragraph, style)
        if style.borders is not None:
            _set_paragraph_borders(paragraph, style.borders)
        if style.shading is not None:
            shading = OxmlElement("w:shd")
            for name, value in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", style.shading)):
                shading.set(qn(name), value)
            paragraph._p.get_or_add_pPr().insert_element_before(shading, *_AFTER_PARAGRAPH_SHADING)
        for stop in style.tabs:
            fmt.tab_stops.add_tab_stop(
                Mm(stop.position_mm), _TAB_ALIGN[stop.align], _TAB_LEADER[stop.leader]
            )
        for tab in tabs:
            fmt.tab_stops.add_tab_stop(
                Mm(tab.position_mm),
                WD_TAB_ALIGNMENT.RIGHT if tab.right_aligned else WD_TAB_ALIGNMENT.LEFT,
            )

    def _write_picture(self, paragraph: Paragraph, run: Run) -> None:
        picture = run.picture
        data = self.extras.images.get(picture.image_id) if picture is not None else None
        if picture is None or data is None or not self.faithful:
            return  # no pixels (a photo's copy): the picture is left out
        target = paragraph.add_run()
        shape = add_picture(paragraph, target, data, picture.width_mm, picture.height_mm)
        if shape is None:
            target._r.getparent().remove(target._r)
        elif picture.crop:
            set_crop(shape._inline, picture.crop)

    def write_runs(self, paragraph: Paragraph, runs: tuple[Run, ...], style: TextStyle) -> None:
        for run in runs:
            if run.picture is not None:  # a picture in the line, at its place among the text
                self._write_picture(paragraph, run)
                continue
            if run.math is not None:  # Word's own equation, as it was
                paragraph._p.append(parse_xml(run.math))
                continue
            if run.vml is not None:  # an old-style drawing in the line, as it was
                paragraph.add_run()._r.append(parse_xml(run.vml))
                continue
            if run.footnote is not None and self.faithful:
                target = paragraph.add_run()
                self._format_run(target, style, run)
                if run.footnote:  # the mark in the text: Word numbers it and sets the note
                    mark = OxmlElement("w:footnoteReference")
                    mark.set(qn("w:id"), str(run.footnote))
                else:  # the number at the start of the note itself
                    mark = OxmlElement("w:footnoteRef")
                target._r.append(mark)
                continue
            if run.field is not None and self.faithful:  # Word counts it on every page
                self._write_field(paragraph, run, style)
                continue
            target = paragraph.add_run(run.text)
            self._format_run(target, style, run)
            if run.clear is not None and self.faithful:  # the text after waits below floats
                for brk in target._r.findall(qn("w:br")):
                    brk.set(qn("w:type"), "textWrapping")
                    brk.set(qn("w:clear"), run.clear)

    def _write_field(self, paragraph: Paragraph, run: Run, style: TextStyle) -> None:
        """PAGE or NUMPAGES as Word writes a field: its code, then the number shown."""
        code = " PAGE " if run.field == "page" else " NUMPAGES "
        for kind, text in (
            ("begin", None),
            (None, code),
            ("separate", None),
            (None, None),
            ("end", None),
        ):
            target = paragraph.add_run(run.text if kind is None and text is None else None)
            self._format_run(target, style, run)
            if kind is not None:
                char = OxmlElement("w:fldChar")
                char.set(qn("w:fldCharType"), kind)
                target._r.append(char)
            elif text is not None:
                instruction = OxmlElement("w:instrText")
                instruction.set(qn("xml:space"), "preserve")
                instruction.text = text
                target._r.append(instruction)

    def write_text(
        self, paragraph: Paragraph, text: str, style: TextStyle, size_pt: float | None = None
    ) -> None:
        self._format_run(paragraph.add_run(text), style, Run(text=text, size_pt=size_pt))

    def _format_run(self, target: DocxRun, style: TextStyle, run: Run) -> None:
        enforce_style = self.mode is RenderMode.STANDARD
        set_font(target, run.font if self.faithful and run.font else style.font)
        own_size = run.size_pt if self.faithful else None
        target.font.size = Pt(own_size or style.size_pt)
        if self.faithful and run.color:
            target.font.color.rgb = RGBColor.from_string(run.color.upper())
        target.font.bold = run.bold or (enforce_style and style.bold)
        target.font.italic = run.italic or (enforce_style and style.italic)
        target.font.all_caps = enforce_style and style.all_caps
        target.font.underline = run.underline
        if self.faithful and run.strike:
            target.font.strike = True
        target.font.superscript = run.superscript
        target.font.subscript = run.subscript
        if self.faithful and run.background:
            shading = OxmlElement("w:shd")
            for name, value in (
                ("w:val", "clear"),
                ("w:color", "auto"),
                ("w:fill", run.background),
            ):
                shading.set(qn(name), value)
            target._r.get_or_add_rPr().insert_element_before(
                shading,
                "w:fitText",
                "w:vertAlign",
                "w:rtl",
                "w:cs",
                "w:em",
                "w:lang",
                "w:eastAsianLayout",
                "w:specVanish",
                "w:oMath",
            )
        if self.faithful and run.kern_pt:
            kern = OxmlElement("w:kern")
            kern.set(qn("w:val"), str(round(run.kern_pt * 2)))
            target._r.get_or_add_rPr().insert_element_before(
                kern,
                "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect",
                "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em",
                "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath",
            )  # fmt: skip
        if self.faithful and run.scale_pct:
            scale = OxmlElement("w:w")
            scale.set(qn("w:val"), str(run.scale_pct))
            target._r.get_or_add_rPr().insert_element_before(
                scale,
                "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect",
                "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em",
                "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath",
            )  # fmt: skip
        if self.faithful and run.spacing_pt:
            spacing = OxmlElement("w:spacing")
            spacing.set(qn("w:val"), str(round(run.spacing_pt * 20)))
            target._r.get_or_add_rPr().insert_element_before(
                spacing,
                "w:w",
                "w:kern",
                "w:position",
                "w:sz",
                "w:szCs",
                "w:highlight",
                "w:u",
                "w:effect",
                "w:vertAlign",
                "w:lang",
            )


def _set_place(paragraph: Paragraph, style: TextStyle) -> None:
    """The original's frame at a fixed place, and its line height off the page's grid."""
    if style.frame:
        set_frame(paragraph, style.frame)
    if style.off_grid:
        paragraph._p.get_or_add_pPr().insert_element_before(
            OxmlElement("w:snapToGrid", {qn("w:val"): "0"}), *_AFTER_SNAP
        )
    if style.no_hyphens:
        paragraph._p.get_or_add_pPr().insert_element_before(
            OxmlElement("w:suppressAutoHyphens"), *_AFTER_NO_HYPHENS
        )


def set_frame(paragraph: Paragraph, attributes: tuple[tuple[str, str], ...]) -> None:
    """The paragraph stands in a frame at a fixed place (w:framePr)."""
    frame = OxmlElement("w:framePr")
    for name, value in attributes:
        frame.set(qn(f"w:{name}"), value)
    paragraph._p.get_or_add_pPr().insert_element_before(frame, *_AFTER_FRAME)


_AFTER_NO_HYPHENS = (
    "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
    "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
    "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
    "w:sectPr", "w:pPrChange",
)  # fmt: skip
_AFTER_SNAP = (
    "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap",
    "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)  # fmt: skip
_AFTER_FRAME = (
    "w:widowControl",
    "w:numPr",
    "w:suppressLineNumbers",
    "w:pBdr",
    "w:shd",
    "w:tabs",
    "w:suppressAutoHyphens",
    "w:kinsoku",
    "w:wordWrap",
    "w:overflowPunct",
    "w:topLinePunct",
    "w:autoSpaceDE",
    "w:autoSpaceDN",
    "w:bidi",
    "w:adjustRightInd",
    "w:snapToGrid",
    "w:spacing",
    "w:ind",
    "w:contextualSpacing",
    "w:mirrorIndents",
    "w:suppressOverlap",
    "w:jc",
    "w:textDirection",
    "w:textAlignment",
    "w:textboxTightWrap",
    "w:outlineLvl",
    "w:divId",
    "w:cnfStyle",
    "w:rPr",
    "w:sectPr",
    "w:pPrChange",
)
_AFTER_PARAGRAPH_SHADING = (
    "w:tabs",
    "w:suppressAutoHyphens",
    "w:kinsoku",
    "w:wordWrap",
    "w:overflowPunct",
    "w:topLinePunct",
    "w:autoSpaceDE",
    "w:autoSpaceDN",
    "w:bidi",
    "w:adjustRightInd",
    "w:snapToGrid",
    "w:spacing",
    "w:ind",
    "w:contextualSpacing",
    "w:mirrorIndents",
    "w:suppressOverlap",
    "w:jc",
    "w:textDirection",
    "w:textAlignment",
    "w:textboxTightWrap",
    "w:outlineLvl",
    "w:divId",
    "w:cnfStyle",
    "w:rPr",
    "w:sectPr",
    "w:pPrChange",
)
"""What follows w:shd in a paragraph's properties, in the schema's order."""


_RULE = BorderLine(size=6, space_pt=1)
"""A paragraph line nothing more is known about: a rule under a letterhead."""


def _set_paragraph_borders(paragraph: Paragraph, borders: CellBorders) -> None:
    """Single lines on the chosen sides of a paragraph (w:pBdr)."""
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:pBdr"))
    if existing is not None:
        p_pr.remove(existing)
    element = OxmlElement("w:pBdr")
    lines = dict(borders.lines)
    for side in ("top", "left", "bottom", "right"):
        if getattr(borders, side):
            line = lines.get(side, _RULE)
            edge = OxmlElement(f"w:{side}")
            edge.set(qn("w:val"), line.style)
            edge.set(qn("w:sz"), str(line.size))
            edge.set(qn("w:space"), str(line.space_pt))
            edge.set(qn("w:color"), line.color or "auto")
            element.append(edge)
    # pBdr sits after the keep/page-break group and before shading, tabs and spacing
    successor = next((c for c in p_pr if c.tag in {qn(f"w:{n}") for n in _AFTER_PBDR}), None)
    if successor is None:
        p_pr.append(element)
    else:
        successor.addprevious(element)


_AFTER_PBDR = (
    "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
    "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid",
    "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
    "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle",
    "rPr", "sectPr", "pPrChange",
)  # fmt: skip


def set_font(run: DocxRun, font_name: str) -> None:
    """Set the font for every script, including Cyrillic and East Asian slots."""
    run.font.name = font_name
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attribute), font_name)


def set_outline_level(paragraph: Paragraph, level: int) -> None:
    """Mark a paragraph as a heading for Word's navigation pane without built-in styles."""
    element = OxmlElement("w:outlineLvl")
    element.set(qn("w:val"), str(max(level - 1, 0)))
    paragraph._p.get_or_add_pPr().append(element)


def _mark_properties(paragraph: Paragraph) -> BaseOxmlElement:
    p_pr = paragraph._p.get_or_add_pPr()
    r_pr: BaseOxmlElement | None = p_pr.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        p_pr.append(r_pr)
    return r_pr


def set_mark_font(paragraph: Paragraph, font: str) -> None:
    """Typeface of the paragraph mark: an empty line is as tall as a line of it."""
    r_pr = _mark_properties(paragraph)
    for old in r_pr.findall(qn("w:rFonts")):
        r_pr.remove(old)
    fonts = OxmlElement("w:rFonts")
    for name in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(name), font)
    r_pr.insert(0, fonts)


def set_mark_size(paragraph: Paragraph, size_pt: float) -> None:
    """Size of the paragraph mark, which sets the height of an empty line."""
    r_pr = _mark_properties(paragraph)
    for old in r_pr.findall(qn("w:sz")):
        r_pr.remove(old)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(round(size_pt * 2)))
    r_pr.append(size)
