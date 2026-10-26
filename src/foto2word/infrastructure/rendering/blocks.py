"""One renderer per block kind. Adding a block kind = adding a class here and to the registry."""

import re
from collections.abc import Mapping
from dataclasses import replace
from io import BytesIO
from typing import Any, Protocol

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.oxml.table import CT_Row
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Mm, Pt, Twips
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from foto2word.domain.document import (
    Alignment,
    BorderLine,
    CellBorders,
    CellPadding,
    CellPlacement,
    FillRowBlock,
    HeadingBlock,
    ImageBlock,
    ImagePlacement,
    LabeledLineBlock,
    ListItemBlock,
    ParagraphBlock,
    ParagraphLayout,
    ParagraphRole,
    RequisitesBlock,
    Run,
    ShapeBlock,
    SignatureLineBlock,
    SpacerBlock,
    SpacerSize,
    TableBlock,
    TableCell,
    TablePosition,
    TableRow,
    TextBoxBlock,
    TwoColumnLineBlock,
    VerticalAlignment,
)
from foto2word.domain.extras import SignatureMode
from foto2word.infrastructure.rendering.formatting import (
    RenderContext,
    Tab,
    set_mark_font,
    set_mark_size,
    set_outline_level,
)
from foto2word.infrastructure.rendering.pictures import (
    add_picture,
    next_drawing_id,
    set_crop,
    set_effect,
)

SIGNATURE_BLANK = "_" * 16
_SPACER_LINES = {SpacerSize.SMALL: 1, SpacerSize.MEDIUM: 2, SpacerSize.LARGE: 4}


class BlockRenderer[B](Protocol):
    def render(self, block: B, ctx: RenderContext) -> None: ...


class HeadingRenderer:
    def render(self, block: HeadingBlock, ctx: RenderContext) -> None:
        styles = ctx.profile.styles
        style = ctx.styled(styles.title if block.level <= 1 else styles.subtitle, block.layout)
        paragraph = ctx.add_paragraph(style, ctx.alignment(style, block.align))
        ctx.write_runs(paragraph, block.runs, style)
        set_outline_level(paragraph, block.level)


class ParagraphRenderer:
    def render(self, block: ParagraphBlock, ctx: RenderContext) -> None:
        styles = ctx.profile.styles
        base = styles.note if block.role is not ParagraphRole.BODY else styles.body
        style = ctx.styled(base, block.layout)
        align = ctx.alignment(style, block.align)
        # Only running text gets the paragraph indent; a line pinned left/centre/right does not.
        measured = ctx.faithful and block.layout is not None
        indent = None if measured or align is Alignment.JUSTIFY else 0.0
        paragraph = ctx.add_paragraph(style, align, first_line_indent_mm=indent)
        ctx.write_runs(paragraph, block.runs, style)


_SYMBOL_FONTS = frozenset({"Symbol", "Wingdings", "Wingdings 2", "Wingdings 3", "Webdings"})
"""Typefaces whose letters are other signs: '•' typed in them is not a bullet."""


class ListItemRenderer:
    def render(self, block: ListItemBlock, ctx: RenderContext) -> None:
        style = ctx.styled(ctx.profile.styles.list_item, block.layout)
        # Nested items shift the first line only, so wrapped lines align with the parent item.
        measured = ctx.faithful and block.layout is not None
        nesting = 0 if measured else block.level * ctx.profile.list_indent_mm
        first_line = style.first_line_indent_mm + nesting
        paragraph = ctx.add_paragraph(
            style, ctx.alignment(style, block.align), first_line_indent_mm=first_line
        )
        # The marker looks like the start of the text: same face and size. With a hanging
        # indent Word puts the text after a tab, aligned with the wrapped lines.
        first = block.runs[0] if block.runs else Run(text="")
        hanging = measured and style.first_line_indent_mm < 0
        gap = "\t" if hanging else " "
        if ctx.faithful and block.marker_gap is not None:  # Word's own: a tab to the next stop
            gap = block.marker_gap
        text, font, size = block.marker, first.font, first.size_pt
        if ctx.faithful:  # the original's marker: its face and size set the line's height too
            size = block.marker_size_pt or size
            if block.marker_glyph is not None:  # the bullet's own sign in its own face
                text, font = block.marker_glyph, block.marker_font or font
            elif block.marker_font and block.marker_font not in _SYMBOL_FONTS:
                font = block.marker_font
        if (
            ctx.faithful
            and block.marker_gap is not None
            and _number(paragraph, text, font or style.font, size or style.size_pt, gap)
        ):
            ctx.write_runs(paragraph, block.runs, style)
            return
        marker = Run(text=f"{text}{gap}", size_pt=size, font=font)
        ctx.write_runs(paragraph, (marker, *block.runs), style)


_GAP_SUFFIX = {"\t": "tab", " ": "space", "": "nothing"}


def _number(paragraph: Paragraph, text: str, font: str, size_pt: float, gap: str) -> bool:
    """Gives the paragraph Word's own list marker showing `text`, as the original's list did:
    Word draws it, so the marker's face does not make the line taller as typed text would.
    False where a list cannot be written (a header, a '%' that Word would read as a number)."""
    numbering_part = getattr(paragraph.part, "numbering_part", None)
    if numbering_part is None or "%" in text or gap not in _GAP_SUFFIX:
        return False
    numbering = numbering_part.element
    key = f"f2w|{text}|{font}|{size_pt}|{_GAP_SUFFIX[gap]}"
    num_id = next(iter(_named_lists(numbering, key)), None)
    if num_id is None:
        abstract_id = 1 + max(
            (int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))),
            default=0,
        )
        num_id = str(
            1 + max((int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))), default=0)
        )
        half_points = round(size_pt * 2)
        w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        face = _xml(font)
        abstract = parse_xml(
            f'<w:abstractNum {w} w:abstractNumId="{abstract_id}"><w:name w:val="{_xml(key)}"/>'
            f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
            f'<w:suff w:val="{_GAP_SUFFIX[gap]}"/><w:lvlText w:val="{_xml(text)}"/>'
            f'<w:lvlJc w:val="left"/><w:rPr><w:rFonts w:ascii="{face}" w:hAnsi="{face}" '
            f'w:cs="{face}" w:eastAsia="{face}" w:hint="default"/>'
            f'<w:b w:val="0"/><w:i w:val="0"/><w:sz w:val="{half_points}"/>'
            f'<w:szCs w:val="{half_points}"/></w:rPr></w:lvl></w:abstractNum>'
        )
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:  # every abstractNum comes before the nums
            first_num.addprevious(abstract)
        else:
            numbering.append(abstract)
        numbering.append(
            parse_xml(
                f'<w:num {w} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/></w:num>'
            )
        )
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = int(num_id)
    return True


def _named_lists(numbering: BaseOxmlElement, key: str) -> set[str]:
    """numIds of the lists this renderer made for the marker `key`."""
    abstract_ids = {
        a.get(qn("w:abstractNumId"))
        for a in numbering.findall(qn("w:abstractNum"))
        if (name := a.find(qn("w:name"))) is not None and name.get(qn("w:val")) == key
    }
    return {
        n.get(qn("w:numId"))
        for n in numbering.findall(qn("w:num"))
        if (ref := n.find(qn("w:abstractNumId"))) is not None
        and ref.get(qn("w:val")) in abstract_ids
    }


def _xml(text: str) -> str:
    return (
        text.replace("&", "&amp;").replace('"', "&quot;").replace("<", "&lt;").replace(">", "&gt;")
    )


class LabeledLineRenderer:
    LABEL_SHARE = 0.35

    def render(self, block: LabeledLineBlock, ctx: RenderContext) -> None:
        style = ctx.styled(ctx.profile.styles.line, block.layout, own_tabs=False)
        tab = Tab(ctx.content_width_mm * self.LABEL_SHARE)
        paragraph = ctx.add_paragraph(style, Alignment.LEFT, first_line_indent_mm=0, tabs=(tab,))
        ctx.write_text(paragraph, f"{block.label}\t{block.value}", style, block.size_pt)


_EDGE_INSIDE_MM = 0.05


class TwoColumnLineRenderer:
    def render(self, block: TwoColumnLineBlock, ctx: RenderContext) -> None:
        style = ctx.styled(ctx.profile.styles.line, block.layout, own_tabs=False)
        # a hair inside the margin: rounded to twips, a stop on the edge may land past it and
        # Word then moves the right part to a line of its own
        tab = Tab(ctx.content_width_mm - _EDGE_INSIDE_MM, right_aligned=True)
        paragraph = ctx.add_paragraph(style, Alignment.LEFT, first_line_indent_mm=0, tabs=(tab,))
        ctx.write_text(paragraph, f"{block.left}\t{block.right}", style, block.size_pt)


class SignatureLineRenderer:
    BLANK_SHARE = 0.38
    NAME_SHARE = 0.70
    SIGNATURE_HEIGHT_MM = 12.0
    PLACEHOLDER = "/подпись/"

    def render(self, block: SignatureLineBlock, ctx: RenderContext) -> None:
        style = ctx.styled(ctx.profile.styles.line, block.layout, own_tabs=False)
        width = ctx.content_width_mm
        tabs = (Tab(width * self.BLANK_SHARE), Tab(width * self.NAME_SHARE))
        paragraph = ctx.add_paragraph(style, Alignment.LEFT, first_line_indent_mm=0, tabs=tabs)
        paragraph.paragraph_format.space_before = Mm(6)
        ctx.write_text(paragraph, f"{block.label}	", style, block.size_pt)
        self._signature_slot(block, ctx, paragraph)
        ctx.write_text(paragraph, f"	{block.name}", style, block.size_pt)

    def _signature_slot(
        self, block: SignatureLineBlock, ctx: RenderContext, paragraph: Paragraph
    ) -> None:
        extras = ctx.extras
        style = ctx.profile.styles.line
        if (
            extras.signature_mode is SignatureMode.OWN
            and extras.signature_png
            and block.id in extras.signed_blocks
        ):
            run = paragraph.add_run()
            run.add_picture(BytesIO(extras.signature_png), height=Mm(self.SIGNATURE_HEIGHT_MM))
        elif extras.signature_mode is SignatureMode.PLACEHOLDER and block.has_signature_mark:
            ctx.write_text(paragraph, self.PLACEHOLDER, style, block.size_pt)
        else:
            ctx.write_text(paragraph, SIGNATURE_BLANK, style, block.size_pt)


class RequisitesRenderer:
    def render(self, block: RequisitesBlock, ctx: RenderContext) -> None:
        style = ctx.styled(ctx.profile.styles.requisites, block.layout)
        right = ctx.faithful and block.align is Alignment.RIGHT
        centred = ctx.faithful and block.align is Alignment.CENTER  # in its own column
        size = block.size_pt if ctx.faithful else None
        font = block.font if ctx.faithful else None
        spacing = block.spacing_pt if ctx.faithful else None
        justified = ctx.faithful and block.align is Alignment.JUSTIFY
        align = Alignment.CENTER if centred else Alignment.JUSTIFY if justified else Alignment.LEFT
        own = ctx.faithful and block.layout is not None  # the original's indents set the wraps
        for line in block.lines:
            paragraph = ctx.add_paragraph(
                style,
                Alignment.RIGHT if right else align,
                left_indent_mm=None if own or not right else 0.0,
                first_line_indent_mm=None if own else 0,
            )
            run = Run(text=line, size_pt=size, font=font, spacing_pt=spacing)
            ctx.write_runs(paragraph, (run,), style)


class ImageRenderer:
    """The picture at its size; without pixels, a visible note of the same height."""

    def render(self, block: ImageBlock, ctx: RenderContext) -> None:
        floating = block.floating if ctx.faithful else None
        line = _picture_line(ctx) if ctx.faithful and block.same_line and not floating else None
        if floating is not None:
            paragraph = _clearing_paragraph(ctx) if block.clear else _anchor_paragraph(ctx)
        else:  # beside the picture before it, or on a line of its own
            style = ctx.styled(ctx.profile.styles.line, block.layout)
            measured = ctx.faithful and block.layout is not None  # the original's own indents
            paragraph = line or ctx.add_paragraph(
                style,
                block.align or Alignment.CENTER,
                left_indent_mm=None if measured else 0,
                first_line_indent_mm=None if measured else 0,
            )
            if block.lead and ctx.faithful:  # the blanks and line breaks before it, in the
                # paragraph's own type: their width decides whether the next picture still fits
                layout = block.layout or ParagraphLayout()
                lead = Run(text=block.lead, size_pt=layout.mark_size_pt, font=layout.mark_font)
                ctx.write_runs(paragraph, (lead,), style)
        data = ctx.extras.images.get(block.id)
        if data is not None:
            run = paragraph.add_run()
            picture = add_picture(paragraph, run, data, block.width_mm, block.height_mm)
            if picture is None:
                run._r.getparent().remove(run._r)
            else:
                if ctx.faithful and block.crop:
                    set_crop(picture._inline, block.crop)
                if ctx.faithful and block.effect_emu and floating is None:
                    set_effect(picture._inline, block.effect_emu)
                if floating is not None:
                    _float(picture._inline, floating)
                    if block.clear:  # the text after waits below it
                        clearing = OxmlElement("w:br")
                        clearing.set(qn("w:type"), "textWrapping")
                        clearing.set(qn("w:clear"), "all")
                        breaking = paragraph.add_run()  # as tall as its paragraph's line: Word
                        # clears only pictures the line reaches; the empty line after is 1 pt
                        layout = block.layout or ParagraphLayout()
                        breaking.font.size = Pt(
                            layout.mark_size_pt or ctx.profile.styles.line.size_pt
                        )
                        breaking.font.name = layout.mark_font
                        breaking._r.append(clearing)
                return
        if floating is not None:
            return  # a floating picture without pixels is left out rather than pushing text
        label = f"[рисунок: {block.description}]" if block.description else "[рисунок]"
        ctx.write_text(paragraph, label, style)
        if ctx.faithful and block.height_mm:
            fmt = paragraph.paragraph_format
            fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            fmt.line_spacing = Mm(block.height_mm)


def _picture_line(ctx: RenderContext) -> Paragraph | None:
    """The paragraph written last, when it holds a picture: the line the next picture joins."""
    paragraphs = getattr(ctx.docx, "paragraphs", None)
    if not paragraphs:
        return None
    last: Paragraph = paragraphs[-1]
    return last if any(True for _ in last._p.iter(qn("w:drawing"))) else None


_WRAP_ELEMENTS = {"none": "wrapNone", "square": "wrapSquare", "top_and_bottom": "wrapTopAndBottom"}


def _anchor_paragraph(ctx: RenderContext) -> Paragraph:
    """Where floating pictures, boxes and lines hang: a paragraph one twip high, shared by
    floating objects that follow each other, so they push the text down by nothing."""
    paragraph = ctx.docx.add_paragraph()
    previous = paragraph._p.getprevious()
    framed = previous is not None and previous.find(f"{qn('w:pPr')}/{qn('w:framePr')}") is not None
    if previous is not None and (_is_anchor(previous) or framed):  # a framed one: as drawn
        paragraph._p.getparent().remove(paragraph._p)  # from inside its frame
        return Paragraph(previous, paragraph._parent)
    fmt = paragraph.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Twips(_ANCHOR_TWIPS)
    paragraph.add_run().font.size = Pt(1)
    return paragraph


_ANCHOR_TWIPS = 1


def _clearing_paragraph(ctx: RenderContext) -> Paragraph:
    """Holds a floating picture the text after waits below, then a clearing line break; the
    line after the break, empty, is 1 pt: what follows it starts right below the picture."""
    paragraph = ctx.docx.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    set_mark_size(paragraph, 1)
    paragraph.add_run().font.size = Pt(1)
    return paragraph


def hang_old_style_drawings(paragraphs: list[Paragraph]) -> None:
    """Old-style (VML) drawings move from their 1-twip paragraph into the paragraph after it,
    which holds them in the original: one more paragraph between two would part their shared
    border and add its own spacing."""
    for paragraph in paragraphs:
        anchor, owner = paragraph._p, paragraph._p.getnext()
        if not _is_anchor(anchor) or owner is None or owner.tag != qn("w:p") or _is_anchor(owner):
            continue
        drawings = [r for r in anchor.findall(qn("w:r")) if r.find(qn("w:pict")) is not None]
        if not drawings:
            continue
        place = 1 if owner.find(qn("w:pPr")) is not None else 0
        for offset, run in enumerate(drawings):
            owner.insert(place + offset, run)
        if next(anchor.iter(qn("w:drawing"), qn("w:pict")), None) is None:
            anchor.getparent().remove(anchor)


def _is_anchor(element: BaseOxmlElement) -> bool:
    spacing = element.find(f"{qn('w:pPr')}/{qn('w:spacing')}")
    return (
        element.tag == qn("w:p")
        and spacing is not None
        and spacing.get(qn("w:line")) == str(_ANCHOR_TWIPS)
        and spacing.get(qn("w:lineRule")) == "exact"
        and not "".join(t.text or "" for t in element.iter(qn("w:t")))
        and next(element.iter(qn("w:br")), None) is None
    )


def _float(inline: BaseOxmlElement, placement: ImagePlacement) -> None:
    """Turn a picture in the line of text into one anchored at the original's place."""
    anchor = OxmlElement(
        "wp:anchor",
        {
            **dict(
                zip(
                    ("distT", "distB", "distL", "distR"),
                    map(str, placement.distances_emu),
                    strict=True,
                )
            ),
            "simplePos": "0",
            "relativeHeight": "251659264",
            "behindDoc": "1" if placement.behind_text else "0",
            "locked": "0",
            "layoutInCell": "1",
            "allowOverlap": "1",
        },
    )
    anchor.append(OxmlElement("wp:simplePos", {"x": "0", "y": "0"}))
    for name, relative, offset, align in (
        ("positionH", placement.horizontal_from, placement.x_mm, placement.x_align),
        ("positionV", placement.vertical_from, placement.y_mm, placement.y_align),
    ):
        position = OxmlElement(f"wp:{name}", {"relativeFrom": relative})
        if align is not None and offset is None:
            value = OxmlElement("wp:align")
            value.text = align
        else:
            value = OxmlElement("wp:posOffset")
            value.text = str(round((offset or 0.0) * _EMU_PER_MM))
        position.append(value)
        anchor.append(position)
    children = {child.tag.split("}")[1]: child for child in inline}
    anchor.append(children["extent"])
    if "effectExtent" in children:
        anchor.append(children["effectExtent"])
    wrap = OxmlElement(f"wp:{_WRAP_ELEMENTS[placement.wrap]}")
    if placement.wrap == "square":
        wrap.set("wrapText", "bothSides")
    anchor.append(wrap)
    for name in ("docPr", "cNvGraphicFramePr", "graphic"):
        if name in children:
            anchor.append(children[name])
    inline.getparent().replace(inline, anchor)


_EMU_PER_MM = 36000


class SpacerRenderer:
    def render(self, block: SpacerBlock, ctx: RenderContext) -> None:
        if ctx.faithful and block.lines:
            for line in block.lines:
                style = ctx.styled(ctx.profile.styles.line, line.layout)
                paragraph = ctx.add_paragraph(style, first_line_indent_mm=0)
                set_mark_size(paragraph, line.size_pt)
                if line.font:
                    set_mark_font(paragraph, line.font)
                if line.breaks:  # each break one more line, as tall as the mark's
                    breaking = paragraph.add_run()
                    breaking.font.size = Pt(line.size_pt)
                    if line.font:
                        breaking.font.name = line.font
                    for _ in range(line.breaks):
                        breaking.add_break()
                    if line.clear is not None:  # what follows waits below the floats
                        for brk in breaking._r.findall(qn("w:br")):
                            brk.set(qn("w:type"), "textWrapping")
                            brk.set(qn("w:clear"), line.clear)
            return
        if ctx.faithful and block.height_pt:
            _gap(ctx, block.height_pt)
            return
        style = ctx.profile.styles.line
        for _ in range(_SPACER_LINES[block.size]):
            ctx.add_paragraph(style, first_line_indent_mm=0)


class TableRenderer:
    def render(self, block: TableBlock, ctx: RenderContext) -> None:
        placements = block.placements()
        n_cols = block.column_count()
        if not block.rows or n_cols == 0:
            return
        table = ctx.docx.add_table(rows=len(block.rows), cols=n_cols)
        if block.grid:
            table.style = "Table Grid"
        table.autofit = False
        total = block.width_mm if ctx.faithful and block.width_mm else ctx.content_width_mm
        self._set_widths(table, self._widths(block, n_cols, total))
        for index, source_row in enumerate(block.rows):
            if ctx.faithful and source_row.height_mm:
                table.rows[index].height = Mm(source_row.height_mm)
                table.rows[index].height_rule = (
                    WD_ROW_HEIGHT_RULE.EXACTLY
                    if source_row.exact_height and not ctx.extras.editable
                    else WD_ROW_HEIGHT_RULE.AT_LEAST
                )
        if ctx.faithful:
            _place_table(table, block)

        style = ctx.profile.styles.table
        for placement in placements:
            row, cell = placement.row, placement.cell
            target = self._merged(table, row, placement.col, cell)
            _set_cell_properties(ctx, target, cell)
            cell_style = ctx.styled(style, cell.layout)
            measured = ctx.faithful and cell.layout is not None
            bold_header = row < block.header_rows and not ctx.faithful  # copy: as printed
            runs = tuple(r.model_copy(update={"bold": r.bold or bold_header}) for r in cell.runs)
            lines = _cell_paragraphs(runs)
            marks = [cell.size_pt, *(r.size_pt or cell.size_pt for r in runs if r.text == "\n")]
            for index, line in enumerate(lines):
                self._nested(ctx, target, cell, index)
                paragraph = target.paragraphs[0] if index == 0 else target.add_paragraph()
                own = cell.line_layouts[index] if index < len(cell.line_layouts) else None
                line_style = ctx.styled(style, own) if own is not None else cell_style
                align = cell.line_aligns[index] if index < len(cell.line_aligns) else cell.align
                ctx.format_paragraph(
                    paragraph, line_style, align, first_line_indent_mm=None if measured else 0
                )
                alike = own if own is not None else cell.layout  # all lines alike: the cell's
                mark = (alike.mark_size_pt if alike is not None else None) or marks[index]
                if ctx.faithful and mark:  # an empty line is as tall as its own mark
                    set_mark_size(paragraph, mark)
                ctx.write_runs(paragraph, line, style)
                _cell_pictures(ctx, paragraph, cell, index)
            self._nested(ctx, target, cell, len(lines))
            _format_merged_rest(ctx, table, placement)
        if ctx.faithful:
            self._collapse_gaps(table, placements)
        _separator(ctx)

    def _nested(self, ctx: RenderContext, target: _Cell, cell: TableCell, line: int) -> None:
        """Tables that stand before this line of the cell (after the last line: at the end)."""
        last = line >= len(_cell_paragraphs(cell.runs))
        for nested in cell.tables:
            if nested.after_line != line and not (last and nested.after_line > line):
                continue
            self.render(nested.table, replace(ctx, docx=target))
            tc = target._tc
            tc.remove(tc.p_lst[-1])  # our separator
            if not last:  # the cell's own next line closes the table, not python-docx's one
                tc.remove(tc.p_lst[-1])
            if line == 0 and not last:
                target.paragraphs[0]._p.addprevious(tc.tbl_lst[-1])

    @staticmethod
    def _collapse_gaps(table: Table, placements: list[CellPlacement]) -> None:
        """Rows shorter than the grid (Word's gridBefore/gridAfter) have no cell there at all;
        the filler cells python-docx creates must not add height or lines."""
        covered = {
            (r, c)
            for p in placements
            for r in range(p.row, p.row + p.cell.row_span)
            for c in range(p.col, p.col + p.cell.col_span)
        }
        for r, row in enumerate(table.rows):
            cells = row.cells
            for c, cell in enumerate(cells):
                if (r, c) in covered:
                    continue
                _set_cell_borders(cell, _NO_LINES)
                filler = cell.paragraphs[0].paragraph_format
                filler.space_before = filler.space_after = Pt(0)
                filler.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                filler.line_spacing = Pt(1)
                set_mark_size(cell.paragraphs[0], 1)
        for r, row in enumerate(table.rows):
            TableRenderer._trim_row(row._tr, r, covered)

    @staticmethod
    def _trim_row(tr: CT_Row, r: int, covered: set[tuple[int, int]]) -> None:
        """Filler cells at the start and end of a short row give way to gridBefore/gridAfter,
        as Word writes such a row: Word refuses a row of more than 63 cells."""
        cells = list(tr.tc_lst)
        starts, col = [], 0
        for tc in cells:
            starts.append(col)
            col += tc.grid_span
        lead = next((i for i, at in enumerate(starts) if (r, at) in covered), len(cells))
        if lead == len(cells):
            return  # a row of fillers only keeps its cells
        last = max(i for i, at in enumerate(starts) if (r, at) in covered)
        before, after = cells[:lead], cells[last + 1 :]
        for tc in (*before, *after):
            tr.remove(tc)
        added = []  # gridBefore, gridAfter, then their widths: first in trPr, in this order
        for name, gone in (("w:gridBefore", before), ("w:gridAfter", after)):
            if gone:
                added.append(OxmlElement(name, {qn("w:val"): str(sum(t.grid_span for t in gone))}))
        for name, gone in (("w:wBefore", before), ("w:wAfter", after)):
            if gone:  # without the width Word lays the row out anew
                twips = sum(t.width.twips for t in gone if t.width is not None)
                added.append(OxmlElement(name, {qn("w:w"): str(twips), qn("w:type"): "dxa"}))
        for index, element in enumerate(added):
            tr.get_or_add_trPr().insert(index, element)

    @staticmethod
    def _merged(table: Table, row: int, col: int, cell: TableCell) -> _Cell:
        last_row = min(row + cell.row_span, len(table.rows)) - 1
        last_col = min(col + cell.col_span, len(table.columns)) - 1
        start = table.cell(row, col)
        if (last_row, last_col) == (row, col):
            return start
        merged: _Cell = start.merge(table.cell(last_row, last_col))
        return merged

    @staticmethod
    def _widths(block: TableBlock, n_cols: int, total_mm: float) -> list[float]:
        weights = list(block.column_widths or ())
        if len(weights) != n_cols or any(w <= 0 for w in weights):
            weights = [1.0] * n_cols
        scale = total_mm / sum(weights)
        return [w * scale for w in weights]

    @staticmethod
    def _set_widths(table: Table, widths_mm: list[float]) -> None:
        for column, width in zip(table._tbl.tblGrid.gridCol_lst, widths_mm, strict=True):
            column.w = Mm(width)
        for row in table.rows:
            for index, cell in enumerate(row.cells):
                cell.width = Mm(widths_mm[index])


_V_ALIGN = {
    VerticalAlignment.TOP: WD_CELL_VERTICAL_ALIGNMENT.TOP,
    VerticalAlignment.CENTER: WD_CELL_VERTICAL_ALIGNMENT.CENTER,
    VerticalAlignment.BOTTOM: WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
}
_TABLE_ALIGN = {
    Alignment.LEFT: WD_TABLE_ALIGNMENT.LEFT,
    Alignment.JUSTIFY: WD_TABLE_ALIGNMENT.LEFT,
    Alignment.CENTER: WD_TABLE_ALIGNMENT.CENTER,
    Alignment.RIGHT: WD_TABLE_ALIGNMENT.RIGHT,
}
_TBL_PR_ORDER = (
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd",
    "tblLayout", "tblCellMar", "tblLook",
)  # fmt: skip
"""Child order the schema prescribes for w:tblPr; Word rejects files that break it."""


def _place_table(table: Table, block: TableBlock) -> None:
    """Position and cell padding measured on the original."""
    if block.align is not None:
        table.alignment = _TABLE_ALIGN[block.align]
    table.autofit = block.autofit
    tbl_pr = table._tbl.tblPr
    if block.autofit and block.width_mm:
        width = OxmlElement("w:tblW")
        width.set(qn("w:w"), str(round(block.width_mm * _TWIPS_PER_MM)))
        width.set(qn("w:type"), "dxa")
        _put(tbl_pr, width)
    for tag, value in (("tblInd", block.indent_mm), ("tblCellSpacing", block.cell_spacing_mm)):
        if value is not None:
            element = OxmlElement(f"w:{tag}")
            element.set(qn("w:w"), str(round(value * _TWIPS_PER_MM)))
            element.set(qn("w:type"), "dxa")
            _put(tbl_pr, element)
    if block.position is not None:
        _put(tbl_pr, _position_element(block.position))
    if block.padding is not None:
        margins = OxmlElement("w:tblCellMar")
        for side in ("top", "left", "bottom", "right"):
            edge = OxmlElement(f"w:{side}")
            edge.set(qn("w:w"), str(round(getattr(block.padding, f"{side}_mm") * _TWIPS_PER_MM)))
            edge.set(qn("w:type"), "dxa")
            margins.append(edge)
        _put(tbl_pr, margins)


_POSITION = {
    "horizontal_anchor": "horzAnchor",
    "vertical_anchor": "vertAnchor",
    "x_align": "tblpXSpec",
    "y_align": "tblpYSpec",
}
_POSITION_MM = {
    "gap_left_mm": "leftFromText",
    "gap_right_mm": "rightFromText",
    "gap_top_mm": "topFromText",
    "gap_bottom_mm": "bottomFromText",
    "x_mm": "tblpX",
    "y_mm": "tblpY",
}


def _position_element(position: TablePosition) -> BaseOxmlElement:
    """w:tblpPr: the table floats at this place and text flows around it."""
    element = OxmlElement("w:tblpPr")
    for field, name in _POSITION_MM.items():
        value = getattr(position, field)
        if value is not None:
            element.set(qn(f"w:{name}"), str(round(value * _TWIPS_PER_MM)))
    for field, name in _POSITION.items():
        value = getattr(position, field)
        if value is not None:
            element.set(qn(f"w:{name}"), value)
    return element


def _put(tbl_pr: BaseOxmlElement, child: BaseOxmlElement) -> None:
    """Replace or insert a w:tblPr child at its schema position."""
    name = child.tag.split("}")[1]
    existing = tbl_pr.find(qn(f"w:{name}"))
    if existing is not None:
        tbl_pr.remove(existing)
    later = _TBL_PR_ORDER[_TBL_PR_ORDER.index(name) + 1 :]
    successor = next((e for e in tbl_pr if e.tag.split("}")[1] in later), None)
    if successor is None:
        tbl_pr.append(child)
    else:
        successor.addprevious(child)


_TWIPS_PER_MM = 1440 / 25.4
_NO_LINES = CellBorders(top=False, bottom=False, left=False, right=False)


def _cell_paragraphs(runs: tuple[Run, ...]) -> list[tuple[Run, ...]]:
    """A run that is just a line break separates the paragraphs of a cell."""
    paragraphs: list[list[Run]] = [[]]
    for run in runs:
        if run.text == "\n":
            paragraphs.append([])
        else:
            paragraphs[-1].append(run)
    return [tuple(p) for p in paragraphs]


def _set_cell_properties(ctx: RenderContext, target: _Cell, cell: TableCell) -> None:
    if cell.borders is not None:
        _set_cell_borders(target, cell.borders)
    if ctx.faithful and cell.v_align is not None:
        target.vertical_alignment = _V_ALIGN[cell.v_align]
    if cell.shading is not None:
        shading = OxmlElement("w:shd")
        for name, value in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", cell.shading)):
            shading.set(qn(name), value)
        target._tc.get_or_add_tcPr().insert_element_before(shading, *_AFTER_SHADING)
    if ctx.faithful and cell.padding is not None:
        margins = OxmlElement("w:tcMar")
        for side in ("top", "left", "bottom", "right"):
            edge = OxmlElement(f"w:{side}")
            edge.set(qn("w:w"), str(round(getattr(cell.padding, f"{side}_mm") * _TWIPS_PER_MM)))
            edge.set(qn("w:type"), "dxa")
            margins.append(edge)
        target._tc.get_or_add_tcPr().insert_element_before(margins, *_AFTER_SHADING[2:])
    if ctx.faithful and cell.text_direction is not None:
        turned = OxmlElement("w:textDirection", {qn("w:val"): cell.text_direction})
        target._tc.get_or_add_tcPr().insert_element_before(turned, *_AFTER_SHADING[3:])
    if ctx.faithful and cell.hide_mark:
        tc_pr = target._tc.get_or_add_tcPr()
        tc_pr.insert_element_before(
            OxmlElement("w:hideMark"), "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge"
        )


_AFTER_SHADING = (
    "w:noWrap",
    "w:tcMar",
    "w:textDirection",
    "w:tcFitText",
    "w:vAlign",
    "w:hideMark",
    "w:headers",
    "w:cellIns",
    "w:cellDel",
    "w:cellMerge",
    "w:tcPrChange",
)
"""What follows w:shd in a cell's properties: Word wants them in the schema's order."""


def drop_needless_separators(root: BaseOxmlElement) -> None:
    """A separator stays between two tables and where a body, cell or margin would otherwise
    end with a table; before a paragraph it is only a 1-pt line the original does not have."""
    for paragraph in list(root.iter(qn("w:p"))):
        if paragraph.get(SEPARATOR_MARK) is None:
            continue
        del paragraph.attrib[SEPARATOR_MARK]
        following = paragraph.getnext()
        if following is not None and following.tag == qn("w:p"):
            paragraph.getparent().remove(paragraph)


def _cell_pictures(ctx: RenderContext, paragraph: Paragraph, cell: TableCell, line: int) -> None:
    """The cell's pictures of this line, put before its text as in the original."""
    if not ctx.faithful:
        return
    for image in cell.images:
        data = ctx.extras.images.get(image.image_id)
        if image.line != line or data is None:
            continue
        run = paragraph.add_run()
        picture = add_picture(paragraph, run, data, image.width_mm, image.height_mm)
        if picture is None:
            run._r.getparent().remove(run._r)
            continue
        if image.crop:
            set_crop(picture._inline, image.crop)
        properties = paragraph._p.pPr
        if properties is not None:
            properties.addnext(run._r)
        else:
            paragraph._p.insert(0, run._r)
        if image.floating is not None:
            _float(picture._inline, image.floating)


def _format_merged_rest(ctx: RenderContext, table: Table, placement: CellPlacement) -> None:
    """The rows a vertically merged cell continues into keep an empty single-spaced line that
    still takes height, as large as in the original."""
    cell = placement.cell
    if not ctx.faithful:
        return
    rows = table._tbl.tr_lst[placement.row + 1 : placement.row + cell.row_span]
    for tr, size_pt in zip(rows, cell.rest_sizes_pt, strict=False):
        start = 0
        for tc in tr.tc_lst:
            if start == placement.col:
                for p in tc.p_lst:
                    paragraph = Paragraph(p, table)
                    paragraph_format = paragraph.paragraph_format
                    paragraph_format.space_before = paragraph_format.space_after = Pt(0)
                    paragraph_format.line_spacing = 1.0
                    set_mark_size(paragraph, size_pt)
                break
            start += tc.grid_span
    if cell.borders is not None and rows:
        _merged_borders(table, placement, cell.borders)


def _merged_borders(table: Table, placement: CellPlacement, borders: CellBorders) -> None:
    """Word draws a vertically merged cell's side lines row by row from the cell of each row,
    its top from the first and its bottom from the last (6488ab1f: the sides were missing
    below the first row)."""
    rows = table._tbl.tr_lst[placement.row : placement.row + placement.cell.row_span]
    for index, tr in enumerate(rows):
        start = 0
        for tc in tr.tc_lst:
            if start == placement.col:
                _set_cell_borders(
                    _Cell(tc, table),
                    borders.model_copy(
                        update={
                            "top": borders.top and index == 0,
                            "bottom": borders.bottom and index == len(rows) - 1,
                        }
                    ),
                )
                break
            start += tc.grid_span


_BLACK_LINE = BorderLine()


def _set_cell_borders(cell: _Cell, borders: CellBorders) -> None:
    """Explicit line or no line on each side, overriding the table style."""
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:tcBorders"))
    if existing is not None:
        tc_pr.remove(existing)
    element = OxmlElement("w:tcBorders")
    lines = dict(borders.lines)
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        if getattr(borders, side):
            line = lines.get(side, _BLACK_LINE)
            edge.set(qn("w:val"), line.style)
            edge.set(qn("w:sz"), str(line.size))
            edge.set(qn("w:color"), line.color or _default_colour(line.style))
        else:
            edge.set(qn("w:val"), "nil")
        element.append(edge)
    tc_pr.append(element)


_SHADED_STYLES = frozenset({"outset", "inset", "threeDEmboss", "threeDEngrave"})


def _default_colour(style: str) -> str:
    """Black, but 'auto' for the raised and sunken lines: only then Word draws them in its light
    and dark greys, as the original's web-page tables (6717f1b5); in black they are thick bars."""
    return "auto" if style in _SHADED_STYLES else "000000"


def _gap(ctx: RenderContext, height_pt: float) -> None:
    """One empty paragraph exactly as tall as the measured gap."""
    paragraph = ctx.docx.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.first_line_indent = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(height_pt)


SEPARATOR_MARK = "f2w-separator"
"""Marks the paragraph after a table until the file is saved; drop_needless_separators then
keeps it only where Word needs one."""


def _separator(ctx: RenderContext) -> None:
    """Word merges two adjacent tables; a 1-pt paragraph keeps them apart almost invisibly."""
    paragraph = ctx.docx.add_paragraph()
    paragraph._p.set(SEPARATOR_MARK, "1")
    fmt = paragraph.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(1)
    paragraph.add_run().font.size = Pt(1)


class FillRowRenderer:
    """An invisible table: the value sits on a cell with a bottom line, the caption under it
    in small type. The blank stays a blank when the user types into it in Word."""

    CAPTION_SHARE = 0.65

    def __init__(self, tables: TableRenderer) -> None:
        self._tables = tables

    def render(self, block: FillRowBlock, ctx: RenderContext) -> None:
        style = ctx.styled(ctx.profile.styles.line, block.layout)
        size = block.size_pt if ctx.faithful and block.size_pt else style.size_pt
        left = style.left_indent_mm if ctx.faithful and block.layout else 0.0
        right = style.right_indent_mm if ctx.faithful and block.layout else 0.0
        table = self.as_table(block, size, max(ctx.content_width_mm - left - right, 20.0))
        self._tables.render(table.model_copy(update={"indent_mm": left or None}), ctx)

    def as_table(self, block: FillRowBlock, size_pt: float, width_mm: float) -> TableBlock:
        caption_size = max(round(size_pt * self.CAPTION_SHARE * 2) / 2, 6.0)
        has_captions = any(slot.caption for slot in block.slots)
        top: list[TableCell] = []
        bottom: list[TableCell] = []
        widths: list[float] = []
        labels_mm = sum(self._label_mm(slot.label, size_pt) for slot in block.slots)
        lines_mm = max(width_mm - labels_mm, width_mm * 0.3)
        weight = sum(max(slot.width, 0.1) for slot in block.slots)
        for slot in block.slots:
            if slot.label:
                top.append(TableCell(runs=(Run(text=slot.label, size_pt=size_pt),), borders=_NONE))
                bottom.append(TableCell(borders=_NONE))
                widths.append(self._label_mm(slot.label, size_pt))
            value = tuple(
                r if r.size_pt else r.model_copy(update={"size_pt": size_pt}) for r in slot.value
            )
            top.append(
                TableCell(
                    runs=value,
                    align=Alignment.CENTER,
                    borders=_UNDERLINE,
                    v_align=VerticalAlignment.BOTTOM,
                    size_pt=size_pt,
                )
            )
            caption = (Run(text=slot.caption, size_pt=caption_size),) if slot.caption else ()
            bottom.append(TableCell(runs=caption, align=Alignment.CENTER, borders=_NONE))
            widths.append(lines_mm * max(slot.width, 0.1) / weight)
        rows = (TableRow(cells=tuple(top)),) + (
            (TableRow(cells=tuple(bottom)),) if has_captions else ()
        )
        return TableBlock(
            id=block.id,
            rows=rows,
            column_widths=tuple(round(w, 1) for w in widths),
            width_mm=round(sum(widths), 1),
            grid=False,
            padding=CellPadding(left_mm=0.5, right_mm=0.5),
        )

    @staticmethod
    def _label_mm(label: str, size_pt: float) -> float:
        """Width of the label at this size: capitals are wider than small letters."""
        if not label:
            return 0.0
        ems = sum(_CHAR_EM.get(_char_class(ch), 0.5) for ch in label)
        return ems * size_pt * 25.4 / 72 + 3.0


_CHAR_EM = {"upper": 0.72, "lower": 0.56, "digit": 0.56, "space": 0.28, "other": 0.4}
"""Character widths in font sizes, taken generously (Arial is wider than Times): a label that
wraps breaks the row, a little extra width does not."""


def _char_class(ch: str) -> str:
    if ch.isupper():
        return "upper"
    if ch.islower():
        return "lower"
    if ch.isdigit():
        return "digit"
    return "space" if ch.isspace() else "other"


_NONE = CellBorders(top=False, bottom=False, left=False, right=False)
_UNDERLINE = CellBorders(top=False, bottom=True, left=False, right=False)


class TextBoxRenderer:
    """A floating text box at the original's place; its blocks are written by the usual
    renderers and then moved into the box. Outside faithful mode they stay in the flow."""

    def __init__(self, registry: Mapping[type, BlockRenderer[Any]]) -> None:
        self._registry = registry

    def render(self, block: TextBoxBlock, ctx: RenderContext) -> None:
        if not ctx.faithful:
            for inner in block.blocks:
                self._registry[type(inner)].render(inner, ctx)
            return
        anchor = _anchor_paragraph(ctx)
        for inner in block.blocks:
            self._registry[type(inner)].render(inner, ctx)
        written = [e for e in anchor._p.itersiblings() if e.tag != qn("w:sectPr")]
        grows = ctx.extras.editable or block.grows
        box = _text_box_element(block, next_drawing_id(anchor), grows=grows)
        content = next(box.iter(qn("w:txbxContent")))
        for element in written:
            content.append(element)
        if not written or written[-1].tag != qn("w:p"):
            content.append(OxmlElement("w:p"))  # a box ends with a paragraph
        run = anchor.add_run()
        run._r.append(box)


_WPS_NS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
_WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _emu(mm: float) -> str:
    return str(round(mm * _EMU_PER_MM))


def _text_box_element(block: TextBoxBlock, drawing_id: int, *, grows: bool) -> BaseOxmlElement:
    """w:drawing with an anchored DrawingML text box; its content is filled in afterwards."""
    padding = block.padding or CellPadding(left_mm=2.54, right_mm=2.54, top_mm=1.27, bottom_mm=1.27)
    outline = (
        '<a:ln w="6350"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln>'
        if block.framed
        else "<a:ln><a:noFill/></a:ln>"
    )
    fit = "<a:spAutoFit/>" if grows else "<a:noAutofit/>"  # grow with the text or not
    fill = (
        f'<a:solidFill><a:srgbClr val="{block.fill}"/></a:solidFill>'
        if block.fill
        else "<a:noFill/>"
    )
    rot = f' rot="{round(block.rotation_deg * 60000)}"' if block.rotation_deg else ""
    drawing = parse_xml(
        f'<w:drawing xmlns:w="{_W_NS}" xmlns:wp="{_WP_NS}" xmlns:a="{_A_NS}" '
        f'xmlns:wps="{_WPS_NS}"><wp:inline><wp:extent cx="{_emu(block.width_mm)}" '
        f'cy="{_emu(block.height_mm)}"/><wp:docPr id="{drawing_id}" name="Text Box {drawing_id}"/>'
        f'<a:graphic><a:graphicData uri="{_WPS_NS}"><wps:wsp><wps:cNvSpPr txBox="1"/>'
        f'<wps:spPr><a:xfrm{rot}><a:off x="0" y="0"/><a:ext cx="{_emu(block.width_mm)}" '
        f'cy="{_emu(block.height_mm)}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f"{fill}{outline}</wps:spPr><wps:txbx><w:txbxContent/></wps:txbx>"
        f'<wps:bodyPr rot="0" vert="{block.text_direction}" wrap="square" '
        f'lIns="{_emu(padding.left_mm)}" tIns="{_emu(padding.top_mm)}" '
        f'rIns="{_emu(padding.right_mm)}" bIns="{_emu(padding.bottom_mm)}" '
        f'anchor="{block.text_anchor}">{fit}</wps:bodyPr>'
        "</wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing>"
    )
    _float(drawing[0], block.floating)
    return drawing


_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


class ShapeRenderer:
    """A drawn line or frame at the original's place; nothing outside faithful mode."""

    def render(self, block: ShapeBlock, ctx: RenderContext) -> None:
        if not ctx.faithful:
            return
        anchor = _anchor_paragraph(ctx)
        if block.vml:  # an old-style shape: as Word wrote it
            anchor.add_run()._r.append(parse_xml(block.vml))
            return
        flips = "".join(
            f' {name}="1"' for name, on in (("flipH", block.flip_h), ("flipV", block.flip_v)) if on
        )
        drawing_id = next_drawing_id(anchor)
        ends = "".join(
            f'<a:{name} type="{kind}"/>'
            for name, kind in (("headEnd", block.head), ("tailEnd", block.tail))
            if kind != "none"
        )
        outline = (
            _without_namespaces(block.path)
            if block.geometry == "custom" and block.path
            else f'<a:prstGeom prst="{block.geometry}"><a:avLst/></a:prstGeom>'
        )
        inside = (
            f'<a:solidFill><a:srgbClr val="{block.fill}"/></a:solidFill>'
            if block.fill
            else "<a:noFill/>"
        )
        stroke = (
            f'<a:solidFill><a:srgbClr val="{block.color or "000000"}"/></a:solidFill>'
            if block.line_pt
            else "<a:noFill/>"
        )
        drawing = parse_xml(
            f'<w:drawing xmlns:w="{_W_NS}" xmlns:wp="{_WP_NS}" xmlns:a="{_A_NS}" '
            f'xmlns:wps="{_WPS_NS}"><wp:inline><wp:extent cx="{_emu(block.width_mm)}" '
            f'cy="{_emu(block.height_mm)}"/><wp:docPr id="{drawing_id}" name="Shape {drawing_id}"/>'
            f'<a:graphic><a:graphicData uri="{_WPS_NS}"><wps:wsp><wps:cNvSpPr/><wps:spPr>'
            f'<a:xfrm{flips}><a:off x="0" y="0"/><a:ext cx="{_emu(block.width_mm)}" '
            f'cy="{_emu(block.height_mm)}"/></a:xfrm>{outline}{inside}'
            f'<a:ln w="{round(block.line_pt * 12700)}">{stroke}{ends}'
            "</a:ln></wps:spPr><wps:bodyPr/>"
            "</wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing>"
        )
        _float(drawing[0], block.floating)
        anchor.add_run()._r.append(drawing)


def _without_namespaces(fragment: str) -> str:
    """An a:custGeom as stored, its namespace declaration left to the drawing around it."""
    return re.sub(r' xmlns:\w+="[^"]*"', "", fragment)


def default_registry() -> Mapping[type, BlockRenderer[Any]]:
    tables = TableRenderer()
    registry: dict[type, BlockRenderer[Any]] = {
        HeadingBlock: HeadingRenderer(),
        ParagraphBlock: ParagraphRenderer(),
        ListItemBlock: ListItemRenderer(),
        TableBlock: tables,
        LabeledLineBlock: LabeledLineRenderer(),
        TwoColumnLineBlock: TwoColumnLineRenderer(),
        SignatureLineBlock: SignatureLineRenderer(),
        RequisitesBlock: RequisitesRenderer(),
        SpacerBlock: SpacerRenderer(),
        ImageBlock: ImageRenderer(),
        FillRowBlock: FillRowRenderer(tables),
    }
    registry[TextBoxBlock] = TextBoxRenderer(registry)
    registry[ShapeBlock] = ShapeRenderer()
    return registry
