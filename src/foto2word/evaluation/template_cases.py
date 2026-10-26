"""Eval cases from real Word templates: the source .docx is the ground truth.

For every template: strip site advertising, render the pages as the "photo", import the
.docx into a DocModel as the expected answer, and measure the round trip
"original .docx -> DocModel -> our .docx" against the original rendering. The round trip
needs no model at all and is the ceiling for what recognition can achieve.
"""

from collections.abc import Callable
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

import docx
import pypdfium2
from PIL import Image, ImageChops, ImageFilter, ImageOps
from rapidfuzz import fuzz

from foto2word.application.ports import DocumentRenderer, Previewer
from foto2word.domain.extras import RenderExtras
from foto2word.domain.profile import Profile, RenderMode
from foto2word.evaluation.cases import EXPECTED_FILE
from foto2word.infrastructure.importing.docx_importer import AD_PHRASES, DocxImporter

_THUMB = (300, 424)


def strip_ads(data: bytes) -> bytes:
    """Remove paragraphs that advertise the source site, in the body, headers and footers."""
    document = docx.Document(BytesIO(data))
    parts: list[Any] = [document]
    for section in document.sections:
        parts += [
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
        ]
    for part in parts:
        for paragraph in list(part.paragraphs):
            if any(phrase in paragraph.text.lower() for phrase in AD_PHRASES):
                element = paragraph._element
                element.getparent().remove(element)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pdf_text(pdf: bytes) -> str:
    document = pypdfium2.PdfDocument(pdf)
    try:
        return "\n".join(page.get_textpage().get_text_range() for page in document)
    finally:
        document.close()


def visual_similarity(a: Image.Image, b: Image.Image) -> float:
    """1.0 = identical page layout at thumbnail scale; ignores small font differences."""
    small_a = ImageOps.grayscale(a).resize(_THUMB)
    small_b = ImageOps.grayscale(b).resize(_THUMB)
    diff = ImageChops.difference(small_a, small_b)
    histogram = diff.histogram()
    mean = sum(i * count for i, count in enumerate(histogram)) / (_THUMB[0] * _THUMB[1])
    return max(0.0, 1.0 - mean / 64)


_INK_WIDTH = 600
_INK_THRESHOLD = 170


def _ink(image: Image.Image) -> Image.Image:
    gray = ImageOps.grayscale(image)
    gray = gray.resize((_INK_WIDTH, round(gray.height * _INK_WIDTH / gray.width)))
    return gray.point(lambda v: 255 if v < _INK_THRESHOLD else 0)


def ink_f1(original: Image.Image, ours: Image.Image, tolerance_px: int = 5) -> float:
    """Do the dark pixels land in the same places? F1 of ink with a small shift tolerance.

    Unlike plain pixel difference, a blank page does not look 'similar' to a filled one.
    """
    a, b = _ink(original), _ink(ours.resize(original.size))
    size = 2 * tolerance_px + 1
    a_near, b_near = a.filter(ImageFilter.MaxFilter(size)), b.filter(ImageFilter.MaxFilter(size))

    def count(image: Image.Image) -> int:
        return image.histogram()[255]

    ink_a, ink_b = count(a), count(b)
    if ink_a == 0 or ink_b == 0:
        return 1.0 if ink_a == ink_b else 0.0
    precision = count(ImageChops.multiply(b, a_near)) / ink_b
    recall = count(ImageChops.multiply(a, b_near)) / ink_a
    return 0.0 if precision + recall == 0 else 2 * precision * recall / (precision + recall)


@dataclass(frozen=True, slots=True)
class RoundTrip:
    case_id: str
    pages_original: int
    pages_ours: int
    text_similarity: float
    visual_similarity: float
    ink_f1: float


class TemplateCaseBuilder:
    def __init__(
        self,
        previewer: Previewer,
        renderer: DocumentRenderer,
        profile_for: Callable[[str], Profile],
        importer: DocxImporter | None = None,
        *,
        dpi: int = 110,
        max_pages: int | None = None,
    ) -> None:
        """max_pages limits the pages turned into images (long documents); all are counted."""
        self._previewer = previewer
        self._renderer = renderer
        self._profile_for = profile_for
        self._importer = importer or DocxImporter()
        self._dpi = dpi
        self._max_pages = max_pages

    def build(self, source: Path, case_dir: Path) -> RoundTrip:
        case_dir.mkdir(parents=True, exist_ok=True)
        clean = strip_ads(source.read_bytes())
        (case_dir / "source.docx").write_bytes(clean)
        original_pdf = self._previewer.to_pdf(clean)
        original_count, originals = _pages(original_pdf, self._dpi, self._max_pages)
        for number, page in enumerate(originals, start=1):
            page.save(case_dir / f"photo-{number}.jpg", quality=90)

        expected = self._importer.convert(docx.Document(BytesIO(clean)))
        (case_dir / EXPECTED_FILE).write_text(expected.model_dump_json(indent=2), "utf-8")

        ours_docx = self._renderer.render(
            expected,
            self._profile_for(expected.document_type.id),
            RenderMode.FAITHFUL,
            RenderExtras(images=self._importer.images),
        )
        (case_dir / "roundtrip.docx").write_bytes(ours_docx)
        ours_pdf = self._previewer.to_pdf(ours_docx)
        ours_count, ours = _pages(ours_pdf, self._dpi, self._max_pages)
        side_by_side(originals[0], ours[0]).save(case_dir / "roundtrip-compare.jpg", quality=85)
        return RoundTrip(
            case_id=case_dir.name,
            pages_original=original_count,
            pages_ours=ours_count,
            text_similarity=fuzz.ratio(_squash(pdf_text(original_pdf)), _squash(pdf_text(ours_pdf)))
            / 100,
            visual_similarity=visual_similarity(originals[0], ours[0]),
            ink_f1=ink_f1(originals[0], ours[0]),
        )


def side_by_side(left: Image.Image, right: Image.Image) -> Image.Image:
    height = max(left.height, right.height)
    canvas = Image.new("RGB", (left.width + right.width + 20, height), "#888")
    canvas.paste(left, (0, 0))
    canvas.paste(right, (left.width + 20, 0))
    return canvas


def _pages(pdf: bytes, dpi: int, limit: int | None) -> tuple[int, list[Image.Image]]:
    """Page count and images of the first `limit` pages (all when None)."""
    document = pypdfium2.PdfDocument(pdf)
    try:
        shown = len(document) if limit is None else min(len(document), limit)
        images = [document[i].render(scale=dpi / 72).to_pil().convert("RGB") for i in range(shown)]
        return len(document), images
    finally:
        document.close()


def _squash(text: str) -> str:
    return " ".join(text.split())
