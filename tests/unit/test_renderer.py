from io import BytesIO

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Length

from foto2word.domain.corrections import CorrectionApplier, accept_all
from foto2word.domain.document import (
    Alignment,
    Document,
    DocumentTypeGuess,
    Orientation,
    ParagraphBlock,
    Run,
    TableBlock,
    TableCell,
    TableRow,
)
from foto2word.domain.profile import RenderMode
from foto2word.infrastructure.profiles.yaml_repository import YamlProfileRepository
from foto2word.infrastructure.rendering.docx_renderer import DocxRenderer


def _render(
    document: Document, profiles: YamlProfileRepository, mode: RenderMode = RenderMode.FAITHFUL
) -> docx.document.Document:
    profile = profiles.for_document_type(document.document_type.id)
    return docx.Document(BytesIO(DocxRenderer().render(document, profile, mode)))


def _mm(length: Length | None) -> int:
    assert length is not None
    return round(float(length.mm))


def _single(*blocks: object, orientation: Orientation = Orientation.PORTRAIT) -> Document:
    return Document.model_validate(
        {
            "document_type": DocumentTypeGuess(id="other", title="x", confidence=1.0),
            "blocks": blocks,
            "orientation": orientation,
        }
    )


def test_act_renders_every_block_with_corrections_applied(
    act: Document, profiles: YamlProfileRepository
) -> None:
    corrected = CorrectionApplier().apply(act, accept_all).document

    output = _render(corrected, profiles)
    texts = [p.text for p in output.paragraphs]

    assert texts[0] == "АКТ"
    assert output.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert "г.Тестовый\t«19» октября 2027 г." in texts
    assert "Председатель комиссии\t________________\tСидоров С.С." in texts
    assert "2. При этом:" in texts
    superscripts = [r.text for p in output.paragraphs for r in p.runs if r.font.superscript]
    assert superscripts == ["2"]


def test_page_setup_comes_from_profile(act: Document, profiles: YamlProfileRepository) -> None:
    section = _render(act, profiles).sections[0]

    assert _mm(section.left_margin) == 20
    assert _mm(section.right_margin) == 10
    assert _mm(section.page_width) == 210


def test_landscape_orientation_swaps_page_size(profiles: YamlProfileRepository) -> None:
    document = _single(
        ParagraphBlock(id="b1", runs=(Run(text="x"),)), orientation=Orientation.LANDSCAPE
    )

    section = _render(document, profiles).sections[0]

    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert _mm(section.page_width) == 297


def test_faithful_mode_keeps_photo_alignment_standard_mode_does_not(
    profiles: YamlProfileRepository,
) -> None:
    document = _single(ParagraphBlock(id="b1", align=Alignment.RIGHT, runs=(Run(text="Подпись"),)))

    faithful = _render(document, profiles, RenderMode.FAITHFUL).paragraphs[0]
    standard = _render(document, profiles, RenderMode.STANDARD).paragraphs[0]

    assert faithful.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert standard.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY


def test_table_with_merged_cells(profiles: YamlProfileRepository) -> None:
    def cell(text: str, **span: int) -> TableCell:
        return TableCell(runs=(Run(text=text),), **span)

    table = TableBlock(
        id="b1",
        header_rows=1,
        rows=(
            TableRow(cells=(cell("Показатель", row_span=2), cell("Значение", col_span=2))),
            TableRow(cells=(cell("до"), cell("после"))),
            TableRow(cells=(cell("Давление"), cell("10"), cell("10"))),
        ),
    )

    output = _render(_single(table), profiles)
    grid = output.tables[0]

    assert (len(grid.rows), len(grid.columns)) == (3, 3)
    assert grid.cell(0, 0).text == grid.cell(1, 0).text == "Показатель"
    assert grid.cell(0, 1).text == grid.cell(0, 2).text == "Значение"
    assert grid.cell(2, 2).text == "10"
    # a copy keeps the header as printed; the standard document makes it bold
    assert not grid.cell(0, 1).paragraphs[0].runs[0].font.bold
    standard = _render(_single(table), profiles, RenderMode.STANDARD).tables[0]
    assert standard.cell(0, 1).paragraphs[0].runs[0].font.bold is True


def test_nested_list_items_shift_only_the_first_line(
    act: Document, profiles: YamlProfileRepository
) -> None:
    paragraphs = {p.text: p for p in _render(act, profiles).paragraphs}

    top = paragraphs["2. При этом:"].paragraph_format
    nested = next(p for t, p in paragraphs.items() if t.startswith("а) ")).paragraph_format

    assert nested.left_indent == top.left_indent
    assert _mm(nested.first_line_indent) == _mm(top.first_line_indent) + 10
