from collections.abc import Sequence
from io import BytesIO

import docx
import pytest

from foto2word.application.convert import ConversionOptions, ConvertDocument
from foto2word.application.ports import Provenance, Understanding, UnderstandingContext
from foto2word.domain.corrections import reject_all
from foto2word.domain.document import Document
from foto2word.domain.images import PageImage
from foto2word.infrastructure.profiles.yaml_repository import YamlProfileRepository
from foto2word.infrastructure.rendering.docx_renderer import DocxRenderer

PAGE = PageImage(b"jpeg", "image/jpeg", "photo.jpg")


class FakeVision:
    def __init__(self, document: Document) -> None:
        self.document = document
        self.calls: list[tuple[Sequence[PageImage], UnderstandingContext]] = []

    def understand(
        self, pages: Sequence[PageImage], context: UnderstandingContext
    ) -> Understanding:
        self.calls.append((pages, context))
        return Understanding(self.document, Provenance("fake", "fake-1", "understand.v1"))


@pytest.fixture
def vision(act: Document) -> FakeVision:
    return FakeVision(act)


@pytest.fixture
def converter(vision: FakeVision, profiles: YamlProfileRepository) -> ConvertDocument:
    return ConvertDocument(vision, DocxRenderer(), profiles)


def test_execute_understands_corrects_and_renders(
    converter: ConvertDocument, vision: FakeVision, act: Document
) -> None:
    result = converter.execute([PAGE], ConversionOptions(document_type_hint="act"))

    assert vision.calls[0][1].document_type_hint == "act"
    assert result.understood == act
    assert result.profile.id == "act"
    assert len(result.corrections.applied) == 2
    assert result.provenance is not None
    texts = [p.text for p in docx.Document(BytesIO(result.output)).paragraphs]
    assert "г.Тестовый\t«19» октября 2027 г." in texts


def test_profile_can_be_forced_and_corrections_rejected(converter: ConvertDocument) -> None:
    options = ConversionOptions(profile_id="generic", decide_correction=reject_all)

    result = converter.execute([PAGE], options)

    assert result.profile.id == "generic"
    assert result.final.blocks == result.understood.blocks


def test_rerender_does_not_call_the_model(
    converter: ConvertDocument, vision: FakeVision, application: Document
) -> None:
    result = converter.rerender(application, ConversionOptions())

    assert vision.calls == []
    assert result.profile.id == "application"
    assert result.provenance is None


def test_at_least_one_page_is_required(converter: ConvertDocument) -> None:
    with pytest.raises(ValueError, match="page"):
        converter.execute([], ConversionOptions())


def test_pictures_are_cut_from_the_photo(profiles: YamlProfileRepository) -> None:
    from PIL import Image

    from foto2word.domain.document import BoundingBox, DocumentTypeGuess, ImageBlock
    from foto2word.infrastructure.images import PillowImageCutter

    photo = BytesIO()
    Image.new("RGB", (200, 100), "white").save(photo, format="JPEG")
    page = PageImage(photo.getvalue(), "image/jpeg", "photo.jpg")
    stamp = ImageBlock(
        id="b1", description="печать", bbox=BoundingBox(x0=0.5, y0=0.5, x1=1.0, y1=1.0)
    )
    document = Document(
        document_type=DocumentTypeGuess(id="other", title="Справка", confidence=1.0),
        blocks=(stamp,),
    )
    converter = ConvertDocument(
        FakeVision(document), DocxRenderer(), profiles, cutter=PillowImageCutter()
    )

    result = converter.execute([page], ConversionOptions())

    pictures = docx.Document(BytesIO(result.output)).inline_shapes
    assert len(pictures) == 1


def test_photos_are_flattened_before_the_model_reads_them(
    act: Document, profiles: YamlProfileRepository
) -> None:
    flat = PageImage(b"flat", "image/jpeg", "photo.jpg")

    class Flattener:
        def rectify(self, page: PageImage) -> PageImage:
            return flat

    vision = FakeVision(act)
    converter = ConvertDocument(vision, DocxRenderer(), profiles, rectifier=Flattener())

    analysis = converter.analyze([PAGE], ConversionOptions())

    assert list(vision.calls[0][0]) == [flat]
    assert analysis.pages == (flat,)


def test_text_lines_are_read_while_the_model_reads_and_every_step_is_timed(
    act: Document, profiles: YamlProfileRepository
) -> None:
    import threading

    from foto2word.domain.ocr import TextLine

    lines_started = threading.Event()
    started_before_answer: list[bool] = []

    class SlowVision(FakeVision):
        def understand(
            self, pages: Sequence[PageImage], context: UnderstandingContext
        ) -> Understanding:
            started_before_answer.append(lines_started.wait(timeout=2))
            return super().understand(pages, context)

    class Lines:
        def read_lines(self, pages: Sequence[PageImage]) -> list[TextLine]:
            lines_started.set()
            return []

    class Layout:
        def estimate(
            self, document: Document, pages: Sequence[PageImage], lines: Sequence[TextLine] = ()
        ) -> Document:
            return document

    converter = ConvertDocument(
        SlowVision(act), DocxRenderer(), profiles, lines=Lines(), layout=Layout()
    )

    result = converter.execute([PAGE], ConversionOptions())

    assert started_before_answer == [True]  # read while the model was reading
    assert {"read", "lines", "layout", "render", "verify"} <= set(result.seconds)
    assert all(seconds >= 0 for seconds in result.seconds.values())
