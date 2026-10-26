import re
from io import BytesIO
from typing import Any

import docx
import pytest
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from foto2word.domain.document import (
    Alignment,
    Block,
    Document,
    HeadingBlock,
    ListItemBlock,
    ParagraphBlock,
    ParagraphLayout,
    RequisitesBlock,
    SpacerBlock,
    TableBlock,
)
from foto2word.domain.extras import RenderExtras
from foto2word.domain.profile import RenderMode
from foto2word.evaluation.template_cases import strip_ads
from foto2word.infrastructure.importing.docx_importer import DocxImporter
from foto2word.infrastructure.profiles.yaml_repository import YamlProfileRepository
from foto2word.infrastructure.rendering.docx_renderer import DocxRenderer


def _required[T](value: T | None) -> T:
    assert value is not None
    return value


def _textual(block: Block) -> HeadingBlock | ParagraphBlock | ListItemBlock:
    assert isinstance(block, HeadingBlock | ParagraphBlock | ListItemBlock)
    return block


def _list_item(block: Block) -> ListItemBlock:
    assert isinstance(block, ListItemBlock)
    return block


def _layout(block: Block) -> ParagraphLayout:
    return _required(_textual(block).layout)


def _bytes(document: docx.document.Document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _import(document: docx.document.Document) -> list[Block]:
    return list(_document(document).blocks)


def _document(source: docx.document.Document) -> Document:
    return DocxImporter().convert(docx.Document(BytesIO(_bytes(source))))


def test_typical_application_layout() -> None:
    source = docx.Document()
    for line in ("Генеральному директору", "ООО «Ромашка»", "Петрову П.П."):
        source.add_paragraph(line).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    source.add_paragraph("")
    title = source.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("ЗАЯВЛЕНИЕ").bold = True
    source.add_paragraph("Прошу предоставить отпуск.")
    source.add_paragraph("г. Москва\t\t\t\t«24» июня 2019 г.")
    source.add_paragraph("25.08.2023            Егорова Е.Е.")
    source.add_paragraph("Больше форм и образцов документов — на ppt.ru")
    linked = source.add_paragraph("В соответствии с ")
    hyperlink = linked._p.makeelement(qn("w:hyperlink"), {})
    hyperlink.append(linked.add_run("Договором")._r)
    linked._p.append(hyperlink)
    linked.add_run(" от 12 сентября")

    blocks = _import(source)

    lines = [b.text() for b in blocks[:3]]  # the addressee's lines as written
    assert lines == ["Генеральному директору", "ООО «Ромашка»", "Петрову П.П."]
    assert all(_textual(b).align is Alignment.RIGHT for b in blocks[:3])
    assert isinstance(blocks[3], SpacerBlock)
    assert isinstance(blocks[4], HeadingBlock)
    assert blocks[4].text() == "ЗАЯВЛЕНИЕ"
    assert isinstance(blocks[5], ParagraphBlock)
    assert isinstance(blocks[6], ParagraphBlock)  # its tabs as written: Word lays them out
    assert blocks[6].text() == "г. Москва				«24» июня 2019 г."
    assert isinstance(blocks[7], ParagraphBlock)
    assert len(blocks) == 9  # the advert is gone
    assert blocks[8].text() == "В соответствии с Договором от 12 сентября"


def test_a_numbered_clause_typed_by_hand_is_a_list_item_not_two_columns() -> None:
    """'1.3.<tab>text' with a hanging indent: the number and its clause, not a line with its
    right part flush with the margin."""
    source = docx.Document()
    for text in ("1.3. \tТорги проводятся на условиях", "2.1.     Претендент обязан:"):
        clause = source.add_paragraph(text)
        clause.paragraph_format.first_line_indent = Mm(-10)

    blocks = _import(source)

    assert [type(b) for b in blocks] == [ListItemBlock, ListItemBlock]
    assert [_list_item(b).marker for b in blocks] == ["1.3.", "2.1."]
    assert [b.text() for b in blocks] == ["Торги проводятся на условиях", "Претендент обязан:"]


def test_requisites_by_left_indent() -> None:
    source = docx.Document()
    for line in ("Директору", "от кассира"):
        source.add_paragraph(line).paragraph_format.left_indent = Mm(90)

    blocks = _import(source)

    assert [b.text() for b in blocks] == ["Директору", "от кассира"]  # paragraphs as written
    assert all(isinstance(b, ParagraphBlock) for b in blocks)
    assert _layout(blocks[0]).left_indent_mm == pytest.approx(90, abs=0.1)


def test_manual_and_automatic_numbering() -> None:
    source = docx.Document()
    source.add_paragraph("1. Покупатель выполнил обязательства.")
    source.add_paragraph("а) подпункт без автонумерации")
    source.add_paragraph("Первый автопункт", style="List Number")
    source.add_paragraph("Второй автопункт", style="List Number")
    source.add_paragraph("Маркированный", style="List Bullet")

    blocks = _import(source)

    assert [(b.marker, b.text()) for b in blocks if isinstance(b, ListItemBlock)] == [
        ("1.", "Покупатель выполнил обязательства."),
        ("а)", "подпункт без автонумерации"),
        ("1.", "Первый автопункт"),
        ("2.", "Второй автопункт"),
        ("•", "Маркированный"),
    ]


def test_table_with_horizontal_and_vertical_merges() -> None:
    source = docx.Document()
    table = source.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "Показатель"
    table.cell(0, 1).merge(table.cell(0, 2)).text = "Значение"
    table.cell(1, 1).text, table.cell(1, 2).text = "до", "после"
    table.cell(2, 0).text, table.cell(2, 1).text, table.cell(2, 2).text = "Давление", "10", "12"

    blocks = _import(source)

    assert len(blocks) == 1
    imported = blocks[0]
    assert isinstance(imported, TableBlock)
    assert [(c.col_span, c.row_span, c.runs[0].text) for c in imported.rows[0].cells] == [
        (1, 2, "Показатель"),
        (2, 1, "Значение"),
    ]
    assert [c.runs[0].text for c in imported.rows[1].cells] == ["до", "после"]
    assert imported.column_count() == 3


def test_round_trip_through_our_renderer_keeps_text(
    profiles: YamlProfileRepository,
) -> None:
    source = docx.Document()
    source.add_paragraph("Генеральному директору").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    source.add_paragraph("Прошу.")
    document = DocxImporter().convert(docx.Document(BytesIO(_bytes(source))))

    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    assert [p.text for p in rendered.paragraphs] == ["Генеральному директору", "Прошу."]


def test_strip_ads_removes_body_and_footer_lines() -> None:
    source = docx.Document()
    source.add_paragraph("Текст")
    source.add_paragraph('ООО "Ppt.ru", ИНН 1234567890')  # a company name, not an advert
    source.add_paragraph("Источник документа: правовой портал ppt.ru")
    source.sections[0].footer.paragraphs[0].text = "Больше форм и образцов документов — на ppt.ru"

    cleaned = docx.Document(BytesIO(strip_ads(_bytes(source))))

    assert [p.text for p in cleaned.paragraphs] == ["Текст", 'ООО "Ppt.ru", ИНН 1234567890']
    assert all("ppt.ru" not in p.text for p in cleaned.sections[0].footer.paragraphs)


def test_ink_f1_separates_blank_from_filled_pages() -> None:
    from PIL import Image, ImageDraw

    from foto2word.evaluation.template_cases import ink_f1

    page = Image.new("RGB", (600, 850), "white")
    ImageDraw.Draw(page).rectangle((100, 100, 500, 140), fill="black")
    shifted = Image.new("RGB", (600, 850), "white")
    ImageDraw.Draw(shifted).rectangle((103, 102, 503, 142), fill="black")
    blank = Image.new("RGB", (600, 850), "white")

    assert ink_f1(page, page) == 1.0
    assert ink_f1(page, shifted) > 0.95
    assert ink_f1(page, blank) == 0.0


def test_tables_keep_widths_and_visibility_of_lines(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    layout = source.add_table(rows=1, cols=2)  # no style: an invisible layout table
    layout.cell(0, 0).text, layout.cell(0, 1).text = "Заказчик:", "Исполнитель:"
    source.add_paragraph("")
    grid = source.add_table(rows=1, cols=2)
    grid.style = "Table Grid"
    grid.cell(0, 0).text, grid.cell(0, 1).text = "a", "b"
    for column, width in zip(grid.columns, (Mm(30), Mm(90)), strict=True):
        for cell in column.cells:
            cell.width = width
    grid._tbl.tblGrid.gridCol_lst[0].w = Mm(30)
    grid._tbl.tblGrid.gridCol_lst[1].w = Mm(90)

    tables = [b for b in _import(source) if isinstance(b, TableBlock)]

    assert [t.grid for t in tables] == [False, True]
    assert tables[1].column_widths == pytest.approx((30.0, 90.0), abs=0.05)
    assert tables[1].width_mm == pytest.approx(120.0, abs=0.05)
    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(
                DocxImporter().convert(docx.Document(BytesIO(_bytes(source)))),
                profiles.get("generic"),
                RenderMode.FAITHFUL,
            )
        )
    )
    styles = [t.style.name if t.style else None for t in rendered.tables]
    assert styles[0] != "Table Grid"
    assert styles[1] == "Table Grid"
    assert round(rendered.tables[1].rows[0].cells[0].width.mm) == 30
    grid_widths = [col.w for col in rendered.tables[1]._tbl.tblGrid.gridCol_lst]
    assert [round(_required(w).mm) for w in grid_widths] == [
        30,
        90,
    ]  # Word's autofit starts from the grid
    preferred = rendered.tables[1]._tbl.tblPr.find(qn("w:tblW"))
    assert preferred.get(qn("w:type")) == "dxa"
    assert round(int(preferred.get(qn("w:w"))) / 56.7) == 120


def test_empty_tables_are_kept_ruled_or_not() -> None:
    source = docx.Document()
    boxes = source.add_table(rows=1, cols=3)  # cells to write in by hand
    boxes.style = "Table Grid"
    source.add_paragraph("")
    source.add_table(rows=1, cols=3)  # no lines, no text: nothing to see, yet it takes room

    tables = [b for b in _import(source) if isinstance(b, TableBlock)]

    assert len(tables) == 2
    assert tables[0].grid
    assert not tables[1].grid


def test_partial_cell_borders_are_kept(profiles: YamlProfileRepository) -> None:
    from docx.oxml import OxmlElement

    source = docx.Document()
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Подпись", "Расшифровка"
    borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    borders.append(bottom)
    table.cell(0, 0)._tc.get_or_add_tcPr().append(borders)

    (imported,) = [b for b in _import(source) if isinstance(b, TableBlock)]

    first, second = imported.rows[0].cells
    assert first.borders is not None
    assert first.borders.bottom
    assert not first.borders.top
    assert second.borders is not None
    assert not second.borders.bottom
    rendered = DocxRenderer().render(
        DocxImporter().convert(docx.Document(BytesIO(_bytes(source)))),
        profiles.get("generic"),
        RenderMode.FAITHFUL,
    )
    assert b'w:bottom w:val="single"' in _document_xml(rendered)


def _document_xml(data: bytes) -> bytes:
    import zipfile

    return zipfile.ZipFile(BytesIO(data)).read("word/document.xml")


def test_text_box_content_is_read_once() -> None:
    from docx.oxml import parse_xml

    source = docx.Document()
    anchor = source.add_paragraph()
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    mc = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    content = "<w:txbxContent><w:p><w:r><w:t>В рамке</w:t></w:r></w:p></w:txbxContent>"
    box = (
        f'<w:r xmlns:w="{w}" xmlns:mc="{mc}"><mc:AlternateContent>'
        f"<mc:Choice>{content}</mc:Choice><mc:Fallback>{content}</mc:Fallback>"
        "</mc:AlternateContent></w:r>"
    )
    anchor._p.append(parse_xml(box))

    texts = [b.text() for b in _import(source)]

    assert texts == ["В рамке"]


def test_fonts_come_from_run_or_document_theme(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    source.add_paragraph("обычный")  # no font set: python-docx's template takes it from its theme
    source.add_paragraph().add_run("явный").font.name = "Arial"

    document = DocxImporter().convert(docx.Document(BytesIO(_bytes(source))))
    fonts = [b.runs[0].font for b in document.blocks if isinstance(b, ParagraphBlock)]

    assert fonts == ["Cambria", "Arial"]
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )
    assert [p.runs[0].font.name for p in rendered.paragraphs] == ["Cambria", "Arial"]
    standard = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.STANDARD))
    )
    assert {p.runs[0].font.name for p in standard.paragraphs} == {"Times New Roman"}


def test_leading_gap_autospacing_and_underlined_signature_line() -> None:
    from docx.oxml import OxmlElement

    source = docx.Document()
    source.add_paragraph("")  # pushes the content down on the page
    web = source.add_paragraph("Генеральному директору")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:beforeAutospacing"), "1")
    web._p.get_or_add_pPr().append(spacing)
    signature = source.add_paragraph("25.08.2023 ")
    signature.add_run("                    ").underline = True
    signature.add_run(" Егорова Е.Е.")

    blocks = _import(source)

    assert isinstance(blocks[0], SpacerBlock)
    assert isinstance(blocks[1], ParagraphBlock)
    assert blocks[1].layout is not None
    assert blocks[1].layout.space_before_pt == 14.0
    assert isinstance(blocks[2], ParagraphBlock)  # a signature line, not «place — date»


def test_empty_paragraphs_become_a_measured_gap_and_alignment_defaults_to_left() -> None:
    source = docx.Document()
    source.add_paragraph("Заявление")
    source.add_paragraph("")
    source.add_paragraph("")
    source.add_paragraph("Прошу предоставить отпуск.")

    blocks = _import(source)

    gap = blocks[1]
    assert isinstance(gap, SpacerBlock)
    assert gap.height_pt is not None
    assert gap.height_pt > 0
    body = blocks[2]
    assert isinstance(body, ParagraphBlock)
    assert body.align == "left"


def _spacing(paragraph: Any, **attributes: str) -> None:
    from docx.oxml import OxmlElement

    element = OxmlElement("w:spacing")
    for name, value in attributes.items():
        element.set(qn(f"w:{name}"), value)
    owner = paragraph.element if hasattr(paragraph, "element") else paragraph._p
    owner.get_or_add_pPr().append(element)


def test_own_zero_autospacing_overrides_the_style_and_gaps_collapse_with_it() -> None:
    source = docx.Document()
    web = source.styles.add_style("Web", 1)
    _spacing(web, beforeAutospacing="1", afterAutospacing="1")
    gap = source.add_paragraph("")
    _spacing(gap, after="200")  # 10 pt
    source.add_paragraph("Генеральному директору", style="Web")
    plain = source.add_paragraph("Прошу предоставить отпуск.", style="Web")
    _spacing(plain, beforeAutospacing="0", afterAutospacing="0")

    blocks = _import(source)

    spacer, first, second = blocks[0], blocks[1], blocks[2]
    assert isinstance(spacer, SpacerBlock)
    alone = docx.Document()
    lone = alone.add_paragraph("")
    _spacing(lone, after="200")
    alone.add_paragraph("Текст")
    lone_spacer = _import(alone)[0]
    assert isinstance(lone_spacer, SpacerBlock)
    assert spacer.height_pt is not None
    assert lone_spacer.height_pt is not None
    assert lone_spacer.height_pt - spacer.height_pt == pytest.approx(10.0)  # max(10, 14), not 24
    assert isinstance(first, ParagraphBlock)
    assert isinstance(second, ParagraphBlock)
    assert first.layout is not None
    assert first.layout.space_before_pt == 14.0
    assert second.layout is not None
    assert second.layout.space_before_pt == 0.0


def test_old_documents_add_up_paragraph_spacing_and_the_render_keeps_that(
    profiles: YamlProfileRepository,
) -> None:
    from docx.oxml import OxmlElement

    from foto2word.domain.document import SpacingRule
    from foto2word.infrastructure.rendering.docx_renderer import DocxRenderer

    source = docx.Document()
    compat = source.settings.element.find(qn("w:compat"))
    compat.insert(0, OxmlElement("w:doNotUseHTMLParagraphAutoSpacing"))
    source.add_paragraph("Текст")

    document = DocxImporter().convert(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    assert document.spacing_rule is SpacingRule.SUM
    assert DocxImporter().convert(docx.Document()).spacing_rule is SpacingRule.LARGER
    assert (
        rendered.settings.element.find(
            f"{qn('w:compat')}/{qn('w:doNotUseHTMLParagraphAutoSpacing')}"
        )
        is not None
    )


def _table_style(source: Any, padding_twips: int) -> Any:
    """A table style that sets Arial 8 and zero spacing, as unified forms (Т-6, М-2) do."""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement

    style = source.styles.add_style("FormTable", WD_STYLE_TYPE.TABLE)
    style.font.name, style.font.size = "Arial", Pt(8)
    style.paragraph_format.space_after = Pt(0)
    tbl_pr = OxmlElement("w:tblPr")
    margins = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:w"), str(padding_twips))
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    tbl_pr.append(margins)
    style.element.append(tbl_pr)
    return style


def test_cells_inherit_the_table_style_and_keep_their_blank_lines(
    profiles: YamlProfileRepository,
) -> None:
    source = docx.Document()
    table = source.add_table(rows=1, cols=2)
    table.style = _table_style(source, padding_twips=28)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    table.cell(0, 0).text = "Форма по ОКУД"
    blank = table.cell(0, 1)
    blank.add_paragraph("")
    blank.add_paragraph("")  # three empty lines make the row three lines tall

    block = next(b for b in _import(source) if isinstance(b, TableBlock))

    label, filler = block.rows[0].cells
    assert label.runs[0].font == "Arial"
    assert label.runs[0].size_pt == 8.0
    assert label.layout is not None
    assert label.layout.space_after_pt == 0.0
    assert [r.text for r in filler.runs] == ["\n", "\n"]
    assert filler.size_pt == 8.0
    assert block.align is Alignment.RIGHT
    assert block.padding is not None
    assert block.padding.left_mm == pytest.approx(0.49, abs=0.01)
    assert block.padding.top_mm == 0.0
    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )
    tbl_pr = rendered.tables[0]._tbl.tblPr
    margin = tbl_pr.find(f"{qn('w:tblCellMar')}/{qn('w:left')}")
    assert margin is not None
    assert margin.get(qn("w:w")) == "28"
    assert rendered.tables[0].alignment == WD_TABLE_ALIGNMENT.RIGHT


def test_what_normal_says_itself_wins_over_the_table_style() -> None:
    """Measured in Word: the table style's spacing wins over the document defaults, but not over
    a value 'Normal' sets itself (an application form: 8 pt after each line of its cells)."""
    source = docx.Document()
    source.styles["Normal"].paragraph_format.space_after = Pt(8)
    table = source.add_table(rows=1, cols=1)
    table.style = _table_style(source, padding_twips=28)
    table.cell(0, 0).text = "от"

    block = next(b for b in _import(source) if isinstance(b, TableBlock))

    (cell,) = block.rows[0].cells
    assert cell.layout is not None
    assert cell.layout.space_after_pt == 8.0
    assert cell.runs[0].font == "Arial"  # Normal names no typeface: the table style's


def test_grid_positions_without_a_cell_take_no_space(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "a", "b"
    table.cell(1, 0).text = "c"
    row = table.rows[1]._tr
    row.remove(row.tc_lst[1])  # Word's gridAfter: the row simply ends early

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    short = rendered.tables[0].rows[1]._tr
    assert len(short.tc_lst) == 1  # no filler cell: Word refuses rows of over 63 cells
    assert _required(short.trPr).find(qn("w:gridAfter")).get(qn("w:val")) == "1"


def test_an_exact_row_height_stays_exact(profiles: YamlProfileRepository) -> None:
    from docx.enum.table import WD_ROW_HEIGHT_RULE

    source = docx.Document()
    table = source.add_table(rows=2, cols=1)
    table.cell(0, 0).text, table.cell(1, 0).text = "a", "b"
    table.rows[0].height, table.rows[0].height_rule = Mm(7), WD_ROW_HEIGHT_RULE.EXACTLY
    table.rows[1].height = Mm(5)  # no rule: Word's default is "at least"

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    rows = rendered.tables[0].rows
    assert rows[0].height_rule == WD_ROW_HEIGHT_RULE.EXACTLY
    assert rows[1].height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST


def test_word_may_widen_columns_only_where_the_original_lets_it(
    profiles: YamlProfileRepository,
) -> None:
    source = docx.Document()
    for autofit in (True, False):
        table = source.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "a", "b"
        table.autofit = autofit

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    assert [t.autofit for t in rendered.tables] == [True, False]


def test_the_rest_of_a_vertically_merged_cell_keeps_its_own_empty_line(
    profiles: YamlProfileRepository,
) -> None:
    source = docx.Document()
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(1, 0).text = "a", "b"
    merged = table.cell(0, 1).merge(table.cell(1, 1))
    merged.paragraphs[0].add_run("code").font.size = Pt(8)
    rest_mark = table.rows[1]._tr.tc_lst[1].p_lst[0].get_or_add_pPr()
    rest_mark.append(
        docx.oxml.parse_xml(f'<w:rPr {docx.oxml.ns.nsdecls("w")}><w:sz w:val="12"/></w:rPr>')
    )

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    rest = rendered.tables[0].rows[1]._tr.tc_lst[1]
    marks = rest.findall(f"{qn('w:p')}/{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:sz')}")
    assert [m.get(qn("w:val")) for m in marks] == ["12"]


def test_an_empty_cell_that_hides_its_mark_stays_hidden(profiles: YamlProfileRepository) -> None:
    from docx.oxml import OxmlElement

    source = docx.Document()
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "a"
    table.cell(0, 1)._tc.get_or_add_tcPr().append(OxmlElement("w:hideMark"))

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    cells = rendered.tables[0].rows[0]._tr.tc_lst
    assert [_required(tc.tcPr).find(qn("w:hideMark")) is not None for tc in cells] == [False, True]


def test_exact_and_at_least_line_heights_are_kept_as_set(profiles: YamlProfileRepository) -> None:
    from docx.enum.text import WD_LINE_SPACING

    source = docx.Document()
    at_least = source.add_paragraph("a")
    at_least.paragraph_format.line_spacing = Pt(6)  # below the natural line: Word keeps it
    at_least.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    exact = source.add_paragraph("b")
    exact.paragraph_format.line_spacing = Pt(9)
    exact.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    lines = [
        (p.paragraph_format.line_spacing_rule, p.paragraph_format.line_spacing)
        for p in rendered.paragraphs
        if p.text in ("a", "b")
    ]
    assert lines == [(WD_LINE_SPACING.AT_LEAST, Pt(6)), (WD_LINE_SPACING.EXACTLY, Pt(9))]


def test_letter_spacing_is_kept(profiles: YamlProfileRepository) -> None:
    from docx.oxml import OxmlElement

    source = docx.Document()
    run = source.add_paragraph().add_run("ПРИКАЗ")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), "60")  # 3 pt between letters
    run._r.get_or_add_rPr().append(spacing)

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    target = next(r for p in rendered.paragraphs for r in p.runs if r.text == "ПРИКАЗ")
    found = _required(target._r.rPr).find(qn("w:spacing"))
    assert found is not None
    assert found.get(qn("w:val")) == "60"


def test_a_table_inside_a_cell_is_kept_in_its_place(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    frame = source.add_table(rows=1, cols=1)  # a form drawn inside one framing cell
    frame.style = "Table Grid"
    cell = frame.cell(0, 0)
    cell.paragraphs[0].text = "Заявление"
    inner = cell.add_table(rows=1, cols=2)
    inner.cell(0, 0).text, inner.cell(0, 1).text = "Фамилия", "Иванов"
    cell.paragraphs[-1].text = "Подпись"

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    assert "Иванов" in " ".join(document.iter_text())
    target = rendered.tables[0].cell(0, 0)
    order = [
        "table"
        if child.tag == qn("w:tbl")
        else "".join(t.text or "" for t in child.iter(qn("w:t")))
        for child in target._tc.iterchildren(qn("w:p"), qn("w:tbl"))
    ]
    assert order == ["Заявление", "table", "Подпись"]
    assert [c.text for c in target.tables[0].rows[0].cells] == ["Фамилия", "Иванов"]


def test_a_floating_table_keeps_its_position(profiles: YamlProfileRepository) -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    source = docx.Document()
    source.add_paragraph("Текст рядом с таблицей")
    table = source.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Код"
    table._tbl.tblPr.insert(
        0,
        parse_xml(
            f'<w:tblpPr {nsdecls("w")} w:leftFromText="180" w:rightFromText="180" '
            'w:vertAnchor="text" w:horzAnchor="margin" w:tblpXSpec="right" w:tblpY="-283"/>'
        ),
    )

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    position = rendered.tables[0]._tbl.tblPr.find(qn("w:tblpPr"))
    assert position is not None
    assert {k.split("}")[1]: v for k, v in position.attrib.items()} == {
        "leftFromText": "180",
        "rightFromText": "180",
        "vertAnchor": "text",
        "horzAnchor": "margin",
        "tblpXSpec": "right",
        "tblpY": "-283",
    }


def _flow_with_breaks(rendered: Any) -> list[str]:
    """Texts of the body paragraphs, with 'BREAK' where a new page starts: a break before a
    paragraph, a page break in it, or a section break after it when the next section starts
    on a new page."""
    body = rendered.element.body
    sections = list(body.iter(qn("w:sectPr")))
    out = []
    for p in body.iter(qn("w:p")):
        if p.find(f"{qn('w:pPr')}/{qn('w:pageBreakBefore')}") is not None or any(
            br.get(qn("w:type")) == "page" for br in p.iter(qn("w:br"))
        ):
            out.append("BREAK")
        text = "".join(t.text or "" for t in p.iter(qn("w:t")))
        if text:
            out.append(text)
        end = p.find(f"{qn('w:pPr')}/{qn('w:sectPr')}")
        if end is not None:
            kind = sections[sections.index(end) + 1].find(qn("w:type"))
            if kind is None or kind.get(qn("w:val")) != "continuous":
                out.append("BREAK")
    return out


def test_page_and_section_breaks_start_a_new_page(profiles: YamlProfileRepository) -> None:
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_BREAK

    source = docx.Document()
    source.add_paragraph("Первый")
    source.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    source.add_paragraph("Второй")
    source.add_section(WD_SECTION.NEW_PAGE)
    source.add_paragraph("Третий")
    source.add_section(WD_SECTION.CONTINUOUS)
    source.add_paragraph("Четвёртый")
    fifth = source.add_paragraph("Пятый")
    fifth.paragraph_format.page_break_before = True

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    assert _flow_with_breaks(rendered) == [
        "Первый",
        "BREAK",
        "Второй",
        "BREAK",
        "Третий",
        "Четвёртый",
        "BREAK",
        "Пятый",
    ]


def _anchored(inline: Any) -> Any:
    """The same picture floating: 50 mm from the page's left edge, 120 mm from its top,
    behind the text, as a stamp over a signature line."""
    from docx.oxml import parse_xml

    wp = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    anchor = parse_xml(
        f'<wp:anchor xmlns:wp="{wp}" distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        'relativeHeight="2" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>1800000</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>4320000</wp:posOffset></wp:positionV>'
        "</wp:anchor>"
    )
    extent, *rest = list(inline)
    anchor.append(extent)
    anchor.append(parse_xml(f'<wp:wrapNone xmlns:wp="{wp}"/>'))
    for child in rest:
        if child.tag.endswith("effectExtent"):
            continue
        anchor.append(child)
    inline.getparent().replace(inline, anchor)
    return anchor


def test_a_floating_picture_stays_where_it_was(profiles: YamlProfileRepository) -> None:
    from PIL import Image

    png = BytesIO()
    Image.new("RGB", (40, 40), "blue").save(png, format="PNG")
    source = docx.Document()
    paragraph = source.add_paragraph("Директор ________ Иванов")
    paragraph.add_run().add_picture(BytesIO(png.getvalue()), width=Mm(30))
    _anchored(next(paragraph._p.iter(qn("wp:inline"))))

    importer = DocxImporter()
    document = importer.convert(docx.Document(BytesIO(_bytes(source))))
    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(
                document,
                profiles.get("generic"),
                RenderMode.FAITHFUL,
                RenderExtras(images=importer.images),
            )
        )
    )

    body = rendered.element.body
    assert "Директор ________ Иванов" in [p.text for p in rendered.paragraphs]
    assert next(body.iter(qn("wp:inline")), None) is None
    floating = next(body.iter(qn("wp:anchor")))
    assert floating.get("behindDoc") == "1"
    horizontal = floating.find(qn("wp:positionH"))
    vertical = floating.find(qn("wp:positionV"))
    assert horizontal.get("relativeFrom") == vertical.get("relativeFrom") == "page"
    assert int(horizontal.find(qn("wp:posOffset")).text) == pytest.approx(1800000, abs=4000)
    assert int(vertical.find(qn("wp:posOffset")).text) == pytest.approx(4320000, abs=4000)
    assert floating.find(qn("wp:wrapNone")) is not None


def test_text_after_a_clearing_break_goes_below_the_floating_picture(
    profiles: YamlProfileRepository,
) -> None:
    """An emblem floating above a heading: the line break with clear="all" after it sends the
    heading below the emblem instead of beside it."""
    from docx.oxml import parse_xml
    from PIL import Image

    png = BytesIO()
    Image.new("RGB", (40, 40), "blue").save(png, format="PNG")
    source = docx.Document()
    paragraph = source.add_paragraph()
    paragraph.add_run().add_picture(BytesIO(png.getvalue()), width=Mm(25))
    _anchored(next(paragraph._p.iter(qn("wp:inline"))))
    paragraph._p.append(parse_xml(f'<w:r xmlns:w="{_W}"><w:br w:clear="all"/></w:r>'))
    source.add_paragraph("МЕСТНАЯ АДМИНИСТРАЦИЯ")

    importer = DocxImporter()
    document = importer.convert(docx.Document(BytesIO(_bytes(source))))
    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(
                document,
                profiles.get("generic"),
                RenderMode.FAITHFUL,
                RenderExtras(images=importer.images),
            )
        )
    )

    anchor = next(rendered.element.body.iter(qn("wp:anchor")))
    holder = next(a for a in anchor.iterancestors() if a.tag == qn("w:p"))
    clearing = [b.get(qn("w:clear")) for b in holder.iter(qn("w:br"))]
    assert clearing == ["all"]
    assert len(list(rendered.element.body.iter(qn("w:br")))) == 1  # not one more empty line


_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
_DML = "http://schemas.openxmlformats.org/drawingml/2006/main"
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _text_box_run(text: str) -> Any:
    """A run holding a text box 40 x 15 mm, 120 mm from the page's left and 200 mm from its
    top, framed: as 'М.П.' marked beside a signature."""
    from docx.oxml import parse_xml

    return parse_xml(
        f'<w:r xmlns:w="{_W}"><w:drawing>'
        f'<wp:anchor xmlns:wp="{_WP}" distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        'relativeHeight="3" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>4320000</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>7200000</wp:posOffset></wp:positionV>'
        '<wp:extent cx="1440000" cy="540000"/><wp:wrapNone/><wp:docPr id="7" name="Box"/>'
        f'<a:graphic xmlns:a="{_DML}"><a:graphicData uri="{_WPS}">'
        f'<wps:wsp xmlns:wps="{_WPS}"><wps:cNvSpPr txBox="1"/>'
        '<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1440000" cy="540000"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        '<a:ln w="6350"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln></wps:spPr>'
        f"<wps:txbx><w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent>"
        "</wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic>"
        "</wp:anchor></w:drawing></w:r>"
    )


def test_a_floating_text_box_keeps_its_place_and_text(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    paragraph = source.add_paragraph("Директор ________ Иванов")
    paragraph._p.append(_text_box_run("М.П."))

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    assert "М.П." in " ".join(document.iter_text())
    body = rendered.element.body
    flow = [p for p in body.iterchildren(qn("w:p"))]
    assert all(
        "М.П." not in "".join(t.text or "" for t in p.iter(qn("w:t")))
        or next(p.iter(qn("w:txbxContent")), None) is not None
        for p in flow
    )
    box = next(body.iter(f"{{{_W}}}txbxContent"))
    assert "".join(t.text or "" for t in box.iter(qn("w:t"))) == "М.П."
    floating = next(body.iter(qn("wp:anchor")))
    horizontal = floating.find(qn("wp:positionH"))
    assert horizontal.get("relativeFrom") == "page"
    assert int(horizontal.find(qn("wp:posOffset")).text) == pytest.approx(4320000, abs=4000)
    extent = floating.find(qn("wp:extent"))
    assert int(extent.get("cx")) == pytest.approx(1440000, abs=4000)
    assert next(floating.iter(f"{{{_DML}}}ln")).find(f"{{{_DML}}}noFill") is None  # framed


def _shape_run(
    geometry: str,
    cx: int,
    cy: int,
    flip: str = "",
    line: str = '<a:solidFill><a:srgbClr val="000000"/></a:solidFill>',
) -> Any:
    """A run holding a drawn shape (no text): 30 mm from the page's left, 100 mm from its top."""
    from docx.oxml import parse_xml

    return parse_xml(
        f'<w:r xmlns:w="{_W}"><w:drawing>'
        f'<wp:anchor xmlns:wp="{_WP}" distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        'relativeHeight="3" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>1080000</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>3600000</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:wrapNone/><wp:docPr id="9" name="Shape"/>'
        f'<a:graphic xmlns:a="{_DML}"><a:graphicData uri="{_WPS}">'
        f'<wps:wsp xmlns:wps="{_WPS}"><wps:cNvSpPr/>'
        f'<wps:spPr><a:xfrm{flip}><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="{geometry}"><a:avLst/></a:prstGeom><a:noFill/>'
        f'<a:ln w="12700">{line}</a:ln></wps:spPr>'
        "<wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>"
    )


def test_drawn_lines_and_frames_keep_their_place(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    paragraph = source.add_paragraph("Подпись")
    paragraph._p.append(_shape_run("line", 2160000, 0))  # a 60 mm rule to sign on
    paragraph._p.append(_shape_run("rect", 720000, 540000))  # a 20 x 15 mm box for a stamp
    paragraph._p.append(_shape_run("straightConnector1", 360000, 360000, ' flipV="1"'))

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    shapes = [
        (
            next(a.iter(f"{{{_DML}}}prstGeom")).get("prst"),
            int(a.find(qn("wp:extent")).get("cx")),
            int(a.find(qn("wp:extent")).get("cy")),
            next(a.iter(f"{{{_DML}}}xfrm")).get("flipV"),
            int(a.find(qn("wp:positionH")).find(qn("wp:posOffset")).text),
        )
        for a in rendered.element.body.iter(qn("wp:anchor"))
    ]
    assert [(g, round(cx / 36000), round(cy / 36000), f) for g, cx, cy, f, _ in shapes] == [
        ("line", 60, 0, None),
        ("rect", 20, 15, None),
        ("line", 10, 10, "1"),
    ]
    assert all(abs(x - 1080000) < 4000 for *_, x in shapes)
    assert all(
        next(a.iter(f"{{{_DML}}}ln")).get("w") == "12700"
        for a in rendered.element.body.iter(qn("wp:anchor"))
    )


def _break_in_last_paragraph(document: Any) -> None:
    """As Word saves it: the section break sits in the section's last paragraph, not in an
    empty one of its own (python-docx adds one)."""
    holder = [p for p in document.element.body.iterchildren(qn("w:p"))][-1]
    previous = holder.getprevious()
    previous.get_or_add_pPr().append(holder.pPr.find(qn("w:sectPr")))
    holder.getparent().remove(holder)


def _columns(section: Any) -> str | None:
    cols = section._sectPr.find(qn("w:cols"))
    return cols.get(qn("w:num")) if cols is not None else None


def test_sections_keep_their_columns_orientation_and_breaks(
    profiles: YamlProfileRepository,
) -> None:
    from docx.enum.section import WD_ORIENT, WD_SECTION

    source = docx.Document()
    source.add_paragraph("Заявление")
    two = source.add_section(WD_SECTION.CONTINUOUS)
    _break_in_last_paragraph(source)
    cols = two._sectPr.find(qn("w:cols"))
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "708")
    source.add_paragraph("Левая колонка")
    source.add_paragraph("Правая колонка")
    wide = source.add_section(WD_SECTION.NEW_PAGE)
    _break_in_last_paragraph(source)
    wide.orientation = WD_ORIENT.LANDSCAPE
    wide.page_width, wide.page_height = Mm(297), Mm(210)
    wide._sectPr.remove(wide._sectPr.find(qn("w:cols")))
    source.add_paragraph("Таблица на альбомном листе")

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    sections = rendered.sections
    assert len(sections) == 3
    assert [s.start_type for s in sections[1:]] == [WD_SECTION.CONTINUOUS, WD_SECTION.NEW_PAGE]
    assert [_columns(s) for s in sections] == [None, "2", None]
    assert sections[2].orientation == WD_ORIENT.LANDSCAPE
    assert round(_required(sections[2].page_width).mm) == 297
    texts = [p.text for p in rendered.paragraphs]
    assert "" not in texts[: texts.index("Таблица на альбомном листе")]  # no empty lines added


def test_page_numbers_and_style_italics_survive_the_round_trip(
    profiles: YamlProfileRepository,
) -> None:
    from docx.oxml import OxmlElement

    source = docx.Document()
    footer = source.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " PAGE ")
    footer._p.append(field)
    quote = source.styles.add_style("Quote italic", 1)
    quote.font.italic = True
    source.add_paragraph("Служебная записка", style="Quote italic")

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    assert document.page_numbers is not None
    assert document.page_numbers.position == "bottom"
    assert document.page_numbers.align is Alignment.RIGHT
    body = document.blocks[0]
    assert isinstance(body, ParagraphBlock | HeadingBlock)
    assert body.runs[0].italic
    codes = [t.text for t in rendered.sections[0].footer._element.iter(qn("w:instrText"))]
    assert codes == [" PAGE "]


def test_numbering_level_indents_and_cell_paragraphs_are_kept(
    profiles: YamlProfileRepository,
) -> None:
    source = docx.Document()
    source.add_paragraph("В связи с увеличением стоимости", style="List Number")
    table = source.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Заказчик"
    table.cell(0, 0).add_paragraph("ООО «Ромашка»")

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    item = document.blocks[0]
    assert isinstance(item, ListItemBlock)
    assert item.layout is not None
    assert (item.layout.first_line_indent_mm or 0) < 0  # hanging: the number has its own column
    assert _word_marker(rendered.paragraphs[0])[::2] == ("1.", "tab")
    assert rendered.paragraphs[0].text == "В связи с увеличением стоимости"
    cell = rendered.tables[0].cell(0, 0)
    assert [p.text for p in cell.paragraphs] == ["Заказчик", "ООО «Ромашка»"]


def test_start_and_end_alignment_of_newer_word_files() -> None:
    from docx.oxml import OxmlElement

    source = docx.Document()
    for text, value in (("Директору", "end"), ("Текст", "start")):
        paragraph = source.add_paragraph(text)
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), value)
        paragraph._p.get_or_add_pPr().append(jc)

    blocks = _import(source)

    assert isinstance(blocks[0], ParagraphBlock)
    assert blocks[0].align is Alignment.RIGHT
    assert isinstance(blocks[1], ParagraphBlock)
    assert blocks[1].align is Alignment.LEFT


def test_pictures_become_image_blocks_and_render_back(profiles: YamlProfileRepository) -> None:
    from PIL import Image

    from foto2word.domain.document import ImageBlock
    from foto2word.domain.extras import RenderExtras

    logo = BytesIO()
    Image.new("RGB", (40, 20), "red").save(logo, format="PNG")
    source = docx.Document()
    source.add_paragraph().add_run().add_picture(BytesIO(logo.getvalue()), width=Mm(40))
    source.add_paragraph("ООО «Ромашка»")
    importer = DocxImporter()

    document = importer.convert(docx.Document(BytesIO(_bytes(source))))

    image = document.blocks[0]
    assert isinstance(image, ImageBlock)
    assert image.width_mm == pytest.approx(40.0, abs=0.1)
    assert importer.images[image.id] == logo.getvalue()
    generic = profiles.get("generic")
    with_picture = docx.Document(
        BytesIO(
            DocxRenderer().render(
                document, generic, RenderMode.FAITHFUL, RenderExtras(images=importer.images)
            )
        )
    )
    without = docx.Document(BytesIO(DocxRenderer().render(document, generic, RenderMode.FAITHFUL)))
    assert len(with_picture.inline_shapes) == 1
    assert without.paragraphs[0].text == "[рисунок]"


def test_tab_stops_with_leaders_and_check_box_symbols(profiles: YamlProfileRepository) -> None:
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.oxml import OxmlElement

    source = docx.Document()
    line = source.add_paragraph("Расположение:\t")
    line.paragraph_format.tab_stops.add_tab_stop(
        Mm(160), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES
    )
    box = source.add_paragraph()
    symbol = OxmlElement("w:sym")
    symbol.set(qn("w:font"), "Wingdings")
    symbol.set(qn("w:char"), "F0A8")
    box.add_run()._r.append(symbol)
    box.add_run(" Прошу зарегистрировать меня")

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    first, second = document.blocks[0], document.blocks[1]
    assert isinstance(first, ParagraphBlock)
    assert first.layout is not None
    (stop,) = first.layout.tabs
    assert (stop.position_mm, stop.align, stop.leader) == (160.0, "right", "underscore")
    assert isinstance(second, ParagraphBlock)
    assert second.text().startswith("☐")
    (rendered_stop,) = rendered.paragraphs[0].paragraph_format.tab_stops
    assert rendered_stop.leader == WD_TAB_LEADER.LINES


def test_letterhead_in_the_header_and_contacts_in_the_footer(
    profiles: YamlProfileRepository,
) -> None:
    from docx.oxml import OxmlElement

    source = docx.Document()
    section = source.sections[0]
    section.header.paragraphs[0].text = "ООО «Ромашка» · ИНН 7701234567"
    footer = section.footer.paragraphs[0]
    footer.text = "Тел. +7 (495) 123-45-67"
    page = section.footer.add_paragraph()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " PAGE ")
    page._p.append(field)
    source.add_paragraph("Письмо")

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    assert [b.text() for b in document.header] == ["ООО «Ромашка» · ИНН 7701234567"]
    assert [b.text() for b in document.footer] == ["Тел. +7 (495) 123-45-67"]
    assert document.header[0].id.startswith("h")
    assert document.page_numbers is not None
    out = rendered.sections[0]
    assert out.header.paragraphs[0].text == "ООО «Ромашка» · ИНН 7701234567"
    assert "Тел. +7 (495) 123-45-67" in [p.text for p in out.footer.paragraphs]


def test_rule_under_a_letterhead_is_kept(profiles: YamlProfileRepository) -> None:
    from docx.oxml import OxmlElement

    source = docx.Document()
    head = source.add_paragraph("ООО «Ромашка»")
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    p_bdr.append(bottom)
    head._p.get_or_add_pPr().insert(0, p_bdr)
    source.add_paragraph("Исх. № 15")

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    first = document.blocks[0]
    assert isinstance(first, ParagraphBlock)
    assert first.layout is not None
    assert first.layout.borders is not None
    assert (first.layout.borders.bottom, first.layout.borders.top) == (True, False)
    lines = rendered.paragraphs[0]._p.find(f"{qn('w:pPr')}/{qn('w:pBdr')}")
    assert lines is not None
    assert [e.tag.split("}")[1] for e in lines] == ["bottom"]


def test_picture_in_a_format_word_files_cannot_take_becomes_a_note(
    profiles: YamlProfileRepository,
) -> None:
    from foto2word.domain.document import DocumentTypeGuess, ImageBlock
    from foto2word.domain.extras import RenderExtras

    document = Document(
        document_type=DocumentTypeGuess(id="other", title="Схема", confidence=1.0),
        blocks=(ImageBlock(id="b1", description="схема", width_mm=50),),
    )
    extras = RenderExtras(images={"b1": b"\x01\x00\x00\x00 EMF drawing"})

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL, extras)
        )
    )

    assert rendered.paragraphs[0].text == "[рисунок: схема]"


def test_fill_row_renders_blanks_on_lines_with_captions_under_them(
    profiles: YamlProfileRepository,
) -> None:
    from foto2word.domain.document import DocumentTypeGuess, FillRowBlock, FillSlot, Run

    row = FillRowBlock(
        id="b1",
        slots=(
            FillSlot(label="Руководитель", caption="должность", width=1.0),
            FillSlot(caption="личная подпись", width=1.0),
            FillSlot(value=(Run(text="П. П. Петров"),), caption="расшифровка подписи", width=1.4),
        ),
    )
    document = Document(
        document_type=DocumentTypeGuess(id="order", title="Приказ", confidence=1.0),
        blocks=(row,),
    )

    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    table = rendered.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Руководитель", "", "", "П. П. Петров"]
    assert [c.text for c in table.rows[1].cells][1:] == [
        "должность",
        "личная подпись",
        "расшифровка подписи",
    ]
    line = table.rows[0].cells[3]._tc.find(f"{qn('w:tcPr')}/{qn('w:tcBorders')}/{qn('w:bottom')}")
    assert line is not None
    assert line.get(qn("w:val")) == "single"
    assert row.text() == "Руководитель должность личная подпись П. П. Петров расшифровка подписи"
    assert row.replace_text("Петров", "Иванов") is not None


def test_drawings_in_the_body_and_the_footer_get_their_own_ids(
    profiles: YamlProfileRepository,
) -> None:
    source = docx.Document()
    source.add_paragraph("Текст")._p.append(_text_box_run("М.П."))
    source.sections[0].footer.paragraphs[0]._p.append(_text_box_run("Стр."))

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    ids = [
        d.get("id")
        for part in (rendered.element, rendered.sections[0].footer._element)
        for d in part.iter(qn("wp:docPr"))
    ]
    assert len(ids) == 2
    assert len(set(ids)) == 2


def test_tab_leaders_make_a_fill_line_and_keep_their_last_blank(
    profiles: YamlProfileRepository,
) -> None:
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER

    source = docx.Document()
    line = source.add_paragraph("дата рождения\tдата смерти\t")
    for position in (Mm(70), Mm(150)):
        line.paragraph_format.tab_stops.add_tab_stop(
            position, WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.LINES
        )

    blocks = _import(source)
    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    assert isinstance(blocks[0], ParagraphBlock)
    assert blocks[0].text() == "дата рождения\tдата смерти\t"
    stops = rendered.paragraphs[0].paragraph_format.tab_stops
    assert [s.leader for s in stops] == [WD_TAB_LEADER.LINES, WD_TAB_LEADER.LINES]


def test_editable_output_lets_rows_lines_and_boxes_grow(profiles: YamlProfileRepository) -> None:
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.enum.text import WD_LINE_SPACING

    source = docx.Document()
    table = source.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "a"
    table.rows[0].height, table.rows[0].height_rule = Mm(7), WD_ROW_HEIGHT_RULE.EXACTLY
    tight = source.add_paragraph("b")
    tight.paragraph_format.line_spacing = Pt(9)
    tight.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    source.add_paragraph("c")._p.append(_text_box_run("М.П."))

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(
                _document(source),
                profiles.get("generic"),
                RenderMode.FAITHFUL,
                RenderExtras(editable=True),
            )
        )
    )

    assert rendered.tables[0].rows[0].height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST
    line = next(p for p in rendered.paragraphs if p.text == "b").paragraph_format
    assert line.line_spacing_rule == WD_LINE_SPACING.AT_LEAST
    body = next(rendered.element.body.iter(f"{{{_WPS}}}bodyPr"))
    assert body.find(f"{{{_DML}}}spAutoFit") is not None  # the box grows with its text


def test_the_header_distance_is_kept(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    section = source.sections[0]
    section.top_margin, section.header_distance = Mm(6), Mm(15)
    section.header.paragraphs[0].text = ""  # the header exists, only an empty line in it
    source.add_paragraph("ДОГОВОР")

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    assert round(_required(rendered.sections[0].header_distance).mm) == 15


@pytest.mark.parametrize("aligned", [True, False])
def test_an_empty_header_takes_room_only_with_its_own_formatting(
    profiles: YamlProfileRepository, aligned: bool
) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    source = docx.Document()
    section = source.sections[0]
    section.top_margin, section.header_distance = Mm(6), Mm(15)
    empty = section.header.paragraphs[0]
    if aligned:  # Word counts an empty header line that carries formatting of its own
        empty.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    source.add_paragraph("ДОГОВОР")

    document = _document(source)

    assert bool(document.header) is aligned


def test_a_bullet_keeps_its_own_typeface(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    item = source.add_paragraph(style="List Bullet")
    item.add_run("Модель: ____").font.size = Pt(9)

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    block = document.blocks[0]
    assert isinstance(block, ListItemBlock)
    assert block.marker == "•"
    assert block.marker_font == "Symbol"  # its line height sets the item's, not the text's
    assert _word_marker(rendered.paragraphs[0])[:2] == ("\uf0b7", "Symbol")


def test_a_bullet_keeps_its_own_sign(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    num_id = _numbered(source, (("bullet", "-"), ("bullet", "")))
    _number(source.add_paragraph("наличными"), num_id, 0)
    _number(source.add_paragraph("картой"), num_id, 1)

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    items = [b for b in document.blocks if isinstance(b, ListItemBlock)]
    assert [item.marker for item in items] == ["-", "▪"]  # a Wingdings square reads as one
    assert _word_marker(rendered.paragraphs[0])[0] == "-"
    assert _word_marker(rendered.paragraphs[1])[0] == ""  # Word's own sign


@pytest.mark.parametrize(("suffix", "gap"), [(None, "\t"), ("space", " "), ("nothing", "")])
def test_a_number_is_followed_by_what_its_level_says(
    profiles: YamlProfileRepository, suffix: str | None, gap: str
) -> None:
    source = docx.Document()
    num_id = _numbered(source, (("decimal", "%1."),))
    if suffix is not None:
        level = source.part.numbering_part.element.find(f".//{qn('w:lvl')}")
        level.append(docx.oxml.parse_xml(f'<w:suff {_W_DECL} w:val="{suffix}"/>'))
    item = source.add_paragraph("ФИО")
    _number(item, num_id, 0)
    item.paragraph_format.left_indent = item.paragraph_format.first_line_indent = 0  # no hanging

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    # a tab even without a hanging indent
    assert _word_marker(rendered.paragraphs[0])[::2] == ("1.", suffix or "tab")
    assert rendered.paragraphs[0].text == "ФИО"


def test_a_picture_bullet_is_not_drawn_as_its_sign() -> None:
    source = docx.Document()
    num_id = _numbered(source, (("bullet", ""),))
    level = source.part.numbering_part.element.find(f".//{qn('w:lvl')}")
    level.append(docx.oxml.parse_xml(f'<w:lvlPicBulletId {_W_DECL} w:val="0"/>'))
    _number(source.add_paragraph("1 смена"), num_id, 0)

    (item,) = [b for b in _import(source) if isinstance(b, ListItemBlock)]

    assert item.marker == "•"
    assert item.marker_glyph is None  # Word shows a picture, not the Symbol dot


def test_text_set_in_capitals_by_its_formatting_reads_in_capitals() -> None:
    source = docx.Document()
    source.add_paragraph().add_run("Основные направления работы:").font.all_caps = True

    (block,) = _import(source)

    assert "ОСНОВНЫЕ НАПРАВЛЕНИЯ РАБОТЫ:" in "".join(r.text for r in _textual(block).runs)


def test_a_list_line_moving_only_its_left_edge_keeps_the_level_hanging_marker() -> None:
    source = docx.Document()
    num_id = _numbered(source, (("bullet", "-"),))
    level = source.part.numbering_part.element.find(f".//{qn('w:lvl')}")
    level.append(
        docx.oxml.parse_xml(f'<w:pPr {_W_DECL}><w:ind w:left="720" w:hanging="360"/></w:pPr>')
    )
    item = source.add_paragraph("история развития естественных наук;")
    _number(item, num_id, 0)
    item._p.get_or_add_pPr().append(docx.oxml.parse_xml(f'<w:ind {_W_DECL} w:left="1070"/>'))

    (block,) = [b for b in _import(source) if isinstance(b, ListItemBlock)]

    assert block.layout is not None
    assert (
        round(_required(_layout(block).left_indent_mm), 1) == 18.9
    )  # the paragraph's own left edge
    assert block.layout.first_line_indent_mm == -6.35  # the level's hanging


def test_text_inside_a_smart_tag_is_kept() -> None:
    source = docx.Document()
    paragraph = source.add_paragraph("место нахождения: ")
    paragraph._p.append(
        docx.oxml.parse_xml(
            f'<w:smartTag {_W_DECL} w:uri="urn:schemas-microsoft-com:office:smarttags" '
            'w:element="metricconverter"><w:smartTagPr><w:attr w:name="ProductID" '
            'w:val="127015, г"/></w:smartTagPr><w:r><w:t>127015, г</w:t></w:r></w:smartTag>'
        )
    )
    paragraph.add_run(".Москва")

    (block,) = _import(source)

    assert "".join(r.text for r in _textual(block).runs) == "место нахождения: 127015, г.Москва"


def test_tabs_at_the_end_of_a_line_are_kept_and_spaces_are_not() -> None:
    source = docx.Document()
    source.add_paragraph("(Ф.И.О.)\t\t\t")
    source.add_paragraph("Директору   ")

    first, second = _import(source)

    tail = "(Ф.И.О.)\t\t\t"
    assert "".join(r.text for r in _textual(first).runs).endswith(tail)  # Word gives them room
    assert "".join(r.text for r in _textual(second).runs) == "Директору"


def test_sizes_written_as_fractions_are_read() -> None:
    source = docx.Document()
    item = source.add_paragraph("Пункт")
    item._p.get_or_add_pPr().append(
        docx.oxml.parse_xml(f'<w:ind {_W_DECL} w:left="0.0" w:firstLine="359.99999999999994"/>')
    )
    table = source.add_table(rows=1, cols=1)
    table._tbl.tblGrid.find(qn("w:gridCol")).set(qn("w:w"), "1608.75")

    blocks = _import(source)  # some writers put fractions where Word puts whole twips

    assert _layout(blocks[0]) is not None
    assert round(_layout(blocks[0]).first_line_indent_mm or 0) == 6


def test_a_document_without_page_settings_gets_words_default_page() -> None:
    source = docx.Document()
    source.add_paragraph("Заявление")
    body = source.element.body
    body.remove(body.find(qn("w:sectPr")))

    document = _document(source)

    (block,) = document.blocks
    assert "".join(r.text for r in _textual(block).runs) == "Заявление"
    assert document.page is not None
    assert round(document.page.width_mm) == 210  # A4, as Word lays it out (measured)
    assert round(document.page.left_mm) == 30
    assert round(document.page.right_mm) == 15


def test_a_section_without_page_size_is_letter_and_without_margins_gets_words() -> None:
    """Measured in Word: a section that leaves out its size is Letter (not A4, as a document
    without any section is); one that leaves out its margins gets 3 / 1.5 / 2 / 2 cm."""
    source = docx.Document()
    source.add_paragraph("Спецификация")
    sect = source.element.body.find(qn("w:sectPr"))
    for name in ("w:pgSz", "w:pgMar"):
        sect.remove(sect.find(qn(name)))

    document = _document(source)

    assert document.page is not None
    assert (document.page.width_mm, document.page.height_mm) == pytest.approx(
        (215.9, 279.4), abs=0.1
    )
    assert round(document.page.left_mm) == 30
    assert round(document.page.top_mm) == 20


def test_pictures_in_the_body_and_a_header_are_numbered_apart(
    profiles: YamlProfileRepository,
) -> None:
    from PIL import Image

    png = BytesIO()
    Image.new("RGB", (40, 20), "navy").save(png, format="PNG")
    source = docx.Document()
    source.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(png.getvalue()))
    source.add_paragraph().add_run().add_picture(BytesIO(png.getvalue()))
    source.add_paragraph("Заявка")

    importer = DocxImporter()
    document = importer.convert(docx.Document(BytesIO(_bytes(source))))
    extras = RenderExtras(images=importer.images)
    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL, extras)
        )
    )

    drawing = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
    stories = (rendered.element, rendered.sections[0].header._element)
    ids = [d.get("id") for story in stories for d in story.iter(drawing)]
    assert len(ids) == 2  # a repeated id makes Word call the file corrupted
    assert len(set(ids)) == len(ids)


def test_a_row_of_merged_continuations_only_keeps_the_rows_below_in_place() -> None:
    source = docx.Document()
    table = source.add_table(rows=3, cols=2)
    for col, text in enumerate(("Раздел", "Наименование показателя")):
        table.cell(0, col).merge(table.cell(1, col)).text = text  # the header spans two rows
    table.cell(2, 0).text, table.cell(2, 1).text = "1", "Количество мероприятий"

    (block,) = [b for b in _import(source) if isinstance(b, TableBlock)]

    assert len(block.rows) == 3  # the second row has no cell of its own, yet it is a row
    assert [(p.row, p.col) for p in block.placements()] == [(0, 0), (0, 1), (2, 0), (2, 1)]


def _spaced(paragraph: Any) -> None:
    """18 pt before, 10 pt after, and 'no spacing between paragraphs of the same style'."""
    paragraph.paragraph_format.space_before, paragraph.paragraph_format.space_after = Pt(18), Pt(10)
    paragraph._p.get_or_add_pPr().append(docx.oxml.parse_xml(f"<w:contextualSpacing {_W_DECL}/>"))


def test_spacing_of_one_style_is_dropped_across_rows_and_around_a_table() -> None:
    source = docx.Document()
    _spaced(source.add_paragraph("служебная записка."))
    table = source.add_table(rows=2, cols=1)
    for row, text in enumerate(("Прошу в отношении помещения", ", расположенного по адресу:")):
        table.cell(row, 0).text = text
        _spaced(table.cell(row, 0).paragraphs[0])
    _spaced(source.add_paragraph("В ходе указанных выше мероприятий"))

    title, block, after = _import(source)

    assert isinstance(block, TableBlock)
    first, second = (row.cells[0].layout for row in block.rows)
    assert _layout(title).space_after_pt == 0  # the table's first line has its style
    assert _required(first).space_before_pt == 0
    assert _required(first).space_after_pt == 0
    assert _required(second).space_before_pt == 0
    assert _required(second).space_after_pt == 0
    assert _layout(after).space_before_pt == 0


def test_pictures_side_by_side_stay_on_one_line(profiles: YamlProfileRepository) -> None:
    from PIL import Image

    from foto2word.domain.document import ImageBlock

    png = BytesIO()
    Image.new("RGB", (40, 20), "orange").save(png, format="PNG")
    source = docx.Document()
    line = source.add_paragraph()
    line.add_run().add_picture(BytesIO(png.getvalue()))
    line.add_run("   ")
    line.add_run().add_picture(BytesIO(png.getvalue()))

    importer = DocxImporter()
    document = importer.convert(docx.Document(BytesIO(_bytes(source))))
    extras = RenderExtras(images=importer.images)
    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL, extras)
        )
    )

    first, second = [b for b in document.blocks if isinstance(b, ImageBlock)]
    assert not first.same_line
    assert second.same_line
    assert second.lead == "   "  # the blanks that set them apart
    with_pictures = [p for p in rendered.paragraphs if p._p.findall(f".//{qn('w:drawing')}")]
    assert len(with_pictures) == 1  # one line, as in the original
    blanks = next(r for r in with_pictures[0].runs if r.text == "   ")
    assert blanks.font.size == Pt(11)  # the paragraph's type, not the profile's 14 pt


def test_a_frame_around_the_page_is_kept(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    source.add_paragraph("Классный час")
    sect_pr = source.sections[0]._sectPr
    sect_pr.find(qn("w:pgMar")).addnext(
        docx.oxml.parse_xml(
            f'<w:pgBorders {_W_DECL} w:offsetFrom="page">'
            + "".join(
                f'<w:{side} w:val="double" w:sz="4" w:space="24" w:color="auto"/>'
                for side in ("top", "left", "bottom", "right")
            )
            + "</w:pgBorders>"
        )
    )

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    assert document.page is not None
    assert [side for side, _ in document.page.borders] == ["top", "left", "bottom", "right"]
    frame = rendered.sections[0]._sectPr.find(qn("w:pgBorders"))
    assert frame is not None
    assert frame.get(qn("w:offsetFrom")) == "page"
    assert frame.find(qn("w:left")).get(qn("w:val")) == "double"


def test_a_line_break_before_a_picture_keeps_its_empty_line() -> None:
    from PIL import Image

    from foto2word.domain.document import ImageBlock

    png = BytesIO()
    Image.new("RGB", (40, 40), "red").save(png, format="PNG")
    source = docx.Document()
    line = source.add_paragraph()
    line.add_run().add_break()  # the emblem sits one line below the top margin
    line.add_run().add_picture(BytesIO(png.getvalue()))

    (picture,) = [b for b in _import(source) if isinstance(b, ImageBlock)]

    assert picture.lead == "\n"


def test_auto_spacing_set_in_the_document_defaults_is_dropped_at_cell_edges() -> None:
    source = docx.Document()
    defaults = source.styles.element.find(
        f"{qn('w:docDefaults')}/{qn('w:pPrDefault')}/{qn('w:pPr')}"
    )
    for old in defaults.findall(qn("w:spacing")):
        defaults.remove(old)
    defaults.append(  # as a file saved from a web page has it
        docx.oxml.parse_xml(
            f'<w:spacing {_W_DECL} w:before="100" w:beforeAutospacing="1" '
            'w:after="100" w:afterAutospacing="1"/>'
        )
    )
    source.add_table(rows=1, cols=1).cell(0, 0).text = "выдан"

    (table,) = [b for b in _import(source) if isinstance(b, TableBlock)]

    layout = table.rows[0].cells[0].layout
    assert layout is not None
    assert layout.space_before_pt == 0  # Word drops it at the cell's edges
    assert layout.space_after_pt == 0


def test_the_paragraph_mark_keeps_its_size_and_face() -> None:
    source = docx.Document()
    named = source.add_paragraph("(подпись)")
    named._p.get_or_add_pPr().append(
        docx.oxml.parse_xml(
            f'<w:rPr {_W_DECL}><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
            '<w:sz w:val="20"/></w:rPr>'
        )
    )
    source.add_paragraph("Заявление")  # its mark takes the face of the original's theme

    first, second = _import(source)

    assert _layout(first).mark_font == "Calibri"
    assert _layout(first).mark_size_pt == 10
    assert _layout(second).mark_font == "Cambria"  # the minor font of this theme, written out
    assert _layout(second).mark_size_pt is not None


def test_pages_are_numbered_from_where_the_original_starts(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    source.add_paragraph("Чертеж проекта межевания")
    sect_pr = source.sections[0]._sectPr
    sect_pr.find(qn("w:cols")).addprevious(
        docx.oxml.parse_xml(f'<w:pgNumType {_W_DECL} w:start="2"/>')
    )

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    assert document.page is not None
    assert document.page.number_start == 2  # an appendix goes on from page 2
    numbering = rendered.sections[0]._sectPr.find(qn("w:pgNumType"))
    assert numbering is not None
    assert numbering.get(qn("w:start")) == "2"


def test_a_cropped_picture_keeps_its_crop(profiles: YamlProfileRepository) -> None:
    from PIL import Image

    from foto2word.domain.document import ImageBlock

    png = BytesIO()
    Image.new("RGB", (80, 40), "teal").save(png, format="PNG")
    source = docx.Document()
    shape = source.add_paragraph().add_run().add_picture(BytesIO(png.getvalue()))
    fill = shape._inline.find(f".//{qn('pic:blipFill')}")
    fill.find(qn("a:blip")).addnext(
        docx.oxml.parse_xml(
            '<a:srcRect xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'l="38321" t="37342" r="21306" b="10559"/>'
        )
    )

    importer = DocxImporter()
    document = importer.convert(docx.Document(BytesIO(_bytes(source))))
    extras = RenderExtras(images=importer.images)
    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL, extras)
        )
    )

    (picture,) = [b for b in document.blocks if isinstance(b, ImageBlock)]
    assert picture.crop == (0.38321, 0.37342, 0.21306, 0.10559)
    rect = rendered.element.find(f".//{qn('a:srcRect')}")
    assert rect is not None  # only the cropped part shows, as in the original
    assert rect.get("l") == "38321"


def test_a_picture_inside_a_line_of_text_stays_in_the_line(profiles: YamlProfileRepository) -> None:
    from PIL import Image

    png = BytesIO()
    Image.new("RGB", (40, 40), "blue").save(png, format="PNG")
    source = docx.Document()
    line = source.add_paragraph()
    line.add_run().add_picture(BytesIO(png.getvalue()))
    line.add_run("Общероссийская физкультурно-спортивная организация")

    importer = DocxImporter()
    document = importer.convert(docx.Document(BytesIO(_bytes(source))))
    extras = RenderExtras(images=importer.images)
    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL, extras)
        )
    )

    (block,) = document.blocks
    assert _textual(block).runs[0].picture is not None  # the logo before the title, in its line
    assert "Общероссийская" in "".join(r.text for r in _textual(block).runs)
    (paragraph,) = [p for p in rendered.paragraphs if p.text.strip()]
    assert paragraph._p.findall(f".//{qn('w:drawing')}")


def test_an_old_style_text_box_at_its_own_place_stays_a_box() -> None:
    from foto2word.domain.document import TextBoxBlock

    source = docx.Document()
    anchor = source.add_paragraph()
    anchor._p.append(
        docx.oxml.parse_xml(
            f'<w:r {_W_DECL} xmlns:v="urn:schemas-microsoft-com:vml"><w:pict>'
            '<v:rect style="position:absolute;margin-left:390pt;margin-top:0;width:68pt;'
            'height:90pt;z-index:1" fillcolor="#5b9bd5"><v:textbox><w:txbxContent>'
            "<w:p><w:r><w:t>ФОТО</w:t></w:r></w:p>"
            "</w:txbxContent></v:textbox></v:rect></w:pict></w:r>"
        )
    )
    source.add_paragraph("Заявление")

    blocks = _import(source)

    (box,) = [b for b in blocks if isinstance(b, TextBoxBlock)]
    assert box.floating is not None
    assert round(box.width_mm) == 24
    assert box.fill == "5B9BD5"
    assert "ФОТО" not in "".join(b.text() for b in blocks if not isinstance(b, TextBoxBlock))


def _field(code: str, shown: str) -> str:
    return (
        f'<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r><w:instrText xml:space="preserve"> {code} </w:instrText></w:r>'
        f'<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>{shown}</w:t></w:r>'
        f'<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )


def test_page_x_of_y_keeps_its_fields(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    source.add_paragraph("Должностная инструкция")
    footer = source.sections[0].footer.paragraphs[0]
    footer._p.append(
        docx.oxml.parse_xml(f'<w:r {_W_DECL}><w:t xml:space="preserve">Страница </w:t></w:r>')
    )
    for xml in (
        _field("PAGE", "1"),
        '<w:r><w:t xml:space="preserve"> из </w:t></w:r>',
        _field("NUMPAGES", "4"),
    ):
        for run in docx.oxml.parse_xml(f"<w:p {_W_DECL}>{xml}</w:p>"):
            footer._p.append(run)

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    assert document.page_numbers is None  # not just a number: the line keeps its words
    runs = [r for b in document.footer for r in getattr(b, "runs", ())]
    assert [r.field for r in runs if r.field] == ["page", "pages"]
    codes = "".join(t.text for t in rendered.sections[0].footer._element.iter(qn("w:instrText")))
    assert "PAGE" in codes
    assert "NUMPAGES" in codes  # Word counts the pages itself


def _word_marker(paragraph: Any) -> tuple[str, str | None, str]:
    """Text, typeface and what follows of the list marker Word draws for the paragraph."""
    numbering = paragraph.part.numbering_part.element
    num_id = paragraph._p.pPr.numPr.numId.val
    (num,) = [n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == str(num_id)]
    abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    (definition,) = [
        a
        for a in numbering.findall(qn("w:abstractNum"))
        if a.get(qn("w:abstractNumId")) == abstract_id
    ]
    level = definition.find(qn("w:lvl"))
    fonts = level.find(f"{qn('w:rPr')}/{qn('w:rFonts')}")
    suffix = level.find(qn("w:suff"))
    return (
        level.find(qn("w:lvlText")).get(qn("w:val")),
        fonts.get(qn("w:ascii")) if fonts is not None else None,
        suffix.get(qn("w:val")) if suffix is not None else "tab",
    )


def _compat_mode(document: Any) -> str | None:
    found = document.settings.element.find(
        f"{qn('w:compat')}/{qn('w:compatSetting')}[@{qn('w:name')}='compatibilityMode']"
    )
    return found.get(qn("w:val")) if found is not None else None


@pytest.mark.parametrize("mode", ["14", "15"])
def test_the_word_version_layout_is_kept(profiles: YamlProfileRepository, mode: str) -> None:
    source = docx.Document()
    source.add_paragraph("Текст")
    setting = source.settings.element.find(
        f"{qn('w:compat')}/{qn('w:compatSetting')}[@{qn('w:name')}='compatibilityMode']"
    )
    setting.set(qn("w:val"), mode)  # Word 2013 breaks lines differently from Word 2010

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    assert _compat_mode(rendered) == mode


def test_each_paragraph_of_a_cell_keeps_its_own_spacing(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    cell = source.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].text = "Общество с ограниченной ответственностью"
    cell.paragraphs[0].paragraph_format.line_spacing = 1.5
    second = cell.add_paragraph("ИНН 9715314185")
    second.paragraph_format.line_spacing = 1.0

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    lines = [p.paragraph_format.line_spacing for p in rendered.tables[0].cell(0, 0).paragraphs]
    assert lines == [1.5, 1.0]


def test_an_old_file_without_a_word_version_keeps_the_old_layout(
    profiles: YamlProfileRepository,
) -> None:
    source = docx.Document()
    source.add_paragraph("Текст")
    compat = source.settings.element.find(qn("w:compat"))
    compat.getparent().remove(compat)  # saved by Word 2007 or older: no compatibility mode

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    assert document.compat_mode == 12
    assert _compat_mode(rendered) == "12"


def test_the_original_compatibility_switches_are_kept_and_no_others(
    profiles: YamlProfileRepository,
) -> None:
    from docx.oxml import OxmlElement

    source = docx.Document()
    source.add_paragraph("Текст")
    compat = source.settings.element.find(qn("w:compat"))
    for child in list(compat):
        compat.remove(child)
    compat.append(OxmlElement("w:doNotExpandShiftReturn"))  # one legacy switch, no settings

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    out = rendered.settings.element.find(qn("w:compat"))
    names = [child.tag.split("}")[1] for child in out]
    assert names == ["doNotExpandShiftReturn", "compatSetting"]
    assert _compat_mode(rendered) == "12"


def _contextual(paragraph: Any) -> None:
    from docx.oxml import OxmlElement

    paragraph._p.get_or_add_pPr().append(OxmlElement("w:contextualSpacing"))


@pytest.mark.parametrize("same_style", [True, False])
def test_no_space_between_paragraphs_of_one_style_when_asked(same_style: bool) -> None:
    source = docx.Document()
    first = source.add_paragraph("*-заявитель дает согласие на использование")
    second = source.add_paragraph(
        "указанных контактных данных", style=None if same_style else "Quote"
    )
    for paragraph in (first, second):
        paragraph.paragraph_format.space_after = Pt(8)
        _contextual(paragraph)

    blocks = _import(source)

    gap = blocks[0].layout.space_after_pt  # type: ignore[union-attr]
    assert gap == (0.0 if same_style else 8.0)


def _note_and_empty_lines(contextual: bool) -> list[Block]:
    source = docx.Document()
    note = source.add_paragraph("для дальнейшей связи с ним.")
    empties = [source.add_paragraph("") for _ in range(2)]
    source.add_paragraph("Заявление", style="Title")
    for paragraph in (note, *empties):
        paragraph.paragraph_format.space_after = Pt(8)
        if contextual:
            _contextual(paragraph)
    return _import(source)


def test_empty_lines_of_one_style_add_no_space_either() -> None:
    plain, asked = _note_and_empty_lines(False), _note_and_empty_lines(True)

    assert asked[0].layout.space_after_pt == 0.0  # type: ignore[union-attr]
    assert isinstance(plain[1], SpacerBlock)
    assert isinstance(asked[1], SpacerBlock)
    # between the two empty lines the 8 pt goes; after the last one it stays (next style differs)
    assert (plain[1].height_pt or 0) - (asked[1].height_pt or 0) == pytest.approx(8, abs=0.2)


def test_a_header_that_starts_with_a_table_starts_with_it(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    header = source.sections[0].header
    table = header.add_table(rows=1, cols=2, width=Mm(170))
    table.cell(0, 0).text, table.cell(0, 1).text = "Карго", "ИНН 9715314185"
    first = header.paragraphs[0]._p
    first.getparent().remove(first)  # the header holds the table first, as Word saves it
    header._element.append(first)
    source.add_paragraph("Заявление")

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    children = list(rendered.sections[0].header._element.iterchildren(qn("w:p"), qn("w:tbl")))
    assert children[0].tag == qn("w:tbl")  # no empty line above the table


def test_a_picture_in_a_cell_stays_in_its_line(profiles: YamlProfileRepository) -> None:
    from PIL import Image

    png = BytesIO()
    Image.new("RGB", (60, 30), "orange").save(png, format="PNG")
    source = docx.Document()
    cell = source.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run().add_picture(BytesIO(png.getvalue()), width=Mm(30))
    cell.add_paragraph("Грузоперевозки по России")

    importer = DocxImporter()
    document = importer.convert(docx.Document(BytesIO(_bytes(source))))
    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(
                document,
                profiles.get("generic"),
                RenderMode.FAITHFUL,
                RenderExtras(images=importer.images),
            )
        )
    )

    paragraphs = rendered.tables[0].cell(0, 0).paragraphs
    assert next(paragraphs[0]._p.iter(qn("wp:inline")), None) is not None  # the logo
    assert paragraphs[1].text == "Грузоперевозки по России"


def test_requisites_keep_their_typeface_and_centring(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    for text in ("Генеральному директору ООО «Карго-Терминал»", "Иванову И.И."):
        line = source.add_paragraph()
        line.add_run(text).font.name = "Segoe UI"
        line.paragraph_format.left_indent = Mm(95)
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    first = rendered.paragraphs[0]
    assert first.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert round(first.paragraph_format.left_indent.mm) == 95
    assert first.runs[0].font.name == "Segoe UI"


@pytest.mark.parametrize(("font", "factor"), [("Times New Roman", 1.15), ("Segoe UI", 1.33)])
def test_an_empty_line_is_as_tall_as_its_typeface_makes_it(font: str, factor: float) -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    source = docx.Document()
    source.add_paragraph("Генеральному директору")
    empty = source.add_paragraph()
    empty.paragraph_format.space_after = Pt(0)
    empty.paragraph_format.line_spacing = 1.0
    empty._p.get_or_add_pPr().append(
        parse_xml(
            f'<w:rPr {nsdecls("w")}><w:rFonts w:ascii="{font}" w:hAnsi="{font}"/>'
            '<w:sz w:val="22"/></w:rPr>'
        )
    )
    source.add_paragraph("от ____")
    for paragraph in source.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)

    spacer = next(b for b in _import(source) if isinstance(b, SpacerBlock))

    assert spacer.height_pt == pytest.approx(11 * factor, abs=0.3)


def test_each_cell_line_keeps_its_alignment_and_an_empty_line_its_size(
    profiles: YamlProfileRepository,
) -> None:
    from docx.oxml import OxmlElement

    source = docx.Document()
    cell = source.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("от Иванова").font.size = Pt(12)
    caption = cell.add_paragraph()
    caption.add_run("(фамилия, имя, отчество)").font.size = Pt(8)
    empty = cell.add_paragraph()
    for line in (caption, empty):
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mark, size = OxmlElement("w:rPr"), OxmlElement("w:sz")
    size.set(qn("w:val"), "16")
    mark.append(size)
    empty._p.get_or_add_pPr().append(mark)

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    lines = rendered.tables[0].cell(0, 0).paragraphs
    assert [p.alignment for p in lines] == [
        WD_ALIGN_PARAGRAPH.LEFT,
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.CENTER,
    ]
    size = lines[2]._p.pPr.find(f"{qn('w:rPr')}/{qn('w:sz')}")
    assert size is not None
    assert size.get(qn("w:val")) == "16"  # an empty line is as tall as its own mark


def test_spaces_that_push_a_line_to_the_right_are_kept(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    source.add_paragraph(" " * 60 + "Заведующему МАДОУ № 97   ")
    source.add_paragraph("\tГалыгиной Г.А.")

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    lines = [p.text for p in rendered.paragraphs if p.text.strip()]
    assert lines == [" " * 60 + "Заведующему МАДОУ № 97", "\tГалыгиной Г.А."]


def test_a_picture_python_docx_cannot_read_goes_in_as_its_own_bytes(
    profiles: YamlProfileRepository,
) -> None:
    from foto2word.domain.document import DocumentTypeGuess, ImageBlock
    from foto2word.domain.extras import RenderExtras

    drawing = b"\x01\x00\x00\x00" + bytes(36) + b" EMF" + bytes(40)
    document = Document(
        document_type=DocumentTypeGuess(id="other", title="Схема", confidence=1.0),
        blocks=(ImageBlock(id="b1", description="схема", width_mm=50, height_mm=20),),
    )
    extras = RenderExtras(images={"b1": drawing})

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL, extras)
        )
    )

    (shape,) = rendered.inline_shapes
    assert round(shape.width.mm) == 50
    assert round(shape.height.mm) == 20
    parts = [p for p in rendered.part.package.iter_parts() if p.content_type == "image/x-emf"]
    assert [p.blob for p in parts] == [drawing]


def test_a_white_cell_line_is_not_drawn(profiles: YamlProfileRepository) -> None:
    from docx.oxml import OxmlElement

    source = docx.Document()
    table = source.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Дата"
    borders = OxmlElement("w:tcBorders")
    for side, color in (("top", "000000"), ("bottom", "FFFFFF")):
        line = OxmlElement(f"w:{side}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:color"), color)
        borders.append(line)
    table.cell(0, 0)._tc.get_or_add_tcPr().append(borders)

    (imported,) = [b for b in _import(source) if isinstance(b, TableBlock)]

    (cell,) = imported.rows[0].cells
    assert cell.borders is not None
    assert cell.borders.top
    assert not cell.borders.bottom  # white on a white page: the original shows no line there


def test_text_inside_content_controls_is_kept() -> None:
    from docx.oxml import OxmlElement

    def wrapped(element: Any) -> Any:
        control, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
        element.addprevious(control)
        content.append(element)
        control.append(content)
        return control

    source = docx.Document()
    wrapped(source.add_paragraph("Заявление")._p)
    line = source.add_paragraph("Дата: ")
    wrapped(line.add_run("20.10.2026")._r)
    cell = source.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("Подпись")
    wrapped(cell.paragraphs[0]._p)

    document = _document(source)

    texts = [str(b.model_dump()) for b in document.blocks]
    assert any("Заявление" in t for t in texts)
    assert any("20.10.2026" in t for t in texts)
    assert any("Подпись" in t for t in texts)


def test_a_title_page_without_a_number_keeps_it_off_the_first_page(
    profiles: YamlProfileRepository,
) -> None:
    from docx.oxml import OxmlElement

    source = docx.Document()
    section = source.sections[0]
    section.different_first_page_header_footer = True  # the first page has its own, empty one
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " PAGE ")
    section.header.paragraphs[0]._p.append(field)
    source.add_paragraph("Гарантийное письмо")

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    assert document.page_numbers is not None
    assert not document.page_numbers.first_page
    out = rendered.sections[0]
    assert out.different_first_page_header_footer
    assert not list(out.first_page_header._element.iter(qn("w:instrText")))
    assert [t.text for t in out.header._element.iter(qn("w:instrText"))] == [" PAGE "]


def test_cell_shading_is_kept(profiles: YamlProfileRepository) -> None:
    from docx.oxml import OxmlElement

    source = docx.Document()
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Дата формирования", ""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), "d9d9d9")
    table.cell(0, 1)._tc.get_or_add_tcPr().append(shading)

    (imported,) = [b for b in _import(source) if isinstance(b, TableBlock)]
    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    first, second = imported.rows[0].cells
    assert first.shading is None
    assert second.shading == "D9D9D9"
    fill = _required(rendered.tables[0].cell(0, 1)._tc.tcPr).find(qn("w:shd"))
    assert fill is not None
    assert fill.get(qn("w:fill")) == "D9D9D9"


def test_highlighted_and_shaded_text_keeps_its_background(
    profiles: YamlProfileRepository,
) -> None:
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml import OxmlElement

    source = docx.Document()
    line = source.add_paragraph("Заполните: ")
    line.add_run("ФИО").font.highlight_color = WD_COLOR_INDEX.YELLOW
    shaded = line.add_run(" дата")
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), "D9D9D9")
    shaded._r.get_or_add_rPr().append(shading)

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    block = document.blocks[0]
    assert isinstance(block, ParagraphBlock)
    assert [r.background for r in block.runs] == [None, "FFFF00", "D9D9D9"]
    fills = [s.get(qn("w:fill")) for s in rendered.paragraphs[0]._p.iter(qn("w:shd"))]
    assert fills == ["FFFF00", "D9D9D9"]


def test_a_grey_bar_behind_a_heading_is_kept(profiles: YamlProfileRepository) -> None:
    from docx.oxml import OxmlElement

    source = docx.Document()
    heading = source.add_paragraph("1. Сведения о заявителе")
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), "D9D9D9")
    heading._p.get_or_add_pPr().append(shading)

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    block = document.blocks[0]
    assert isinstance(block, ListItemBlock)  # "1. ..." is a numbered line
    assert block.layout is not None
    assert block.layout.shading == "D9D9D9"
    fill = _required(rendered.paragraphs[0]._p.pPr).find(qn("w:shd"))
    assert fill is not None
    assert fill.get(qn("w:fill")) == "D9D9D9"


def test_an_underlined_blank_at_the_end_of_a_line_stays() -> None:
    source = docx.Document()
    line = source.add_paragraph("Подпись: ")
    line.add_run("\t\t").font.underline = True
    source.add_paragraph("Дата:   ")

    blocks = _import(source)

    texts = ["".join(r.text for r in b.runs) for b in blocks if isinstance(b, ParagraphBlock)]
    assert texts == ["Подпись: \t\t", "Дата:"]  # the line to sign on is seen; bare spaces not


_W_DECL = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _numbered(source: docx.document.Document, levels: tuple[tuple[str, str], ...]) -> str:
    """A list definition with these (numFmt, lvlText) levels; returns its numId."""
    numbering = source.part.numbering_part.element
    levels_xml = "".join(
        f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
        f'<w:lvlText w:val="{text}"/></w:lvl>'
        for i, (fmt, text) in enumerate(levels)
    )
    definition = f'<w:abstractNum {_W_DECL} w:abstractNumId="90">{levels_xml}</w:abstractNum>'
    numbering.insert(0, docx.oxml.parse_xml(definition))
    numbering.append(
        docx.oxml.parse_xml(f'<w:num {_W_DECL} w:numId="90"><w:abstractNumId w:val="90"/></w:num>')
    )
    return "90"


def _number(paragraph: Any, num_id: str, level: int) -> None:
    paragraph._p.get_or_add_pPr().append(
        docx.oxml.parse_xml(
            f'<w:numPr {_W_DECL}><w:ilvl w:val="{level}"/><w:numId w:val="{num_id}"/></w:numPr>'
        )
    )


def test_multi_level_numbers_count_in_the_body_and_in_cells() -> None:
    source = docx.Document()
    num_id = _numbered(source, (("decimal", "%1."), ("decimal", "%1.%2.")))
    _number(source.add_paragraph("Сведения о заявителе"), num_id, 0)
    cell = source.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("Фамилия, имя, отчество:")
    _number(cell.paragraphs[0], num_id, 1)
    _number(source.add_paragraph("Состав сведений"), num_id, 0)
    _number(source.add_paragraph("Сведения о заявителе"), num_id, 1)

    blocks = _import(source)

    markers = [b.marker for b in blocks if isinstance(b, ListItemBlock)]
    assert markers == ["1.", "2.", "2.1."]
    (table,) = [b for b in blocks if isinstance(b, TableBlock)]
    assert table.rows[0].cells[0].text().startswith("1.1.\tФамилия")


def test_a_numbered_cell_line_hangs_its_number_like_the_list_level() -> None:
    source = docx.Document()
    numbering = source.part.numbering_part.element
    numbering.insert(
        0,
        docx.oxml.parse_xml(
            f'<w:abstractNum {_W_DECL} w:abstractNumId="91"><w:lvl w:ilvl="0"><w:start w:val="1"/>'
            '<w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/>'
            '<w:pPr><w:ind w:left="360" w:hanging="360"/></w:pPr></w:lvl></w:abstractNum>'
        ),
    )
    numbering.append(
        docx.oxml.parse_xml(f'<w:num {_W_DECL} w:numId="91"><w:abstractNumId w:val="91"/></w:num>')
    )
    cell = source.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("Организационный этап")
    _number(cell.paragraphs[0], "91", 0)

    (table,) = [b for b in _import(source) if isinstance(b, TableBlock)]

    layout = table.rows[0].cells[0].layout
    assert layout is not None
    assert layout.left_indent_mm == pytest.approx(6.35, abs=0.06)
    assert layout.first_line_indent_mm == pytest.approx(-6.35, abs=0.06)


def test_default_tab_stops_keep_the_original_step(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    source.settings.element.find(qn("w:defaultTabStop")).set(qn("w:val"), "708")
    source.add_paragraph("Фамилия\tИмя")

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    assert document.default_tab_pt == pytest.approx(35.4)
    step = rendered.settings.element.find(qn("w:defaultTabStop"))
    assert step.get(qn("w:val")) == "708"


def test_a_line_with_a_wide_gap_stays_a_paragraph() -> None:
    """Measured on 411 files: a line kept with its own tabs and spaces lands where Word puts
    it (156 better, 5 slightly worse); rebuilt as two columns it lost its typeface, its line
    breaks, and wrapped otherwise when a hundred spaces did not fit on one line."""
    source = docx.Document()
    title = source.add_paragraph("Образец заявления" + " " * 20 + "изменение формы ГИА")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mixed = source.add_paragraph("выдан" + " " * 10)
    mixed.add_run("федеральным учреждением").bold = True
    source.add_paragraph("г. Москва" + " " * 30 + "20 октября 2026 г.")

    kinds = [b.kind for b in _import(source)]

    assert kinds == ["paragraph", "paragraph", "paragraph"]


def test_lines_pushed_along_by_spaces_are_not_a_requisites_block() -> None:
    """'Приложение 1' after ninety spaces: in Word the spaces take a line of their own, so the
    lines stay as written instead of a block that drops the spaces."""
    source = docx.Document()
    for line in (" " * 90 + "Приложение 1", "к Положению о конкурсе"):
        source.add_paragraph(line).paragraph_format.left_indent = Mm(80)

    blocks = _import(source)

    assert not [b for b in blocks if isinstance(b, RequisitesBlock)]
    assert blocks[0].text().startswith(" " * 90)


def test_requisites_keep_letter_spacing_and_justification(
    profiles: YamlProfileRepository,
) -> None:
    source = docx.Document()
    for line in ("Членам комиссии по предупреждению", "г. Грозного"):
        paragraph = source.add_paragraph()
        paragraph.paragraph_format.left_indent = Mm(100)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = paragraph.add_run(line)
        spacing = docx.oxml.OxmlElement("w:spacing")
        spacing.set(qn("w:val"), "-12")
        run._r.get_or_add_rPr().append(spacing)

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    block = document.blocks[0]
    assert _textual(block).runs[0].spacing_pt == pytest.approx(-0.6)
    assert _textual(block).align is Alignment.JUSTIFY
    first = rendered.paragraphs[0]
    assert first.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    letters = _required(first.runs[0]._r.rPr).find(qn("w:spacing"))
    assert letters is not None
    assert letters.get(qn("w:val")) == "-12"


def test_a_right_aligned_line_ending_in_a_tab_leader_keeps_its_blank() -> None:
    source = docx.Document()
    line = source.add_paragraph("от\t")
    line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    line.paragraph_format.tab_stops.add_tab_stop(
        Mm(110), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES
    )

    (block,) = _import(source)

    assert isinstance(block, ParagraphBlock)
    assert "".join(r.text for r in block.runs) == "от\t"


def test_tables_stay_apart_without_a_spacer_before_text(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    source.add_table(rows=1, cols=1).cell(0, 0).text = "ИНН"
    source.add_paragraph("(ИНН при наличии)")
    source.add_table(rows=1, cols=1).cell(0, 0).text = "КПП"
    source.add_table(rows=1, cols=1).cell(0, 0).text = "ОГРН"

    rendered = DocxRenderer().render(
        _document(source), profiles.get("generic"), RenderMode.FAITHFUL
    )

    body = docx.Document(BytesIO(rendered)).element.body
    kinds = [
        "tbl" if child.tag == qn("w:tbl") else "p:" + "".join(t.text for t in child.iter(qn("w:t")))
        for child in body
        if child.tag in (qn("w:tbl"), qn("w:p"))
    ]
    # text follows a table directly; two tables need a paragraph between them, and so does the end
    assert kinds == ["tbl", "p:(ИНН при наличии)", "tbl", "p:", "tbl", "p:"]
    assert b"f2w-" not in _document_xml(rendered)


def test_a_cell_keeps_its_own_margins(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "информационную систему ______", "x"
    table.cell(0, 0)._tc.get_or_add_tcPr().append(
        docx.oxml.parse_xml(
            f'<w:tcMar {_W_DECL}><w:left w:w="0" w:type="dxa"/><w:right w:w="0" w:type="dxa"/>'
            "</w:tcMar>"
        )
    )

    (imported,) = [b for b in _import(source) if isinstance(b, TableBlock)]
    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    first, second = imported.rows[0].cells
    assert first.padding is not None
    assert (first.padding.left_mm, first.padding.right_mm, first.padding.top_mm) == (0, 0, 0)
    assert second.padding is None  # the table's margins
    margins = _required(rendered.tables[0].cell(0, 0)._tc.tcPr).find(qn("w:tcMar"))
    assert margins is not None
    assert margins.find(qn("w:left")).get(qn("w:w")) == "0"


def test_a_font_missing_here_keeps_its_description_for_word_to_substitute(
    profiles: YamlProfileRepository,
) -> None:
    import zipfile

    source = docx.Document()
    source.add_paragraph().add_run("Заявителем может быть").font.name = "Liberation Sans"
    part = next(
        p for p in source.part.package.iter_parts() if str(p.partname).endswith("fontTable.xml")
    )
    part._blob = part.blob.replace(
        b"</w:fonts>",
        b'<w:font w:name="Liberation Sans"><w:panose1 w:val="020B0604020202020204"/>'
        b'<w:charset w:val="CC"/><w:family w:val="swiss"/><w:pitch w:val="variable"/></w:font>'
        b"</w:fonts>",
    )

    document = _document(source)
    rendered = DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL)

    (face,) = [f for f in document.fonts if f.name == "Liberation Sans"]
    assert (face.family, face.panose) == ("swiss", "020B0604020202020204")
    table = zipfile.ZipFile(BytesIO(rendered)).read("word/fontTable.xml")
    assert b'w:name="Liberation Sans"' in table
    assert b'w:val="020B0604020202020204"' in table


def test_a_font_keeps_the_scripts_it_covers() -> None:
    source = docx.Document()
    source.add_paragraph().add_run("Уважаемые родители").font.name = "Source Sans Pro"
    parts = source.part.package.iter_parts()
    part = next(p for p in parts if str(p.partname).endswith("fontTable.xml"))
    part._blob = part.blob.replace(
        b"</w:fonts>",
        b'<w:font w:name="Source Sans Pro"><w:charset w:val="00"/><w:family w:val="swiss"/>'
        b'<w:pitch w:val="variable"/><w:sig w:usb0="600002F7" w:usb1="02000001" w:usb2="0" '
        b'w:usb3="0" w:csb0="0000019F" w:csb1="0"/></w:font></w:fonts>',
    )

    (face,) = [f for f in _document(source).fonts if f.name == "Source Sans Pro"]

    assert dict(face.signature)["usb0"] == "600002F7"


def test_auto_spacing_does_not_push_a_cell_text_from_the_cell_edges() -> None:
    source = docx.Document()
    cell = source.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("Предоставление земельных участков")
    cell.add_paragraph("в аренду")
    for paragraph in cell.paragraphs:
        paragraph._p.get_or_add_pPr().append(
            docx.oxml.parse_xml(
                f'<w:spacing {_W_DECL} w:before="100" w:beforeAutospacing="1" '
                'w:after="100" w:afterAutospacing="1"/>'
            )
        )

    (table,) = [b for b in _import(source) if isinstance(b, TableBlock)]

    layouts = table.rows[0].cells[0].line_layouts
    assert layouts[0].space_before_pt == 0  # Word drops it at the top of a cell
    assert (layouts[0].space_after_pt or 0) > 0
    assert (layouts[1].space_before_pt or 0) > 0
    assert layouts[1].space_after_pt == 0  # and at its bottom


def test_a_cell_drops_spacing_between_lines_of_one_style_that_ask_so() -> None:
    source = docx.Document()
    cell = source.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("Пример:")
    for text in ("полный рабочий день;", "оформление по ТК;"):
        item = cell.add_paragraph(text, style="List Paragraph")
        item.paragraph_format.space_after = Pt(6)
        item._p.get_or_add_pPr().append(docx.oxml.parse_xml(f"<w:contextualSpacing {_W_DECL}/>"))

    (table,) = [b for b in _import(source) if isinstance(b, TableBlock)]

    layouts = table.rows[0].cells[0].line_layouts
    assert layouts[1].space_after_pt == 0  # 'don't add space between paragraphs of one style'
    assert layouts[2].space_after_pt == 6  # the last one has no neighbour of its style


_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def test_an_old_style_watermark_floats_behind_the_text() -> None:
    from PIL import Image

    from foto2word.domain.document import ImageBlock

    picture = BytesIO()
    Image.new("RGB", (40, 20), "red").save(picture, format="PNG")
    source = docx.Document()
    run = source.add_paragraph().add_run()
    run.add_picture(BytesIO(picture.getvalue()))
    rel_id = next(run._r.iter(f"{{{_A_NS}}}blip")).get(qn("r:embed"))
    for drawing in run._r.findall(qn("w:drawing")):
        run._r.remove(drawing)
    run._r.append(
        docx.oxml.parse_xml(
            f'<w:pict {_W_DECL} xmlns:v="urn:schemas-microsoft-com:vml" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<v:shape style="position:absolute;margin-left:0;margin-top:0;width:186pt;'
            "height:263pt;z-index:-251653120;mso-position-horizontal:center;"
            "mso-position-horizontal-relative:margin;mso-position-vertical:center;"
            f'mso-position-vertical-relative:margin"><v:imagedata r:id="{rel_id}"/></v:shape>'
            "</w:pict>"
        )
    )
    source.add_paragraph("Что знают челябинцы о переписи населения?")

    image = next(b for b in _import(source) if isinstance(b, ImageBlock))

    assert image.floating is not None
    assert image.floating.behind_text
    assert (image.floating.horizontal_from, image.floating.x_align) == ("margin", "center")
    assert (image.floating.vertical_from, image.floating.y_align) == ("margin", "center")
    assert image.width_mm == pytest.approx(65.62, abs=0.01)


def test_lines_keep_their_colour_width_and_kind(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    cell = source.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Организация"
    cell._tc.get_or_add_tcPr().append(
        docx.oxml.parse_xml(
            f'<w:tcBorders {_W_DECL}><w:bottom w:val="single" w:sz="4" w:color="A6A6A6"/>'
            "</w:tcBorders>"
        )
    )
    note = source.add_paragraph("Если производительность неизвестна")
    note._p.get_or_add_pPr().append(
        docx.oxml.parse_xml(
            f'<w:pBdr {_W_DECL}><w:left w:val="single" w:sz="24" w:space="4" w:color="4FC3BD"/>'
            "</w:pBdr>"
        )
    )

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    table = next(b for b in document.blocks if isinstance(b, TableBlock))
    borders = table.rows[0].cells[0].borders
    assert borders is not None
    assert dict(borders.lines)["bottom"].color == "A6A6A6"
    edge = _required(rendered.tables[0].cell(0, 0)._tc.tcPr).find(
        f"{qn('w:tcBorders')}/{qn('w:bottom')}"
    )
    assert edge.get(qn("w:color")) == "A6A6A6"
    bar = _required(rendered.paragraphs[-1]._p.pPr).find(f"{qn('w:pBdr')}/{qn('w:left')}")
    assert (bar.get(qn("w:sz")), bar.get(qn("w:space")), bar.get(qn("w:color"))) == (
        "24",
        "4",
        "4FC3BD",
    )


def test_empty_paragraphs_stay_as_they_are_for_word_to_lay_out(
    profiles: YamlProfileRepository,
) -> None:
    from foto2word.domain.document import SpacerBlock

    source = docx.Document()
    source.add_paragraph("ОПИСЬ ДОКУМЕНТОВ")
    for half_points in ("16", "28"):
        empty = source.add_paragraph()
        empty._p.get_or_add_pPr().append(
            docx.oxml.parse_xml(f'<w:rPr {_W_DECL}><w:sz w:val="{half_points}"/></w:rPr>')
        )
    source.add_paragraph("Настоящим подтверждается")

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    (gap,) = [b for b in document.blocks if isinstance(b, SpacerBlock)]
    assert [line.size_pt for line in gap.lines] == [8, 14]
    sizes = [
        _required(p._p.pPr).find(f"{qn('w:rPr')}/{qn('w:sz')}").get(qn("w:val"))
        for p in rendered.paragraphs
        if not p.text
    ]
    assert sizes == ["16", "28"]


def test_a_line_after_empty_ones_of_its_style_drops_its_contextual_spacing() -> None:
    source = docx.Document()
    for text in ("краткосрочным кредитам (займам)", "", "«___» ____________20____года"):
        paragraph = source.add_paragraph(text)
        paragraph._p.get_or_add_pPr().append(
            docx.oxml.parse_xml(
                f'<w:spacing {_W_DECL} w:before="100" w:beforeAutospacing="1" w:after="0"/>'
            )
        )
        paragraph._p.get_or_add_pPr().append(
            docx.oxml.parse_xml(f"<w:contextualSpacing {_W_DECL}/>")
        )

    blocks = _import(source)

    last = blocks[-1]
    assert isinstance(last, ParagraphBlock)
    assert last.layout is not None
    assert last.layout.space_before_pt == 0  # its neighbour is the empty line of its style


def test_right_aligned_requisites_keep_the_indents_that_wrap_them(
    profiles: YamlProfileRepository,
) -> None:
    source = docx.Document()
    line = source.add_paragraph("Приложение 1 к документации об аукционе в электронной форме")
    line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    line.paragraph_format.left_indent = Mm(127)
    line.paragraph_format.first_line_indent = Mm(26)

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    fmt = rendered.paragraphs[0].paragraph_format
    assert round(fmt.left_indent.mm) == 127
    assert round(fmt.first_line_indent.mm) == 26


def test_a_header_of_several_empty_lines_takes_its_room() -> None:
    source = docx.Document()
    section = source.sections[0]
    section.top_margin, section.header_distance = Mm(2.5), Mm(12.5)
    for _ in range(2):  # measured in Word: one bare empty line is ignored, more are not
        section.header.add_paragraph()
    source.add_paragraph("СПИСОК АФФИЛИРОВАННЫХ ЛИЦ")

    document = _document(source)

    assert document.header


def test_the_line_a_framed_page_number_hangs_from_takes_its_room() -> None:
    """Measured in Word (74ff206f): the bare line after a framed number moves the text down by
    a line (37.2 mm against 32.4 mm without the header), unlike a lone bare line."""
    source = docx.Document()
    section = source.sections[0]
    section.top_margin, section.header_distance = Mm(20), Mm(20)
    framed = section.header.paragraphs[0]
    framed._p.get_or_add_pPr().append(
        docx.oxml.parse_xml(f'<w:framePr {_W_DECL} w:wrap="around" w:hAnchor="margin" w:y="1"/>')
    )
    section.header.add_paragraph()
    source.add_paragraph("ОПИСЬ ДОКУМЕНТОВ")

    document = _document(source)

    assert document.header  # the frame hangs from it: Word keeps its room


def test_an_old_style_picture_sized_in_pixels_keeps_its_size() -> None:
    from PIL import Image

    from foto2word.domain.document import ImageBlock

    picture = BytesIO()
    Image.new("RGB", (40, 20), "red").save(picture, format="PNG")
    source = docx.Document()
    run = source.add_paragraph().add_run()
    run.add_picture(BytesIO(picture.getvalue()))
    rel_id = next(run._r.iter(f"{{{_A_NS}}}blip")).get(qn("r:embed"))
    for drawing in run._r.findall(qn("w:drawing")):
        run._r.remove(drawing)
    run._r.append(
        docx.oxml.parse_xml(
            f'<w:pict {_W_DECL} xmlns:v="urn:schemas-microsoft-com:vml" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            f'<v:shape style="width:675px;height:450px"><v:imagedata r:id="{rel_id}"/></v:shape>'
            "</w:pict>"
        )
    )

    image = next(b for b in _import(source) if isinstance(b, ImageBlock))

    assert image.width_mm == pytest.approx(178.59, abs=0.01)  # 675 px = 506.25 pt
    assert image.height_mm == pytest.approx(119.06, abs=0.01)


def test_a_line_break_in_a_cell_stays_inside_its_paragraph(
    profiles: YamlProfileRepository,
) -> None:
    source = docx.Document()
    line = source.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    line.add_run("Директору МБОУ СОШ № 15").font.size = Pt(12)
    line.add_run().add_break()  # a run of its own, as Word often writes it
    line.add_run("Корытцевой Марине Михайловне").font.size = Pt(12)

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    paragraphs = rendered.tables[0].cell(0, 0).paragraphs
    assert [p.text for p in paragraphs] == ["Директору МБОУ СОШ № 15\nКорытцевой Марине Михайловне"]


def _with_footnote(source: docx.document.Document, text: str, note: str) -> None:
    """A paragraph with a footnote mark at its end, and the note in a footnotes part."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part

    paragraph = source.add_paragraph(text)
    paragraph._p.append(
        docx.oxml.parse_xml(
            f'<w:r {_W_DECL}><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
            '<w:footnoteReference w:id="1"/></w:r>'
        )
    )
    notes = (
        f'<w:footnotes {_W_DECL}><w:footnote w:id="1"><w:p><w:r><w:rPr>'
        '<w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r>'
        f'<w:r><w:t xml:space="preserve"> {note}</w:t></w:r></w:p></w:footnote></w:footnotes>'
    ).encode()
    part = Part(
        PackURI("/word/footnotes.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        notes,
        source.part.package,
    )
    source.part.relate_to(part, RT.FOOTNOTES)


def test_footnotes_keep_their_marks_and_notes(profiles: YamlProfileRepository) -> None:
    import zipfile

    source = docx.Document()
    _with_footnote(source, "Состав сведений", "Разделы, в которые Заявитель не вносит изменения")

    document = _document(source)
    rendered = DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL)

    (note,) = document.footnotes
    assert note.id == 1
    text = "".join(r.text for p in note.paragraphs for r in p.runs)
    assert "Разделы, в которые Заявитель" in text
    block = document.blocks[0]
    assert isinstance(block, ParagraphBlock)
    assert [r.footnote for r in block.runs] == [None, 1]
    archive = zipfile.ZipFile(BytesIO(rendered))
    assert b'<w:footnoteReference w:id="1"/>' in archive.read("word/document.xml")
    notes = archive.read("word/footnotes.xml")
    assert "Разделы, в которые Заявитель".encode() in notes
    assert b"<w:footnoteRef/>" in notes
    assert b"footnotes" in archive.read("word/_rels/document.xml.rels")


def test_a_column_break_and_columns_of_their_own_widths_are_kept(
    profiles: YamlProfileRepository,
) -> None:
    from docx.enum.text import WD_BREAK

    source = docx.Document()
    sect_pr = source.sections[0]._sectPr
    sect_pr.remove(sect_pr.find(qn("w:cols")))
    sect_pr.append(
        docx.oxml.parse_xml(
            f'<w:cols {_W_DECL} w:num="2" w:space="340" w:equalWidth="0">'
            '<w:col w:w="4536" w:space="340"/><w:col w:w="4479"/></w:cols>'
        )
    )
    source.add_paragraph("Заявление поступило:")
    ending = source.add_paragraph("почтовая связь")
    ending.add_run().add_break(WD_BREAK.COLUMN)
    source.add_paragraph("(наименование органа)")

    document = _document(source)
    rendered = DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL)

    first, _, third = document.blocks
    assert first.section is not None
    assert [round(w) for w, _ in first.section.column_widths_mm] == [80, 79]
    assert third.column_break_before
    xml = _document_xml(rendered)
    assert b'w:type="column"' in xml
    assert b'w:equalWidth="0"' in xml


def test_a_paragraph_starting_with_a_column_break_starts_the_next_column_whole(
    profiles: YamlProfileRepository,
) -> None:
    """Word sets it, spacing before included, in the next column: the break goes inside it,
    and an empty one holding only the break is the next column's first line (9144a589)."""
    from docx.enum.text import WD_BREAK

    source = docx.Document()
    source.sections[0]._sectPr.find(qn("w:cols")).set(qn("w:num"), "2")
    source.add_paragraph("Гарантийный талон")
    title = source.add_paragraph()
    title.paragraph_format.space_before = Pt(4)
    title.add_run().add_break(WD_BREAK.COLUMN)
    title.add_run("ПАСПОРТ")

    document = _document(source)
    rendered = DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL)

    assert document.blocks[1].column_break_before
    body = docx.Document(BytesIO(rendered)).paragraphs
    passport = next(p for p in body if "ПАСПОРТ" in p.text)
    assert passport._p.find(f"{qn('w:r')}/{qn('w:br')}") is not None  # the break inside it
    assert len(body) == 2  # no paragraph of its own for the break

    empty = docx.Document()
    empty.sections[0]._sectPr.find(qn("w:cols")).set(qn("w:num"), "2")
    empty.add_paragraph("Гарантийный талон")
    empty.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)
    empty.add_paragraph("_smart light.")
    assert any(b.column_break_before for b in _document(empty).blocks[1:2])


def test_a_clearing_line_break_keeps_the_text_below_a_floating_table(
    profiles: YamlProfileRepository,
) -> None:
    """f8e56995: <w:br w:clear="all"/> after a floating table; a plain break let the title and
    the text flow up beside the table."""
    source = docx.Document()
    floating = source.add_table(rows=1, cols=1)
    floating._tbl.tblPr.append(
        docx.oxml.parse_xml(
            f'<w:tblpPr {_W_DECL} w:vertAnchor="text" w:horzAnchor="margin" w:tblpXSpec="right"/>'
        )
    )
    floating.cell(0, 0).text = "Директору школы"
    source.add_paragraph().add_run()._r.append(
        docx.oxml.parse_xml(f'<w:br {_W_DECL} w:type="textWrapping" w:clear="all"/>')
    )
    title = source.add_paragraph("ЗАЯВЛЕНИЕ")
    title.add_run()._r.append(docx.oxml.parse_xml(f'<w:br {_W_DECL} w:clear="left"/>'))
    title.add_run("Прошу принять")

    rendered = DocxRenderer().render(
        _document(source), profiles.get("generic"), RenderMode.FAITHFUL
    )

    xml = _document_xml(rendered)
    assert b'w:clear="all"' in xml  # the empty line's break
    assert b'w:clear="left"' in xml  # the break inside the text


def test_breaks_written_outside_any_run_still_break_the_line() -> None:
    """58a67267: a web export writes '</w:r><w:br/><w:br/><w:r>' in the paragraph itself; Word
    shows the empty lines, python-docx sees no run there."""
    source = docx.Document()
    paragraph = source.add_paragraph("Соревнования проходили в три этапа.")
    for _ in range(2):
        paragraph._p.append(docx.oxml.parse_xml(f"<w:br {_W_DECL}/>"))
    paragraph.add_run("Команду представляли 7 человек.")

    block = _document(source).blocks[0]

    assert isinstance(block, ParagraphBlock)
    assert "".join(r.text for r in block.runs) == (
        "Соревнования проходили в три этапа.\n\nКоманду представляли 7 человек."
    )


def test_a_hidden_mark_after_a_visible_break_leaves_one_line_between_tables(
    profiles: YamlProfileRepository,
) -> None:
    """c9ad5c61: between two tables a paragraph whose mark is hidden holds a visible break;
    Word shows the line the break ends, so the gap is a full empty line, not a 1-pt hairline."""
    source = docx.Document()
    source.add_table(rows=1, cols=1).cell(0, 0).text = "Довольны совместной работой:"
    between = source.add_paragraph()
    between._p.get_or_add_pPr().append(docx.oxml.parse_xml(f"<w:rPr {_W_DECL}><w:vanish/></w:rPr>"))
    between.add_run().add_break()
    source.add_table(rows=1, cols=1).cell(0, 0).text = "Характеристика"

    rendered = DocxRenderer().render(
        _document(source), profiles.get("generic"), RenderMode.FAITHFUL
    )

    between_tables = re.findall(rb"</w:tbl>(.*?)<w:tbl>", _document_xml(rendered), re.S)
    assert len(between_tables) == 1
    assert b'w:lineRule="exact"' not in between_tables[0]


def _table_sizes(*, styled: bool, normal_pt: int) -> list[float | None]:
    source = docx.Document()
    defaults = source.styles.element.find(qn("w:docDefaults"))
    rpr = defaults.find(f"{qn('w:rPrDefault')}/{qn('w:rPr')}")
    for old in rpr.findall(qn("w:sz")):
        rpr.remove(old)
    rpr.append(docx.oxml.parse_xml(f'<w:sz {_W_DECL} w:val="22"/>'))
    compat = source.settings.element.find(qn("w:compat"))
    if compat is not None:
        source.settings.element.remove(compat)
    source.styles["Normal"].font.size = docx.shared.Pt(normal_pt)
    table = source.add_table(rows=1, cols=1)
    if styled:
        table.style = "Table Grid"
    table.cell(0, 0).text = "Наименование организации"
    blocks = [b for b in _document(source).blocks if isinstance(b, TableBlock)]
    return [r.size_pt for r in blocks[0].rows[0].cells[0].runs]


def test_a_table_style_without_a_size_lets_the_defaults_win_over_a_12_pt_normal() -> None:
    """ed34515b: Normal 12 pt, document defaults 11 pt, 'Table Grid' with no size: Word sets
    the cells in 11 pt. Not in a table without a style, nor with Normal at 14 pt."""
    assert _table_sizes(styled=True, normal_pt=12) == [11.0]
    assert _table_sizes(styled=False, normal_pt=12) == [12.0]
    assert _table_sizes(styled=True, normal_pt=14) == [14.0]


def test_raised_and_sunken_cell_lines_keep_the_automatic_colour(
    profiles: YamlProfileRepository,
) -> None:
    """6717f1b5: Word draws outset lines of colour 'auto' in light and dark greys; in black
    they are thick bars."""
    source = docx.Document()
    table = source.add_table(rows=1, cols=2)
    for cell in table.rows[0].cells:
        cell._tc.get_or_add_tcPr().append(
            docx.oxml.parse_xml(
                f'<w:tcBorders {_W_DECL}><w:top w:val="outset" w:sz="6" w:color="auto"/>'
                '<w:bottom w:val="single" w:sz="4" w:color="auto"/></w:tcBorders>'
            )
        )
        cell.text = "нет"

    rendered = DocxRenderer().render(
        _document(source), profiles.get("generic"), RenderMode.FAITHFUL
    )

    xml = _document_xml(rendered)
    assert b'w:val="outset" w:sz="6" w:color="auto"' in xml
    assert b'w:val="single" w:sz="4" w:color="000000"' in xml


def test_keep_with_next_and_keep_lines_hold_the_page_breaks(
    profiles: YamlProfileRepository,
) -> None:
    source = docx.Document()
    heading = source.add_paragraph("3. Уточненные сведения о заявителе")
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.keep_together = True
    source.add_paragraph("3.1. Фамилия, имя, отчество")

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    fmt = rendered.paragraphs[0].paragraph_format
    assert fmt.keep_with_next
    assert fmt.keep_together
    assert not rendered.paragraphs[1].paragraph_format.keep_with_next


def test_a_theme_font_on_the_run_wins_over_the_style_typeface() -> None:
    source = docx.Document()
    plain = source.styles.add_style("Plain Text Test", 1)
    plain.font.name = "Consolas"
    paragraph = source.add_paragraph(style="Plain Text Test")
    run = paragraph.add_run("Основы толкования")
    run._r.get_or_add_rPr().insert(
        0,
        docx.oxml.parse_xml(
            f'<w:rFonts {_W_DECL} w:asciiTheme="majorHAnsi" w:hAnsiTheme="majorHAnsi"/>'
        ),
    )

    block = _import(source)[0]

    assert isinstance(block, ParagraphBlock)
    assert block.runs[0].font not in (None, "Consolas")  # the theme's heading typeface


def test_cyrillic_text_takes_the_high_ansi_typeface() -> None:
    source = docx.Document()
    style = source.styles.add_style("Заголовок сообщения (текст)", 2)  # a character style
    style.element.get_or_add_rPr().insert(
        0, docx.oxml.parse_xml(f'<w:rFonts {_W_DECL} w:ascii="Arial Black" w:hAnsi="Arial Black"/>')
    )
    paragraph = source.add_paragraph()
    for text in ("Кому:", "To:"):
        run = paragraph.add_run(text, style=style)
        run.bold = text == "To:"  # two runs, not one
        run._r.get_or_add_rPr().insert(
            0, docx.oxml.parse_xml(f'<w:rFonts {_W_DECL} w:hAnsi="Times New Roman"/>')
        )

    block = _import(source)[0]

    assert isinstance(block, ParagraphBlock)
    assert [r.font for r in block.runs] == ["Times New Roman", "Arial Black"]


def test_a_text_box_keeps_its_turn_fill_and_centred_text(profiles: YamlProfileRepository) -> None:
    from docx.oxml import parse_xml

    xml = _text_box_run("ОПРОСНЫЙ ЛИСТ").xml
    xml = xml.replace("<a:xfrm>", '<a:xfrm rot="16200000">').replace(
        "<a:ln w=", '<a:solidFill><a:srgbClr val="50C2BD"/></a:solidFill><a:ln w='
    )
    xml = xml.replace("<wps:bodyPr/>", '<wps:bodyPr anchor="ctr"/>')
    source = docx.Document()
    source.add_paragraph("Ливневые очистные сооружения")._p.append(parse_xml(xml))

    document = _document(source)
    rendered = DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL)

    box = next(b for b in document.blocks if b.kind == "text_box")
    assert (box.rotation_deg, box.fill, box.text_anchor) == (270, "50C2BD", "ctr")
    xml_out = _document_xml(rendered)
    assert b'rot="16200000"' in xml_out
    assert b'<a:srgbClr val="50C2BD"/>' in xml_out
    assert b'anchor="ctr"' in xml_out


def test_theme_colours_paint_lines_and_text_as_word_does() -> None:
    source = docx.Document()
    table = source.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Общая площадь"
    table.cell(0, 0)._tc.get_or_add_tcPr().append(
        docx.oxml.parse_xml(
            f'<w:tcBorders {_W_DECL}><w:bottom w:val="single" w:sz="4" w:color="000000" '
            'w:themeColor="text1" w:themeTint="80"/></w:tcBorders>'
        )
    )
    note = source.add_paragraph()
    run = note.add_run("Примечания")
    run._r.get_or_add_rPr().append(
        docx.oxml.parse_xml(f'<w:color {_W_DECL} w:val="000000" w:themeColor="accent1"/>')
    )

    blocks = _import(source)

    (table_block,) = [b for b in blocks if isinstance(b, TableBlock)]
    borders = table_block.rows[0].cells[0].borders
    assert borders is not None
    bottom = dict(borders.lines)["bottom"].color
    assert bottom is not None
    assert bottom != "000000"  # text1 tinted half way to white: a grey line
    assert int(bottom[:2], 16) > 100
    paragraph = blocks[-1]
    assert isinstance(paragraph, ParagraphBlock)
    assert paragraph.runs[0].color not in (None, "000000")  # the theme's accent, not black


def test_a_heading_keeps_the_colour_of_its_paragraph_style() -> None:
    source = docx.Document()
    source.add_heading("1. Цели создания", 1)  # Heading 1: the theme's accent, darkened

    (heading,) = _import(source)

    colour = _textual(heading).runs[0].color
    assert colour is not None  # Word's own value for it: 365F91
    assert [int(colour[i : i + 2], 16) for i in (0, 2, 4)] == pytest.approx([54, 95, 145], abs=2)


def test_a_row_that_skips_grid_columns_keeps_its_cells_in_place(
    profiles: YamlProfileRepository,
) -> None:
    source = docx.Document()
    table = source.add_table(rows=1, cols=3)
    tr = table.rows[0]._tr
    first = tr.tc_lst[0]
    tr.remove(first)  # the row starts at the second grid column
    tr.get_or_add_trPr().append(docx.oxml.parse_xml(f'<w:gridBefore {_W_DECL} w:val="1"/>'))
    tr.tc_lst[0].append(docx.oxml.parse_xml(f"<w:p {_W_DECL}><w:r><w:t>ОГРН</w:t></w:r></w:p>"))

    (block,) = [b for b in _import(source) if isinstance(b, TableBlock)]
    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    assert block.rows[0].skip == 1
    assert [p.col for p in block.placements()] == [1, 2]
    row = rendered.tables[0].rows[0]._tr
    assert _required(row.trPr).find(qn("w:gridBefore")).get(qn("w:val")) == "1"
    assert "ОГРН" in "".join(row.tc_lst[0].itertext())


def test_framed_paragraphs_and_old_style_lines_keep_their_places(
    profiles: YamlProfileRepository,
) -> None:
    """A payment order as 1C writes it: each field a paragraph in a frame at a fixed place
    (w:framePr), the lines to fill drawn as old-style (VML) lines on the page."""
    from docx.oxml import parse_xml

    from foto2word.domain.document import ShapeBlock

    source = docx.Document()
    for text, frame in (
        (
            "Поступ. в банк плат.",
            'w:w="2300" w:wrap="auto" w:hAnchor="margin" w:x="3781" w:y="1261"',
        ),
        ("", 'w:w="420" w:h="220" w:wrap="auto" w:hAnchor="margin" w:x="10681" w:y="2101"'),
    ):
        paragraph = source.add_paragraph(text)
        paragraph._p.get_or_add_pPr().insert(0, parse_xml(f"<w:framePr {_W_DECL} {frame}/>"))
        if not text:  # the box for the payment kind: an empty frame with a line around it
            paragraph._p.get_or_add_pPr().append(
                parse_xml(
                    f'<w:pBdr {_W_DECL}><w:top w:val="single" w:sz="4" w:space="1"/>'
                    '<w:bottom w:val="single" w:sz="4" w:space="1"/></w:pBdr>'
                )
            )
    holder = source.paragraphs[0]._p
    holder.append(
        parse_xml(
            f'<w:r {_W_DECL}><w:pict><v:line xmlns:v="urn:schemas-microsoft-com:vml" '
            'style="position:absolute;mso-position-horizontal-relative:page;'
            'mso-position-vertical-relative:page" from="59pt,76pt" to="173pt,76pt"/>'
            "</w:pict></w:r>"
        )
    )
    source.add_paragraph("Текст после рамок")

    importer = DocxImporter()
    document = importer.convert(docx.Document(BytesIO(_bytes(source))))
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    frames = [p._p.find(f"{qn('w:pPr')}/{qn('w:framePr')}") for p in rendered.paragraphs]
    framed = [
        (p.text, f.get(qn("w:x")), f.get(qn("w:y")))
        for p, f in zip(rendered.paragraphs, frames, strict=True)
        if f is not None
    ]
    assert framed == [("Поступ. в банк плат.", "3781", "1261"), ("", "10681", "2101")]
    (line,) = [b for b in document.blocks if isinstance(b, ShapeBlock)]
    assert line.geometry == "line"
    assert line.floating.horizontal_from == line.floating.vertical_from == "page"
    assert (line.floating.x_mm, line.floating.y_mm) == pytest.approx((20.81, 26.81), abs=0.05)
    assert (line.width_mm, line.height_mm) == pytest.approx((40.22, 0.0), abs=0.05)
    assert "Текст после рамок" in [p.text for p in rendered.paragraphs]


def test_a_frame_keeps_what_its_style_says_and_the_paragraph_leaves_out(
    profiles: YamlProfileRepository,
) -> None:
    from docx.oxml import parse_xml

    source = docx.Document()
    note = source.styles.add_style("Заметка", 1)
    note.element.get_or_add_pPr().insert(
        0, parse_xml(f'<w:framePr {_W_DECL} w:wrap="around" w:vAnchor="text" w:hAnchor="text"/>')
    )
    paragraph = source.add_paragraph("Дата заполнения", style="Заметка")
    paragraph._p.get_or_add_pPr().append(
        parse_xml(f'<w:framePr {_W_DECL} w:w="9166" w:hAnchor="page" w:x="1591" w:y="404"/>')
    )

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    frame = _required(
        next(p for p in rendered.paragraphs if p.text == "Дата заполнения")._p.pPr
    ).find(qn("w:framePr"))
    assert frame.get(qn("w:vAnchor")) == "text"  # the style's, the paragraph did not say
    assert frame.get(qn("w:hAnchor")) == "page"  # the paragraph's own wins


def test_an_old_style_line_inside_a_group_is_not_read_as_a_page_line() -> None:
    """In a VML group the ends are in the group's own units, not points: the group is kept
    whole, as written."""
    from docx.oxml import parse_xml

    from foto2word.domain.document import ShapeBlock

    source = docx.Document()
    source.add_paragraph("Дата")._p.append(
        parse_xml(
            f'<w:r {_W_DECL}><w:pict><v:group xmlns:v="urn:schemas-microsoft-com:vml" '
            'style="position:absolute;width:55pt;height:1pt" coordsize="1104,10">'
            '<v:line style="position:absolute" from="0,5" to="1104,5"/></v:group></w:pict></w:r>'
        )
    )

    blocks = _import(source)

    (group,) = [b for b in blocks if isinstance(b, ShapeBlock)]
    assert group.geometry == "custom"
    assert group.vml is not None
    assert "<v:group" in group.vml


def test_an_old_style_freeform_line_is_kept_as_word_wrote_it(
    profiles: YamlProfileRepository,
) -> None:
    """A PDF converter's rule to sign on: a VML shape with a path, wrapping top and bottom. Word
    sends the text below a thin VML line otherwise than below the same line in DrawingML, so the
    shape goes back as it was."""
    from docx.oxml import parse_xml

    from foto2word.domain.document import ShapeBlock

    source = docx.Document()
    source.add_paragraph("Руководителю")
    source.add_paragraph()._p.append(
        parse_xml(
            f'<w:r {_W_DECL}><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml" '
            'xmlns:w10="urn:schemas-microsoft-com:office:word" '
            'style="position:absolute;margin-left:399pt;margin-top:13.5pt;width:168pt;'
            'height:.1pt;z-index:-1;mso-position-horizontal-relative:page" '
            'coordsize="2133600,1270" path="m,l2133600,e" filled="f" strokeweight=".5pt">'
            '<w10:wrap type="topAndBottom"/></v:shape></w:pict></w:r>'
        )
    )
    source.add_paragraph("Заявление")

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    (line,) = [b for b in document.blocks if isinstance(b, ShapeBlock)]
    assert line.width_mm == pytest.approx(59.27, abs=0.01)
    assert line.floating.wrap == "top_and_bottom"
    (shape,) = rendered.element.body.iter("{urn:schemas-microsoft-com:vml}shape")
    assert shape.get("path") == "m,l2133600,e"


def test_old_style_word_art_in_the_line_is_kept(profiles: YamlProfileRepository) -> None:
    """A schedule's title in WordArt (a VML text path standing in the line): the paragraph is
    not an empty line, and the WordArt goes back as Word wrote it."""
    from docx.oxml import parse_xml

    source = docx.Document()
    source.add_paragraph()._p.append(
        parse_xml(
            f'<w:r {_W_DECL}><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml" '
            'type="#_x0000_t136" style="width:464pt;height:28pt" fillcolor="silver">'
            '<v:textpath string="Расписание занятий"/></v:shape></w:pict></w:r>'
        )
    )
    source.add_paragraph("Первое полугодие")

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    title = document.blocks[0]
    assert title.text() == "Расписание занятий"
    (path,) = rendered.element.body.iter("{urn:schemas-microsoft-com:vml}textpath")
    assert path.get("string") == "Расписание занятий"


def test_an_old_style_picture_is_not_kept_as_a_drawing() -> None:
    """A VML watermark is a picture (v:imagedata): its image is read as one, not copied as a
    drawing whose link to the image would be lost."""
    from docx.oxml import parse_xml

    from foto2word.domain.document import ShapeBlock

    source = docx.Document()
    source.add_paragraph("Текст")._p.append(
        parse_xml(
            f'<w:r {_W_DECL}><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            'type="#_x0000_t75" style="position:absolute;width:300pt;height:50pt">'
            '<v:imagedata r:id="rIdWm"/></v:shape></w:pict></w:r>'
        )
    )

    assert not [b for b in _import(source) if isinstance(b, ShapeBlock)]


def test_hidden_text_and_pictures_are_left_out_as_word_prints_them() -> None:
    from docx.enum.style import WD_STYLE_TYPE
    from PIL import Image

    from foto2word.domain.document import ImageBlock

    png = BytesIO()
    Image.new("RGB", (40, 40), "blue").save(png, format="PNG")
    source = docx.Document()
    line = source.add_paragraph("Мужчины и женщины!")
    hidden = line.add_run()
    hidden.add_picture(BytesIO(png.getvalue()), width=Mm(150))
    hidden.font.hidden = True
    note = source.styles.add_style("Скрытое", WD_STYLE_TYPE.CHARACTER)
    note.font.hidden = True
    line.add_run(" черновик", style="Скрытое")
    shown = line.add_run(" !")
    shown.font.hidden = False

    blocks = _import(source)

    assert not [b for b in blocks if isinstance(b, ImageBlock)]
    assert blocks[0].text() == "Мужчины и женщины! !"


def test_the_page_colour_is_kept(profiles: YamlProfileRepository) -> None:
    from docx.oxml import parse_xml

    source = docx.Document()
    source.element.insert(
        0,
        parse_xml(
            f'<w:background {_W_DECL} w:color="DEEAF6" w:themeColor="accent1" w:themeTint="33"/>'
        ),
    )
    source.settings.element.insert(0, parse_xml(f"<w:displayBackgroundShape {_W_DECL}/>"))
    source.add_paragraph("Плакат")

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    assert document.page_background is not None  # the theme's accent, nearly white
    background = rendered.element.find(qn("w:background"))
    assert background is not None
    assert background.get(qn("w:color")) == document.page_background
    assert rendered.settings.element.find(qn("w:displayBackgroundShape")) is not None


def test_an_empty_line_with_a_hidden_mark_is_gone_unless_it_parts_two_tables() -> None:
    from docx.oxml import parse_xml

    def hide_mark(paragraph: Any) -> None:
        paragraph._p.get_or_add_pPr().append(parse_xml(f"<w:rPr {_W_DECL}><w:vanish/></w:rPr>"))

    source = docx.Document()
    source.add_paragraph("Сведения")
    hide_mark(source.add_paragraph())
    source.add_table(rows=1, cols=1).cell(0, 0).text = "Первая"
    hide_mark(source.add_paragraph())  # between two tables: Word keeps it, or they would join
    source.add_table(rows=1, cols=1).cell(0, 0).text = "Вторая"

    blocks = _import(source)

    kinds = [type(b).__name__ for b in blocks]
    assert kinds[:2] == ["ParagraphBlock", "TableBlock"]  # no empty line before the table
    assert kinds.count("TableBlock") == 2
    assert kinds.index("TableBlock", 2) > 2  # something stands between the tables


def test_a_text_box_hangs_after_the_empty_lines_above_its_paragraph() -> None:
    """A flowchart drawn over empty lines: each box hangs from its own line, so the empty lines
    before it stay before it, or every box would climb to the top."""
    from foto2word.domain.document import TextBoxBlock

    source = docx.Document()
    source.add_paragraph("БЛОК-СХЕМА")
    source.add_paragraph()
    source.add_paragraph()
    source.add_paragraph()._p.append(_text_box_run("Подача заявления"))

    kinds = [type(b) for b in _import(source)]

    assert kinds.index(SpacerBlock) < kinds.index(TextBoxBlock)


def test_an_arrow_keeps_its_head_and_colour(profiles: YamlProfileRepository) -> None:
    """A flowchart's connector: a blue line with a triangle at its end."""
    source = docx.Document()
    source.add_paragraph("Схема")._p.append(
        _shape_run(
            "straightConnector1",
            0,
            720000,
            line='<a:solidFill><a:srgbClr val="4F81BD"/></a:solidFill><a:tailEnd type="triangle"/>',
        )
    )

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    line = next(rendered.element.body.iter(f"{{{_DML}}}ln"))
    assert line.find(f"{{{_DML}}}tailEnd").get("type") == "triangle"
    assert line.find(f"{{{_DML}}}headEnd") is None
    assert next(line.iter(f"{{{_DML}}}srgbClr")).get("val") == "4F81BD"


def test_a_page_number_in_a_frame_stands_where_the_frame_does() -> None:
    from docx.oxml import parse_xml

    source = docx.Document()
    number = source.sections[0].header.paragraphs[0]
    number._p.get_or_add_pPr().append(
        parse_xml(
            f'<w:framePr {_W_DECL} w:wrap="around" w:vAnchor="text" w:xAlign="right" w:y="1"/>'
        )
    )
    number._p.append(
        parse_xml(f'<w:fldSimple {_W_DECL} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>')
    )
    source.add_paragraph("Информационное сообщение")

    document = _document(source)

    assert document.page_numbers is not None
    assert document.page_numbers.align is Alignment.RIGHT


def test_an_empty_paragraph_with_a_line_break_is_two_lines_high() -> None:
    """As HTML pages turned into Word files write their blank lines: <p><br></p>."""
    source = docx.Document()
    source.add_paragraph("Классный час")
    source.add_paragraph()
    source.add_paragraph("Подготовила учитель")
    source.add_paragraph().add_run().add_break()
    source.add_paragraph("2019 г")

    gaps = [b.height_pt for b in _import(source) if isinstance(b, SpacerBlock)]

    assert len(gaps) == 2
    assert 11 < _required(gaps[1]) - _required(gaps[0]) < 18  # one more line of 11 pt type


def test_a_group_is_taken_apart_into_its_pieces_at_their_places() -> None:
    """A logo and a caption grouped (wpg): the group 100 x 50 mm at 20 mm / 30 mm on the page,
    its own units twice as fine; the picture its left half, the caption its right half."""
    from docx.oxml import parse_xml
    from PIL import Image

    from foto2word.domain.document import ImageBlock, TextBoxBlock

    png = BytesIO()
    Image.new("RGB", (40, 40), "blue").save(png, format="PNG")
    source = docx.Document()
    holder = source.add_paragraph()
    holder.add_run().add_picture(BytesIO(png.getvalue()), width=Mm(10))
    rel_id = next(holder._p.iter(f"{{{_DML}}}blip")).get(qn("r:embed"))
    for drawing in holder._p.findall(f".//{qn('w:drawing')}"):
        drawing.getparent().remove(drawing)
    mm = 36000
    holder._p.append(
        parse_xml(
            f'<w:r xmlns:w="{_W}"><w:drawing>'
            f'<wp:anchor xmlns:wp="{_WP}" distT="0" distB="0" distL="0" distR="0" simplePos="0" '
            'relativeHeight="3" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
            '<wp:simplePos x="0" y="0"/>'
            '<wp:positionH relativeFrom="page">'
            f"<wp:posOffset>{20 * mm}</wp:posOffset></wp:positionH>"
            '<wp:positionV relativeFrom="page">'
            f"<wp:posOffset>{30 * mm}</wp:posOffset></wp:positionV>"
            f'<wp:extent cx="{100 * mm}" cy="{50 * mm}"/><wp:wrapNone/><wp:docPr id="7" name="G"/>'
            f'<a:graphic xmlns:a="{_DML}"><a:graphicData '
            'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup">'
            "<wpg:wgp "
            'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
            f'xmlns:wps="{_WPS}" '
            'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            f'<wpg:cNvGrpSpPr/><wpg:grpSpPr><a:xfrm><a:off x="0" y="0"/>'
            f'<a:ext cx="{100 * mm}" cy="{50 * mm}"/><a:chOff x="0" y="0"/>'
            f'<a:chExt cx="{200 * mm}" cy="{100 * mm}"/></a:xfrm></wpg:grpSpPr>'
            '<pic:pic><pic:nvPicPr><pic:cNvPr id="1" name="logo"/><pic:cNvPicPr/></pic:nvPicPr>'
            f'<pic:blipFill><a:blip r:embed="{rel_id}"/><a:stretch><a:fillRect/></a:stretch>'
            f'</pic:blipFill><pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{100 * mm}" '
            f'cy="{100 * mm}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
            "</pic:pic>"
            f'<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="{100 * mm}" y="0"/>'
            f'<a:ext cx="{100 * mm}" cy="{100 * mm}"/></a:xfrm><a:prstGeom prst="rect">'
            f"<a:avLst/></a:prstGeom></wps:spPr><wps:txbx><w:txbxContent><w:p><w:r><w:t>"
            "Подпись</w:t></w:r></w:p></w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp>"
            "</wpg:wgp></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>"
        )
    )

    blocks = _import(source)

    (picture,) = [b for b in blocks if isinstance(b, ImageBlock)]
    (caption,) = [b for b in blocks if isinstance(b, TextBoxBlock)]
    assert (picture.width_mm, picture.height_mm) == pytest.approx((50, 50), abs=0.1)
    assert picture.floating is not None
    assert (picture.floating.x_mm, picture.floating.y_mm) == pytest.approx((20, 30), abs=0.1)
    assert (caption.width_mm, caption.height_mm) == pytest.approx((50, 50), abs=0.1)
    assert (caption.floating.x_mm, caption.floating.y_mm) == pytest.approx((70, 30), abs=0.1)


_OMML = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def test_an_equation_keeps_its_place_in_the_line(profiles: YamlProfileRepository) -> None:
    """Ом·мм²/м written as Word's own equation (OMML): kept as an equation where it stands."""
    from docx.oxml import parse_xml

    source = docx.Document()
    line = source.add_paragraph("Удельное сопротивление вольфрама 0,055 ")
    line._p.append(
        parse_xml(
            f'<m:oMath xmlns:m="{_OMML}"><m:f><m:num><m:r><m:t>Ом·мм²</m:t></m:r></m:num>'
            "<m:den><m:r><m:t>м</m:t></m:r></m:den></m:f></m:oMath>"
        )
    )
    line.add_run(".")

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    paragraph = next(p for p in rendered.paragraphs if "вольфрама" in p.text)
    kinds = [c.tag.split("}")[1] for c in paragraph._p if c.tag != qn("w:pPr")]
    assert kinds == ["r", "oMath", "r"]
    assert "".join(t.text for t in paragraph._p.iter(f"{{{_OMML}}}t")) == "Ом·мм²м"


def test_a_filled_shape_of_any_outline_keeps_its_fill_and_outline(
    profiles: YamlProfileRepository,
) -> None:
    """A poster's panels: a navy banner drawn as a free outline (custGeom), filled, no line."""
    path = (
        '<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/>'
        '<a:pathLst><a:path w="100" h="50"><a:moveTo><a:pt x="0" y="0"/></a:moveTo>'
        '<a:lnTo><a:pt x="100" y="0"/></a:lnTo><a:lnTo><a:pt x="100" y="50"/></a:lnTo>'
        "<a:close/></a:path></a:pathLst></a:custGeom>"
    )
    source = docx.Document()
    run = _shape_run("rect", 3600000, 1800000, line="<a:noFill/>")
    sp_pr = next(run.iter(f"{{{_WPS}}}spPr"))
    sp_pr.replace(
        next(sp_pr.iter(f"{{{_DML}}}prstGeom")),
        docx.oxml.parse_xml(path.replace("<a:custGeom>", f'<a:custGeom xmlns:a="{_DML}">', 1)),
    )
    sp_pr.replace(
        next(sp_pr.iter(f"{{{_DML}}}noFill")),
        docx.oxml.parse_xml(
            f'<a:solidFill xmlns:a="{_DML}"><a:srgbClr val="122E52"/></a:solidFill>'
        ),
    )
    source.add_paragraph("Плакат")._p.append(run)

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    shape = next(rendered.element.body.iter(f"{{{_WPS}}}spPr"))
    assert shape.find(f"{{{_DML}}}custGeom") is not None
    assert len(list(shape.iter(f"{{{_DML}}}lnTo"))) == 2
    assert shape.find(f"{{{_DML}}}solidFill/{{{_DML}}}srgbClr").get("val") == "122E52"
    assert shape.find(f"{{{_DML}}}ln/{{{_DML}}}noFill") is not None


def test_a_bold_place_and_date_line_keeps_its_bold() -> None:
    source = docx.Document()
    line = source.add_paragraph()
    line.add_run("г. Малоярославец\t\t\t\t28.03.2019 года").bold = True

    (block,) = _import(source)

    assert isinstance(block, ParagraphBlock)  # its tabs as they are: no place for bold in
    assert block.runs  # a two-column line
    assert all(run.bold for run in block.runs)


def test_auto_spacing_is_not_put_between_items_of_a_list() -> None:
    """A list from a web page: its items have 'auto' spacing, which a browser (and Word) puts
    around the list, not between its items."""
    from docx.oxml import parse_xml

    source = docx.Document()
    source.add_paragraph("Для проведения платежа:")
    items = [
        source.add_paragraph(text, style="List Bullet") for text in ("Минск", "БГМУ", "Оплата")
    ]
    for item in items:
        item._p.get_or_add_pPr().append(
            parse_xml(
                f'<w:spacing {_W_DECL} w:before="100" w:beforeAutospacing="1" w:after="100" '
                'w:afterAutospacing="1"/>'
            )
        )

    layouts = [_layout(b) for b in _import(source)][1:]

    assert [(_required(x).space_before_pt, _required(x).space_after_pt) for x in layouts] == [
        (14.0, 0.0),
        (0.0, 0.0),
        (0.0, 14.0),
    ]


def test_a_page_number_in_a_frame_keeps_its_frame(profiles: YamlProfileRepository) -> None:
    from docx.oxml import parse_xml

    source = docx.Document()
    number = source.sections[0].header.paragraphs[0]
    number._p.get_or_add_pPr().append(
        parse_xml(
            f'<w:framePr {_W_DECL} w:wrap="around" w:vAnchor="text" w:xAlign="right" w:y="1"/>'
        )
    )
    number._p.append(
        parse_xml(f'<w:fldSimple {_W_DECL} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>')
    )
    source.add_paragraph("Протокол")

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    header = rendered.sections[0].header.paragraphs[0]
    frame = header._p.find(f"{qn('w:pPr')}/{qn('w:framePr')}")
    assert frame is not None
    assert frame.get(qn("w:xAlign")) == "right"


def test_a_picture_before_a_page_break_stays_on_its_page() -> None:
    """A sample of a certificate: the picture, then a page break and the next caption, all in
    one paragraph. The picture belongs to the page before the break."""
    from docx.enum.text import WD_BREAK
    from PIL import Image

    from foto2word.domain.document import ImageBlock

    png = BytesIO()
    Image.new("RGB", (40, 40), "green").save(png, format="PNG")
    source = docx.Document()
    source.add_paragraph("Страница 1")
    holder = source.add_paragraph()
    holder.add_run().add_picture(BytesIO(png.getvalue()), width=Mm(100))
    _anchored(next(holder._p.iter(qn("wp:inline"))))
    holder.add_run().add_break(WD_BREAK.PAGE)
    holder.add_run("Страница 2")

    blocks = _import(source)

    picture = next(b for b in blocks if isinstance(b, ImageBlock))
    caption = next(b for b in blocks if b.text() == "Страница 2")
    assert not picture.page_break_before
    assert caption.page_break_before


def test_a_file_without_paragraph_defaults_gets_words_own() -> None:
    """Measured in Word (PHPWord files, no pPrDefault, an empty Normal): Word gives such
    paragraphs its own defaults, 8 pt after and lines of 1.08, whatever the file's Word version."""
    source = docx.Document()
    defaults = source.styles.element.find(qn("w:docDefaults"))
    defaults.remove(defaults.find(qn("w:pPrDefault")))
    normal = source.styles["Normal"].element
    for spacing in normal.iter(qn("w:spacing")):
        spacing.getparent().remove(spacing)
    source.add_paragraph("Заявление")

    (block,) = _import(source)

    assert _layout(block).space_after_pt == pytest.approx(8.0)
    assert _layout(block).line_spacing == pytest.approx(259 / 240, abs=0.01)


def test_a_raw_line_feed_inside_text_is_not_a_line_break() -> None:
    """PHPWord writes text with its line feeds inside w:t; Word does not break the line there
    (only w:br does), so no empty line appears under the paragraph."""
    source = docx.Document()
    paragraph = source.add_paragraph()
    run = paragraph.add_run()
    text = run._r.makeelement(qn("w:t"), {})
    text.text = "Питание от литиевой батареи 6 V (тип 123).\n"
    run._r.append(text)

    (block,) = _import(source)

    assert "\n" not in block.text()


def test_kerning_of_the_style_is_kept(profiles: YamlProfileRepository) -> None:
    """Normal kerns pairs of letters from 1 pt up (w:kern 2): the text is a little narrower and
    breaks its lines where the original does."""
    from docx.oxml import parse_xml

    source = docx.Document()
    source.styles["Normal"].element.get_or_add_rPr().append(
        parse_xml(f'<w:kern {_W_DECL} w:val="2"/>')
    )
    source.add_paragraph("Проректор по учебно-методической работе")

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    run = next(r for p in rendered.paragraphs for r in p.runs if r.text.startswith("Проректор"))
    kern = run._r.find(f"{qn('w:rPr')}/{qn('w:kern')}")
    assert kern is not None
    assert kern.get(qn("w:val")) == "2"


def test_letters_squeezed_in_width_stay_so(profiles: YamlProfileRepository) -> None:
    """A protocol's list set at 95% width (w:w): its lines fit where the same text at 100%
    wraps."""
    from docx.oxml import parse_xml

    source = docx.Document()
    run = source.add_paragraph().add_run("1.Виткина Н.В., начальник управления образования")
    run._r.get_or_add_rPr().append(parse_xml(f'<w:w {_W_DECL} w:val="95"/>'))
    run.font.size = Pt(14)  # w:sz goes after w:w, as Word orders them

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    (block,) = document.blocks
    assert [r.scale_pct for r in _textual(block).runs] == [95]
    out = next(r for p in rendered.paragraphs for r in p.runs if r.text.startswith("1.Виткина"))
    scale = out._r.find(f"{qn('w:rPr')}/{qn('w:w')}")
    assert scale is not None
    assert scale.get(qn("w:val")) == "95"


def test_the_grid_of_lines_is_kept(profiles: YamlProfileRepository) -> None:
    """A Japanese questionnaire lays its lines on Word's grid (w:docGrid of type lines)."""
    from docx.oxml import parse_xml

    source = docx.Document()
    sect_pr = source.sections[0]._sectPr
    for old in sect_pr.findall(qn("w:docGrid")):
        sect_pr.remove(old)
    sect_pr.append(parse_xml(f'<w:docGrid {_W_DECL} w:type="linesAndChars" w:linePitch="323"/>'))
    source.add_paragraph("Q1 Укажите свой пол.")

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    grid = rendered.sections[0]._sectPr.find(qn("w:docGrid"))
    assert grid.get(qn("w:type")) == "linesAndChars"
    assert grid.get(qn("w:linePitch")) == "323"


def test_a_cell_turned_on_its_side_stays_so(profiles: YamlProfileRepository) -> None:
    """Headers of narrow columns written bottom to top (w:textDirection btLr)."""
    from docx.oxml import parse_xml

    source = docx.Document()
    cell = source.add_table(rows=1, cols=2).cell(0, 1)
    cell.text = "Площадь (кв. м)"
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:textDirection {_W_DECL} w:val="btLr"/>'))

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    turned = _required(rendered.tables[0].cell(0, 1)._tc.tcPr).find(qn("w:textDirection"))
    assert turned is not None
    assert turned.get(qn("w:val")) == "btLr"
    assert _required(rendered.tables[0].cell(0, 0)._tc.tcPr).find(qn("w:textDirection")) is None


def test_an_empty_first_line_loses_its_auto_spacing_above() -> None:
    """As with text: Word drops the auto spacing above the very first paragraph (3511c7ef)."""
    from docx.oxml import parse_xml

    source = docx.Document()
    first = source.add_paragraph()
    first._p.get_or_add_pPr().append(
        parse_xml(f'<w:spacing {_W_DECL} w:before="100" w:beforeAutospacing="1"/>')
    )
    source.add_paragraph("Ректору")

    blocks = _import(source)

    spacer = blocks[0]
    assert isinstance(spacer, SpacerBlock)
    assert spacer.lines[0].layout.space_before_pt == 0.0


def test_a_number_alone_in_a_cell_has_no_tab_after_it() -> None:
    """Word's own number takes no line for its tab; a typed tab past a narrow cell's edge does."""
    source = docx.Document()
    cell = source.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].style = source.styles["List Number"]

    block = next(b for b in _import(source) if isinstance(b, TableBlock))

    (only,) = block.rows[0].cells
    assert only.runs
    assert not only.runs[0].text.endswith(("	", " "))


def test_the_document_defaults_indent_every_paragraph() -> None:
    """A first-line indent set only in w:pPrDefault (9129e2b7): Word indents every line so."""
    from docx.oxml import parse_xml

    source = docx.Document()
    defaults = source.styles.element.find(qn("w:docDefaults"))
    holder = defaults.find(qn("w:pPrDefault"))
    if holder is None:
        holder = parse_xml(f"<w:pPrDefault {_W_DECL}><w:pPr/></w:pPrDefault>")
        defaults.append(holder)
    holder.find(qn("w:pPr")).append(parse_xml(f'<w:ind {_W_DECL} w:firstLine="709"/>'))
    source.add_paragraph("Субъект Российской Федерации;")

    (block,) = _import(source)

    assert _layout(block).first_line_indent_mm == pytest.approx(12.5, abs=0.1)


def test_the_default_fonts_are_named_not_the_themes(profiles: YamlProfileRepository) -> None:
    """A theme face inherited from the defaults would win over a run's own (t07's spaces)."""
    source = docx.Document()
    source.add_paragraph("г. Субъект" + " " * 100 + "27.10.2026 г.")

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    fonts = rendered.styles.element.find(
        f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}/{qn('w:rFonts')}"
    )
    assert fonts is not None
    assert not [k for k in fonts.attrib if k.endswith("Theme") or k.endswith("theme")]


def test_a_line_off_the_grid_stays_off_it(profiles: YamlProfileRepository) -> None:
    """A WPS file with a grid of lines whose paragraphs switch it off (w:snapToGrid 0): they
    keep their own line height (4db39a7d: lines twice as tall)."""
    from docx.oxml import parse_xml

    source = docx.Document()
    paragraph = source.add_paragraph("По проекту внесения изменений")
    paragraph._p.get_or_add_pPr().insert(0, parse_xml(f'<w:snapToGrid {_W_DECL} w:val="0"/>'))

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    snap = _required(rendered.paragraphs[0]._p.pPr).find(qn("w:snapToGrid"))
    assert snap is not None
    assert snap.get(qn("w:val")) == "0"


def test_the_document_defaults_align_every_paragraph() -> None:
    """w:jc in w:pPrDefault only (1bacda5c): Word justifies every paragraph that says nothing."""
    from docx.oxml import parse_xml

    source = docx.Document()
    defaults = source.styles.element.find(qn("w:docDefaults"))
    holder = defaults.find(qn("w:pPrDefault"))
    if holder is None:
        holder = parse_xml(f"<w:pPrDefault {_W_DECL}><w:pPr/></w:pPrDefault>")
        defaults.append(holder)
    holder.find(qn("w:pPr")).append(parse_xml(f'<w:jc {_W_DECL} w:val="both"/>'))
    source.add_paragraph("Заполняется на каждое рабочее место")

    (block,) = _import(source)

    assert _textual(block).align is Alignment.JUSTIFY


def test_a_header_starts_without_its_auto_spacing() -> None:
    """As at the top of the body: the header's first line keeps no auto spacing above it."""
    from docx.oxml import parse_xml

    source = docx.Document()
    line = source.sections[0].header.paragraphs[0]
    line.text = "Приложение №1 к"
    line._p.get_or_add_pPr().append(
        parse_xml(f'<w:spacing {_W_DECL} w:before="100" w:beforeAutospacing="1"/>')
    )
    source.add_paragraph("Текст")

    header = DocxImporter().convert(docx.Document(BytesIO(_bytes(source)))).header

    assert _layout(header[0]).space_before_pt == 0.0


def test_the_binding_room_is_kept(profiles: YamlProfileRepository) -> None:
    """w:gutter beside the left margin: the text starts further in (2c58ddf4: 9 mm)."""
    source = docx.Document()
    source.sections[0].gutter = Mm(9)
    source.add_paragraph("Список растений")

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    assert document.page is not None
    assert document.page.gutter_mm == pytest.approx(9.0, abs=0.05)
    assert _required(rendered.sections[0].gutter).mm == pytest.approx(9.0, abs=0.05)


def test_a_line_of_pictures_joins_its_neighbours_of_its_style() -> None:
    """A picture alone in a list paragraph between two list paragraphs: 'no space between
    paragraphs of the same style' holds across it (ef045c01)."""
    from PIL import Image

    from foto2word.domain.document import ImageBlock

    png = BytesIO()
    Image.new("RGB", (40, 40), "red").save(png, format="PNG")
    source = docx.Document()
    style = source.styles["List Paragraph"]
    style.paragraph_format.space_after = Pt(10)
    style.element.get_or_add_pPr().append(docx.oxml.OxmlElement("w:contextualSpacing"))
    source.add_paragraph("Организационный момент.", style=style)
    holder = source.add_paragraph(style=style)
    holder.add_run().add_picture(BytesIO(png.getvalue()), width=Mm(40))
    source.add_paragraph("Задание.", style=style)

    blocks = _import(source)

    first, picture = blocks[0], next(b for b in blocks if isinstance(b, ImageBlock))
    assert _layout(first).space_after_pt == 0.0
    assert picture.layout is not None
    assert picture.layout.space_after_pt == 0.0


def test_lines_joined_by_their_style_stay_lines() -> None:
    """Requisites lines with 'no space between paragraphs of the same style': as a block they
    would each keep their 10 pt after (f1eb3c05)."""
    source = docx.Document()
    for line in ("Предварительно утвержден Советом", "директоров ОАО", "Протокол № 3"):
        paragraph = source.add_paragraph(line)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph._p.get_or_add_pPr().append(docx.oxml.OxmlElement("w:contextualSpacing"))

    blocks = _import(source)

    assert not [b for b in blocks if isinstance(b, RequisitesBlock)]


def test_an_empty_last_section_after_columns_is_kept() -> None:
    """A list in two columns, then a continuous section with only an empty line: Word balances
    the columns before that section, so it stays (2c58ddf4)."""
    from docx.enum.section import WD_SECTION

    source = docx.Document()
    source.add_paragraph("Список растений")
    columns = source.add_section(WD_SECTION.CONTINUOUS)
    columns._sectPr.append(docx.oxml.OxmlElement("w:cols"))
    columns._sectPr.find(qn("w:cols")).set(qn("w:num"), "2")
    for name in ("Пастушья сумка", "Редька дикая"):
        source.add_paragraph(name)
    closing = source.add_section(WD_SECTION.CONTINUOUS)  # a copy of the last: one column again
    closing._sectPr.remove(closing._sectPr.find(qn("w:cols")))
    source.add_paragraph("")

    blocks = _import(source)

    last = blocks[-1]
    assert last.section is not None
    assert last.section.columns == 1


def test_room_for_a_pictures_shadow_is_kept(profiles: YamlProfileRepository) -> None:
    """A screenshot with a shadow (wp:effectExtent 12 pt above, 28 pt below): Word gives its
    line that much more height (e5eb3dd7)."""
    from docx.oxml import parse_xml
    from PIL import Image

    from foto2word.domain.document import ImageBlock

    png = BytesIO()
    Image.new("RGB", (40, 40), "blue").save(png, format="PNG")
    source = docx.Document()
    source.add_paragraph().add_run().add_picture(BytesIO(png.getvalue()), width=Mm(60))
    inline = next(source.element.body.iter(qn("wp:inline")))
    inline.find(qn("wp:extent")).addnext(
        parse_xml(
            '<wp:effectExtent xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/'
            'wordprocessingDrawing" l="0" t="152400" r="0" b="356235"/>'
        )
    )

    importer = DocxImporter()
    document = importer.convert(docx.Document(BytesIO(_bytes(source))))
    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(
                document,
                profiles.get("generic"),
                RenderMode.FAITHFUL,
                RenderExtras(images=importer.images),
            )
        )
    )

    picture = next(b for b in document.blocks if isinstance(b, ImageBlock))
    assert picture.effect_emu == (0, 152400, 0, 356235)
    effect = next(rendered.element.body.iter(qn("wp:effectExtent")))
    assert (effect.get("t"), effect.get("b")) == ("152400", "356235")


def test_an_empty_line_before_a_table_of_its_style_keeps_no_space_after() -> None:
    """'No space between paragraphs of the same style' holds from an empty line to the table's
    first line (b5f6f7e4: everything below 10 pt lower)."""
    source = docx.Document()
    source.add_paragraph("г. Махачкала")
    empty = source.add_paragraph()
    empty.paragraph_format.space_after = Pt(10)
    empty._p.get_or_add_pPr().append(docx.oxml.OxmlElement("w:contextualSpacing"))
    source.add_table(rows=1, cols=1).cell(0, 0).text = "Председательствовал:"

    blocks = _import(source)

    spacer = next(b for b in blocks if isinstance(b, SpacerBlock))
    assert spacer.lines[-1].layout.space_after_pt == 0.0


def test_a_floating_pictures_wrap_distances_are_kept(profiles: YamlProfileRepository) -> None:
    """The text keeps 9 pt from a photo box on its left and right (wp:anchor distL, distR)."""
    from PIL import Image

    png = BytesIO()
    Image.new("RGB", (40, 40), "gray").save(png, format="PNG")
    source = docx.Document()
    paragraph = source.add_paragraph()
    paragraph.add_run().add_picture(BytesIO(png.getvalue()), width=Mm(25))
    anchor = _anchored(next(paragraph._p.iter(qn("wp:inline"))))
    anchor.set("distL", "114300")
    anchor.set("distR", "114300")
    source.add_paragraph("Анкета кандидата в члены Молодежного парламента")

    importer = DocxImporter()
    document = importer.convert(docx.Document(BytesIO(_bytes(source))))
    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(
                document,
                profiles.get("generic"),
                RenderMode.FAITHFUL,
                RenderExtras(images=importer.images),
            )
        )
    )

    placed = next(rendered.element.body.iter(qn("wp:anchor")))
    assert (placed.get("distL"), placed.get("distR")) == ("114300", "114300")


def test_automatic_hyphenation_and_its_language_are_kept(profiles: YamlProfileRepository) -> None:
    """Word hyphenates the original (w:autoHyphenation) by the rules of its language (ru-RU);
    a paragraph that says no (w:suppressAutoHyphens) stays unhyphenated."""
    from docx.oxml import parse_xml

    source = docx.Document()
    source.settings.element.append(parse_xml(f"<w:autoHyphenation {_W_DECL}/>"))
    lang = source.styles.element.find(
        f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}/{qn('w:lang')}"
    )
    lang.set(qn("w:val"), "ru-RU")
    source.add_paragraph("Список используемых идентификаторов данных клиентов")
    kept = source.add_paragraph("ФЛГМ-40М.ПС")
    kept._p.get_or_add_pPr().append(parse_xml(f"<w:suppressAutoHyphens {_W_DECL}/>"))

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    assert rendered.settings.element.find(qn("w:autoHyphenation")) is not None
    defaults = rendered.styles.element.find(
        f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}/{qn('w:lang')}"
    )
    assert defaults.get(qn("w:val")) == "ru-RU"
    last = next(p for p in rendered.paragraphs if p.text == "ФЛГМ-40М.ПС")
    assert _required(last._p.pPr).find(qn("w:suppressAutoHyphens")) is not None


def test_a_line_break_before_a_typed_number_stays() -> None:
    """'<br/>1. Общие принципы': an empty line, then the clause; the number is no list marker
    that would cut the empty line with it (79dce606)."""
    source = docx.Document()
    paragraph = source.add_paragraph()
    paragraph.add_run().add_break()
    paragraph.add_run("1. Общие принципы работы автопилотов")

    (block,) = _import(source)

    assert block.text().startswith("\n1. Общие")


def test_a_text_box_picture_is_not_also_its_paragraphs() -> None:
    """A logo inside a text box anchored in a paragraph is read with the box only, and its
    fallback copy not at all (40fbac66: two logos drawn six times)."""
    from docx.oxml import parse_xml
    from PIL import Image

    from foto2word.domain.document import ImageBlock

    png = BytesIO()
    Image.new("RGB", (40, 40), "green").save(png, format="PNG")
    source = docx.Document()
    holder = source.add_paragraph()
    holder.add_run().add_picture(BytesIO(png.getvalue()), width=Mm(30))
    inline = next(holder._p.iter(qn("wp:inline")))
    content = parse_xml(f"<w:txbxContent {_W_DECL}><w:p><w:r/></w:p></w:txbxContent>")
    run = inline.getparent().getparent()
    content.find(qn("w:p")).append(run)  # the picture now sits inside a box's content
    holder._p.append(content)
    source.add_paragraph("Республиканское унитарное предприятие")

    blocks = _import(source)

    assert len([b for b in blocks if isinstance(b, ImageBlock)]) == 1  # the box's, once


def test_bold_in_both_styles_cancels_out_and_strike_is_kept(
    profiles: YamlProfileRepository,
) -> None:
    """Bold toggles in Word: a bold paragraph style with a bold character style gives normal
    text; the run's own setting wins outright, and its strike-through is kept (277a1330)."""
    from docx.enum.style import WD_STYLE_TYPE

    source = docx.Document()
    para = source.styles.add_style("Основной текст (3)", WD_STYLE_TYPE.PARAGRAPH)
    para.font.bold = True
    char = source.styles.add_style("Основной текст (3) + 0 pt", WD_STYLE_TYPE.CHARACTER)
    char.font.bold = True
    paragraph = source.add_paragraph(style=para)
    paragraph.add_run("Прошу выдать мне разрешение ", style=char)
    struck = paragraph.add_run("в пределах квоты", style=char)
    struck.font.strike = True

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    runs = _textual(document.blocks[0]).runs
    assert [r.bold for r in runs] == [False, False]
    assert [r.strike for r in runs] == [False, True]
    assert rendered.paragraphs[0].runs[-1].font.strike


def test_a_rows_own_table_lines_stand_in_for_the_tables() -> None:
    """A row that switches the table's lines off (w:tblPrEx): its cells draw only their own,
    no top line over the blanks of the row above (39086d44)."""
    from docx.oxml import parse_xml

    source = docx.Document()
    table = source.add_table(rows=2, cols=1)
    table.style = "Table Grid"
    row = table.rows[1]._tr
    row.insert(
        0,
        parse_xml(
            f'<w:tblPrEx {_W_DECL}><w:tblBorders><w:top w:val="none"/><w:left w:val="none"/>'
            '<w:bottom w:val="none"/><w:right w:val="none"/><w:insideH w:val="none"/>'
            '<w:insideV w:val="none"/></w:tblBorders></w:tblPrEx>'
        ),
    )

    block = next(b for b in _import(source) if isinstance(b, TableBlock))

    upper, lower = block.rows[0].cells[0], block.rows[1].cells[0]
    assert upper.borders is not None
    assert upper.borders.top
    assert lower.borders is not None
    assert not lower.borders.top


def test_a_one_line_cells_mark_keeps_its_own_size(profiles: YamlProfileRepository) -> None:
    """'через Единый портал<br/>' with a 1 pt mark: the empty line after the break is 1 pt,
    not a whole line of the cell's text (c7ece4a1, ba30e53e)."""
    from docx.oxml import parse_xml

    source = docx.Document()
    cell = source.add_table(rows=1, cols=1).cell(0, 0)
    paragraph = cell.paragraphs[0]
    paragraph.add_run("через Единый портал государственных").font.size = Pt(12)
    paragraph.runs[0].add_break()
    paragraph._p.get_or_add_pPr().append(parse_xml(f'<w:rPr {_W_DECL}><w:sz w:val="2"/></w:rPr>'))

    document = _document(source)
    rendered = docx.Document(
        BytesIO(DocxRenderer().render(document, profiles.get("generic"), RenderMode.FAITHFUL))
    )

    mark = rendered.tables[0].cell(0, 0).paragraphs[0]._p.pPr.find(f"{qn('w:rPr')}/{qn('w:sz')}")
    assert mark is not None
    assert mark.get(qn("w:val")) == "2"


def test_an_empty_line_with_a_break_keeps_it_when_rendered(profiles: YamlProfileRepository) -> None:
    source = docx.Document()
    source.add_paragraph("Заявка")
    source.add_paragraph().add_run().add_break()
    source.add_table(rows=1, cols=1).cell(0, 0).text = "ФИО участника"

    rendered = docx.Document(
        BytesIO(
            DocxRenderer().render(_document(source), profiles.get("generic"), RenderMode.FAITHFUL)
        )
    )

    assert sum(1 for _ in rendered.element.body.iter(qn("w:br"))) == 1
