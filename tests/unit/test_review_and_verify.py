import base64
import json
import sys
from collections.abc import Sequence
from io import BytesIO
from types import SimpleNamespace
from typing import Any, cast

import httpx
import openai
import pytest
from PIL import Image

from foto2word.application.convert import ConversionOptions, ConvertDocument
from foto2word.application.ports import (
    Provenance,
    Reviewed,
    Understanding,
    UnderstandingContext,
    Usage,
)
from foto2word.application.review import ReviewLoop
from foto2word.application.verify import DocumentVerifier, FlagKind
from foto2word.domain.catalog import DEFAULT_CATALOG
from foto2word.domain.document import (
    Correction,
    Document,
    ExtractedField,
    FieldKind,
    ParagraphBlock,
    Run,
)
from foto2word.domain.images import PageImage
from foto2word.domain.patches import BlockPatch, DocumentPatcher, PatchOp, ReviewResult
from foto2word.infrastructure.ocr.yandex import ENDPOINT, SecondReaderError, YandexVisionReader
from foto2word.infrastructure.preview import WordPreviewer, detect_previewer, pdf_to_pages
from foto2word.infrastructure.profiles.yaml_repository import YamlProfileRepository
from foto2word.infrastructure.providers.openai_provider import OpenAIVisionModel
from foto2word.infrastructure.providers.schema import strict_json_schema
from foto2word.infrastructure.rendering.docx_renderer import DocxRenderer

PAGE = PageImage(b"jpeg", "image/jpeg", "photo.jpg")
CONTEXT = UnderstandingContext(DEFAULT_CATALOG)
PROV = Provenance("fake", "m", "v1", Usage(10, 5))


def _paragraph(block_id: str, text: str) -> ParagraphBlock:
    return ParagraphBlock(id=block_id, runs=(Run(text=text),))


# --- patches ------------------------------------------------------------------------------


def test_patcher_applies_replace_insert_delete_and_skips_invalid(act: Document) -> None:
    review = ReviewResult(
        patches=(
            BlockPatch(
                op=PatchOp.REPLACE, target_id="b5", block=_paragraph("x", "Комиссия:"), reason="r"
            ),
            BlockPatch(
                op=PatchOp.INSERT_AFTER, target_id="b1", block=_paragraph("r1", "Новый"), reason="r"
            ),
            BlockPatch(
                op=PatchOp.INSERT_AFTER, target_id="", block=_paragraph("r0", "Первый"), reason="r"
            ),
            BlockPatch(op=PatchOp.DELETE, target_id="b4", reason="r"),
            BlockPatch(op=PatchOp.DELETE, target_id="nope", reason="r"),
            BlockPatch(
                op=PatchOp.INSERT_AFTER, target_id="b1", block=_paragraph("b2", "dup"), reason="r"
            ),
        ),
        missing_corrections=(Correction(id="c9", block_id="b10", before="а", after="б"),),
    )

    outcome = DocumentPatcher().apply(act, review)

    ids = [b.id for b in outcome.document.blocks]
    assert ids[:3] == ["r0", "b1", "r1"]
    assert "b4" not in ids
    assert outcome.document.blocks[ids.index("b5")].text() == "Комиссия:"
    assert len(outcome.applied) == 4
    assert len(outcome.skipped) == 2
    assert outcome.document.corrections[-1].id == "c9"


def test_review_result_strict_schema_is_valid() -> None:
    schema = json.dumps(strict_json_schema(ReviewResult))

    assert '"oneOf"' not in schema
    assert '"default"' not in schema


# --- verification -------------------------------------------------------------------------


class StaticReader:
    def __init__(self, text: str) -> None:
        self.text = text

    def read_text(self, pages: Sequence[PageImage]) -> str:
        return self.text


def test_verifier_flags_invalid_missing_and_low_confidence(application: Document) -> None:
    broken = application.model_copy(
        update={
            "fields": (
                *application.fields,
                ExtractedField(kind=FieldKind.INN, value="000000000001", block_id="b1"),
                ExtractedField(kind=FieldKind.DATE, value="01.01.2020", block_id="b4"),
            ),
            "blocks": (
                *application.blocks,
                ParagraphBlock(id="b99", confidence=0.5, runs=(Run(text="??"),)),
            ),
        }
    )

    report = DocumentVerifier().verify(broken)

    kinds = {(f.kind, f.value) for f in report.flags}
    assert (FlagKind.INVALID_VALUE, "000000000001") in kinds
    assert (FlagKind.NOT_IN_TEXT, "000000000001") in kinds
    assert (FlagKind.NOT_IN_TEXT, "01.01.2020") in kinds
    assert any(f.kind is FlagKind.LOW_CONFIDENCE and f.block_id == "b99" for f in report.flags)
    assert not report.checked_with_second_reader


def test_clean_documents_have_no_flags(act: Document, application: Document) -> None:
    assert DocumentVerifier().verify(act).ok
    assert DocumentVerifier().verify(application).ok


def test_second_reader_mismatch_only_for_numbers(application: Document) -> None:
    ocr = "\n".join(application.iter_text()).replace("000000000000", "000000000011")

    report = DocumentVerifier(StaticReader(ocr)).verify(application, [PAGE])

    assert report.checked_with_second_reader
    assert [(f.kind, f.value) for f in report.flags] == [
        (FlagKind.SECOND_READER_MISMATCH, "000000000000")
    ]


# --- review loop --------------------------------------------------------------------------


class FakePreviewer:
    def __init__(self) -> None:
        self.calls = 0

    def preview(self, rendered: bytes) -> list[PageImage]:
        self.calls += 1
        assert rendered[:2] == b"PK"  # a real .docx
        return [PageImage(b"png", "image/png", "render-1.png")]

    def to_pdf(self, rendered: bytes) -> bytes:
        return b"%PDF-fake"


class ScriptedReviewer:
    def __init__(self, results: list[ReviewResult]) -> None:
        self.results = results
        self.seen: list[Document] = []

    def review(
        self,
        originals: Sequence[PageImage],
        rendered: Sequence[PageImage],
        document: Document,
        context: UnderstandingContext,
    ) -> Reviewed:
        self.seen.append(document)
        return Reviewed(self.results.pop(0), PROV)


def _loop(reviewer: ScriptedReviewer, previewer: FakePreviewer) -> ReviewLoop:
    return ReviewLoop(reviewer, previewer, DocxRenderer(), max_rounds=3)


def test_review_loop_stops_when_reviewer_is_satisfied(
    act: Document, profiles: YamlProfileRepository
) -> None:
    fix = BlockPatch(op=PatchOp.DELETE, target_id="b4", reason="лишний пробел")
    reviewer = ScriptedReviewer([ReviewResult(patches=(fix,)), ReviewResult()])
    previewer = FakePreviewer()

    report = _loop(reviewer, previewer).run([PAGE], act, profiles.get("act"), CONTEXT)

    assert len(report.rounds) == 1
    assert previewer.calls == 2
    assert "b4" not in [b.id for b in report.document.blocks]
    assert report.usage == Usage(20, 10)
    assert reviewer.seen[1] == report.document


def test_review_loop_stops_when_patches_change_nothing(
    act: Document, profiles: YamlProfileRepository
) -> None:
    useless = BlockPatch(op=PatchOp.DELETE, target_id="missing", reason="?")
    reviewer = ScriptedReviewer([ReviewResult(patches=(useless,))] * 3)

    report = _loop(reviewer, FakePreviewer()).run([PAGE], act, profiles.get("act"), CONTEXT)

    assert len(report.rounds) == 1
    assert report.document == act


class FakeVision:
    def __init__(self, document: Document) -> None:
        self.document = document

    def understand(
        self, pages: Sequence[PageImage], context: UnderstandingContext
    ) -> Understanding:
        return Understanding(self.document, PROV)


def test_convert_with_review_and_verification(
    act: Document, profiles: YamlProfileRepository
) -> None:
    reviewer = ScriptedReviewer([ReviewResult()])
    converter = ConvertDocument(
        FakeVision(act),
        DocxRenderer(),
        profiles,
        review_loop=_loop(reviewer, FakePreviewer()),
    )

    with_review = converter.execute([PAGE], ConversionOptions(review=True))
    without = converter.execute([PAGE], ConversionOptions())

    assert converter.can_review
    assert with_review.review is not None
    assert without.review is None
    assert with_review.verification.ok


# --- OpenAI review request ----------------------------------------------------------------


def test_openai_review_sends_both_photo_and_render(act: Document) -> None:
    captured: dict[str, Any] = {}

    def create(**kwargs: Any) -> SimpleNamespace:
        captured.update(kwargs)
        return SimpleNamespace(
            status="completed",
            output_text=json.dumps({"patches": [], "missing_corrections": []}),
            model="gpt-test",
            incomplete_details=None,
            usage=SimpleNamespace(input_tokens=100, output_tokens=20),
        )

    client = cast(openai.OpenAI, SimpleNamespace(responses=SimpleNamespace(create=create)))
    model = OpenAIVisionModel(client, "gpt-test")
    render = PageImage(b"png", "image/png", "render-1.png")

    reviewed = model.review([PAGE], [render], act, CONTEXT)

    assert reviewed.result == ReviewResult()
    assert reviewed.provenance.usage == Usage(100, 20)
    texts = [c["text"] for c in captured["input"][0]["content"] if c["type"] == "input_text"]
    assert texts[0].startswith("Оригинал")
    assert texts[1].startswith("Собранный документ")
    assert '"АКТ"' in texts[2]
    assert captured["text"]["format"]["name"] == "reviewresult"


# --- second reader ------------------------------------------------------------------------


def test_yandex_reader_posts_image_and_reads_full_text() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert str(request.url) == ENDPOINT
        assert request.headers["Authorization"] == "Api-Key key"
        body = json.loads(request.content)
        assert base64.b64decode(body["content"]) == b"jpeg"
        assert body["mimeType"] == "JPEG"
        return httpx.Response(200, json={"result": {"textAnnotation": {"fullText": "ИНН 123"}}})

    client = httpx.Client(transport=httpx.MockTransport(handler))

    assert YandexVisionReader("key", "folder", client).read_text([PAGE]) == "ИНН 123"


def test_yandex_reader_errors_are_wrapped() -> None:
    client = httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(403)))

    with pytest.raises(SecondReaderError):
        YandexVisionReader("key", "folder", client).read_text([PAGE])


# --- preview ------------------------------------------------------------------------------


def _pdf(pages: int) -> bytes:
    images = [Image.new("RGB", (300, 400), "white") for _ in range(pages)]
    buffer = BytesIO()
    images[0].save(buffer, format="PDF", save_all=True, append_images=images[1:])
    return buffer.getvalue()


def test_pdf_to_pages_renders_each_page() -> None:
    pages = pdf_to_pages(_pdf(3), dpi=72, max_pages=2)

    assert [p.source_name for p in pages] == ["render-1.jpg", "render-2.jpg"]
    assert Image.open(BytesIO(pages[0].data)).size == (300, 400)


@pytest.mark.skipif(sys.platform != "win32", reason="Word automation is Windows-only")
def test_word_previewer_renders_real_docx(act: Document, profiles: YamlProfileRepository) -> None:
    previewer = detect_previewer()
    if not isinstance(previewer, WordPreviewer):
        pytest.skip("Microsoft Word is not installed")
    rendered = DocxRenderer().render(act, profiles.get("act"), ConversionOptions().mode)

    pages = previewer.preview(rendered)

    assert len(pages) == 1
    assert Image.open(BytesIO(pages[0].data)).width > 500


def test_yandex_reader_returns_lines_with_boxes_on_the_page() -> None:
    answer = {
        "result": {
            "textAnnotation": {
                "width": "1000",
                "height": "2000",
                "fullText": "ИНН 123",
                "blocks": [
                    {
                        "lines": [
                            {
                                "text": "ИНН 123",
                                "boundingBox": {
                                    "vertices": [
                                        {"x": "100", "y": "200"},
                                        {"x": "100", "y": "240"},
                                        {"x": "500", "y": "240"},
                                        {"x": "500", "y": "200"},
                                    ]
                                },
                            }
                        ]
                    }
                ],
            }
        }
    }
    client = httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(200, json=answer)))

    (line,) = YandexVisionReader("key", "folder", client).read_lines([PAGE])

    assert (line.page, line.text) == (1, "ИНН 123")
    assert (line.box.x0, line.box.y0, line.box.x1, line.box.y1) == (0.1, 0.1, 0.5, 0.12)


def test_missing_fonts_count_only_what_the_text_uses() -> None:
    import docx

    from foto2word.infrastructure.preview import missing_fonts

    source = docx.Document()
    source.add_paragraph("Текст").runs[0].font.name = "Редкий Шрифт"
    source.styles.add_style("Неиспользуемый", 1).font.name = "Другой Редкий"
    buffer = BytesIO()
    source.save(buffer)

    missing = missing_fonts(buffer.getvalue(), frozenset({"calibri", "cambria", "times new roman"}))

    assert "редкий шрифт" in missing
    assert "другой редкий" not in missing  # a style no text uses
