from collections.abc import Sequence
from io import BytesIO

import docx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image, ImageDraw

from foto2word.application.convert import ConvertDocument
from foto2word.application.ports import Provenance, Understanding, UnderstandingContext
from foto2word.application.sessions import SessionService
from foto2word.domain.catalog import DEFAULT_CATALOG
from foto2word.domain.dates import DateShifter
from foto2word.domain.document import (
    Document,
    DocumentTypeGuess,
    ParagraphBlock,
    Run,
    TableBlock,
    TableCell,
    TableRow,
)
from foto2word.domain.errors import InvalidImageError
from foto2word.domain.extras import RenderExtras, SignatureMode
from foto2word.domain.images import PageImage
from foto2word.domain.profile import RenderMode
from foto2word.infrastructure.exports import scan_pdf, tables_xlsx
from foto2word.infrastructure.images import PillowImageLoader
from foto2word.infrastructure.profiles.yaml_repository import YamlProfileRepository
from foto2word.infrastructure.rendering.docx_renderer import DocxRenderer
from foto2word.infrastructure.session_store import InMemorySessionStore
from foto2word.infrastructure.signature import prepare_signature
from foto2word.web.app import create_app


def _doc(*texts: str) -> Document:
    return Document(
        document_type=DocumentTypeGuess(id="other", title="x", confidence=1.0),
        blocks=tuple(ParagraphBlock(id=f"b{i}", runs=(Run(text=t),)) for i, t in enumerate(texts)),
    )


def _png(draw_ink: bool = True, transparent: bool = False) -> bytes:
    mode, background = ("RGBA", (0, 0, 0, 0)) if transparent else ("RGB", "white")
    image = Image.new(mode, (300, 100), background)
    if draw_ink:
        ImageDraw.Draw(image).line((20, 70, 120, 20, 200, 80, 280, 30), fill="black", width=5)
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


# --- dates ----------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("в отопительном периоде 2026гг.", [("2026", "2027")]),
        ("«12» Января 2026г", [("«12» Января 2026", "«12» Января 2027")]),
        ("Протокол от 05.03.2025 № 3", []),  # 'от' = a past reference
        ("Срок до 31.12.2026 включительно", [("31.12.2026", "31.12.2027")]),
        ("Паспорт выдан 13.05.2022 УМВД", []),
        ("«29» февраля 2024 г.", [("«29» февраля 2024", "«28» февраля 2025")]),
        ("ИНН 000000000000, тел 80000000000", []),
        ("с 2026 года по 2026 год", []),  # ambiguous repeated fragment is skipped
    ],
)
def test_date_shifter(text: str, expected: list[tuple[str, str]]) -> None:
    proposals = DateShifter().propose(_doc(text), years=1)

    assert [(p.before, p.after) for p in proposals] == expected


def test_date_shift_proposals_apply_cleanly(act: Document) -> None:
    proposals = DateShifter().propose(act, years=1)

    assert {p.block_id for p in proposals} == {"b2", "b3"}
    assert DateShifter().propose(act, years=0) == ()


# --- signature --------------------------------------------------------------------------------


@pytest.mark.parametrize("transparent", [False, True])
def test_prepare_signature_makes_tight_transparent_png(transparent: bool) -> None:
    png = prepare_signature(_png(transparent=transparent))

    image = Image.open(BytesIO(png))
    assert image.mode == "RGBA"
    corner = image.getpixel((0, 0))
    assert isinstance(corner, tuple)
    assert corner[3] < 255  # background became transparent
    assert image.width < 300


def test_prepare_signature_rejects_empty_images() -> None:
    with pytest.raises(InvalidImageError):
        prepare_signature(_png(draw_ink=False))
    with pytest.raises(InvalidImageError):
        prepare_signature(b"not an image")


def _render(document: Document, profiles: YamlProfileRepository, extras: RenderExtras) -> bytes:
    return DocxRenderer().render(document, profiles.get("act"), RenderMode.FAITHFUL, extras)


def test_own_signature_goes_only_to_chosen_lines(
    act: Document, profiles: YamlProfileRepository
) -> None:
    extras = RenderExtras(
        signature_mode=SignatureMode.OWN,
        signature_png=prepare_signature(_png()),
        signed_blocks=frozenset({"b19"}),
    )

    output = docx.Document(BytesIO(_render(act, profiles, extras)))

    assert len(output.inline_shapes) == 1
    lines = [p.text for p in output.paragraphs if p.text.startswith(("Председатель", "Члены"))]
    assert lines[-1] == "Члены комиссии:\t________________\tПетров П.П."


def test_placeholder_only_where_original_was_signed(
    act: Document, profiles: YamlProfileRepository
) -> None:
    extras = RenderExtras(signature_mode=SignatureMode.PLACEHOLDER)

    texts = [p.text for p in docx.Document(BytesIO(_render(act, profiles, extras))).paragraphs]

    assert "Председатель комиссии\t/подпись/\tСидоров С.С." in texts
    assert "Представитель ТСО\t________________\t" in texts


def test_scan_appendix_adds_pages(act: Document, profiles: YamlProfileRepository) -> None:
    photo = PillowImageLoader().load_bytes(_png(), "p.png")

    output = docx.Document(BytesIO(_render(act, profiles, RenderExtras(appendix=(photo, photo)))))

    assert len(output.inline_shapes) == 2
    assert "Приложение. Фото оригинала" in [p.text for p in output.paragraphs]


# --- exports ----------------------------------------------------------------------------------


def test_scan_pdf_has_a_page_per_photo() -> None:
    photo = PillowImageLoader().load_bytes(_png(), "p.png")

    pdf = scan_pdf([photo, photo])

    assert pdf.startswith(b"%PDF")
    assert pdf.count(b"/Type /Page\n") + pdf.count(b"/Type /Page ") + pdf.count(b"/Type/Page") >= 1


def test_tables_xlsx_keeps_numbers_codes_and_merges() -> None:
    table = TableBlock(
        id="t1",
        header_rows=1,
        rows=(
            TableRow(cells=(TableCell(runs=(Run(text="Показатель"),), col_span=2),)),
            TableRow(
                cells=(TableCell(runs=(Run(text="код"),)), TableCell(runs=(Run(text="007"),)))
            ),
            TableRow(cells=(TableCell(runs=(Run(text="Т"),)), TableCell(runs=(Run(text="14,3"),)))),
        ),
    )
    document = _doc("x").model_copy(update={"blocks": (table,)})

    data = tables_xlsx(document)

    assert data is not None
    sheet = load_workbook(BytesIO(data)).active
    assert sheet is not None
    assert sheet["A1"].value == "Показатель"
    assert sheet["A1"].font.bold
    assert "A1:B1" in [str(r) for r in sheet.merged_cells.ranges]
    assert sheet["B2"].value == "007"
    assert sheet["B3"].value == pytest.approx(14.3)
    assert tables_xlsx(_doc("no tables")) is None


# --- through the web API ----------------------------------------------------------------------


class FakeVision:
    def __init__(self, document: Document) -> None:
        self.document = document

    def understand(
        self, pages: Sequence[PageImage], context: UnderstandingContext
    ) -> Understanding:
        return Understanding(self.document, Provenance("fake", "fake-1", "v1"))


class FakePreviewer:
    def preview(self, rendered: bytes) -> list[PageImage]:
        return []

    def to_pdf(self, rendered: bytes) -> bytes:
        return b"%PDF-fake"


@pytest.fixture
def client(act: Document, profiles: YamlProfileRepository) -> TestClient:
    service = SessionService(
        ConvertDocument(FakeVision(act), DocxRenderer(), profiles),
        InMemorySessionStore(),
        DEFAULT_CATALOG,
        run=lambda job: job(),
        previewer=FakePreviewer(),
    )
    return TestClient(create_app(service, PillowImageLoader(), DEFAULT_CATALOG))


def _session(client: TestClient) -> str:
    response = client.post("/api/sessions", files=[("photos", ("p.png", _png(), "image/png"))])
    return str(response.json()["id"])


def test_web_shift_years_and_refill_fields(client: TestClient) -> None:
    sid = _session(client)

    view = client.post(
        f"/api/sessions/{sid}/decisions",
        json={"shift_years": 1, "field_values": {"2": "Иванов И.И."}},
    ).json()

    texts = {b["id"]: b["text"] for b in view["blocks"]}
    # The pencil already moved 2026 -> 2027; shifting by a year moves the result once more.
    assert "отопительном периоде 2028 г." in texts["b2"]
    assert texts["b6"].endswith("Иванов И.И.")
    assert view["fields"][2]["current"] == "Иванов И.И."
    assert {c["after"] for c in view["date_changes"]} >= {"2028"}


def test_web_signature_requires_consent_and_signs(client: TestClient) -> None:
    sid = _session(client)
    files = {"image": ("s.png", _png(), "image/png")}

    refused = client.post(f"/api/sessions/{sid}/signature", files=files)
    accepted = client.post(
        f"/api/sessions/{sid}/signature", files=files, data={"own_signature": "true"}
    )
    view = client.post(
        f"/api/sessions/{sid}/decisions",
        json={"signature_mode": "own", "signed_blocks": ["b19", "b20"], "attach_scan": True},
    ).json()

    assert refused.status_code == 400
    assert accepted.json()["has_signature"] is True
    assert view["signed_blocks"] == ["b19", "b20"]
    document = docx.Document(BytesIO(client.get(f"/api/sessions/{sid}/document.docx").content))
    assert len(document.inline_shapes) == 3  # two signatures + one photo of the original


def test_web_exports(client: TestClient) -> None:
    sid = _session(client)

    assert client.get(f"/api/sessions/{sid}/document.pdf").content == b"%PDF-fake"
    assert client.get(f"/api/sessions/{sid}/scan.pdf").content.startswith(b"%PDF")
    assert client.get(f"/api/sessions/{sid}/tables.xlsx").status_code == 404  # the act has none


def test_web_editable_output_is_a_decision(client: TestClient) -> None:
    sid = _session(client)

    view = client.post(f"/api/sessions/{sid}/decisions", json={"editable": True}).json()
    again = client.post(f"/api/sessions/{sid}/decisions", json={}).json()

    assert view["editable"] is True
    assert again["editable"] is False
