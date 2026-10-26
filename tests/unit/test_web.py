from collections.abc import Callable, Sequence
from datetime import UTC, datetime, timedelta
from io import BytesIO

import docx
import pytest
from fastapi.testclient import TestClient
from PIL import Image

from foto2word.application.convert import ConvertDocument
from foto2word.application.ports import Provenance, Understanding, UnderstandingContext
from foto2word.application.sessions import Session, SessionService, SessionStatus
from foto2word.domain.catalog import DEFAULT_CATALOG
from foto2word.domain.document import Document
from foto2word.domain.errors import VisionModelError
from foto2word.domain.images import PageImage
from foto2word.infrastructure.images import PillowImageLoader
from foto2word.infrastructure.profiles.yaml_repository import YamlProfileRepository
from foto2word.infrastructure.rendering.docx_renderer import DocxRenderer
from foto2word.infrastructure.session_store import InMemorySessionStore
from foto2word.web.app import UploadLimits, create_app


class FakeVision:
    def __init__(self, document: Document | None) -> None:
        self.document = document

    def understand(
        self, pages: Sequence[PageImage], context: UnderstandingContext
    ) -> Understanding:
        if self.document is None:
            raise VisionModelError("модель недоступна")
        return Understanding(self.document, Provenance("fake", "fake-1", "v1"))


class FakePreviewer:
    def preview(self, rendered: bytes) -> list[PageImage]:
        return [PageImage(b"\xff\xd8preview", "image/jpeg", "render-1.jpg")]

    def to_pdf(self, rendered: bytes) -> bytes:
        return b"%PDF-fake"


def _client(
    document: Document | None,
    profiles: YamlProfileRepository,
    run: Callable[[Callable[[], None]], None] = lambda job: job(),
) -> TestClient:
    converter = ConvertDocument(FakeVision(document), DocxRenderer(), profiles)
    service = SessionService(
        converter, InMemorySessionStore(), DEFAULT_CATALOG, run=run, previewer=FakePreviewer()
    )
    app = create_app(service, PillowImageLoader(), DEFAULT_CATALOG, UploadLimits(max_files=2))
    return TestClient(app)


def _photo() -> tuple[str, bytes, str]:
    buffer = BytesIO()
    Image.new("RGB", (60, 80), "white").save(buffer, format="PNG")
    return ("page.png", buffer.getvalue(), "image/png")


@pytest.fixture
def client(act: Document, profiles: YamlProfileRepository) -> TestClient:
    return _client(act, profiles)


def _create(client: TestClient) -> dict[str, object]:
    response = client.post("/api/sessions", files=[("photos", _photo())])
    assert response.status_code == 201, response.text
    return dict(response.json())


def test_index_and_catalog(client: TestClient) -> None:
    page = client.get("/classic").text
    assert "Сфотографируйте документ" in page
    assert "/static/app.js?v=" in page
    assert client.get("/static/app.js").status_code == 200
    assert {"id": "act", "title": "Акт"} in client.get("/api/catalog").json()


def test_full_flow_upload_decide_download(client: TestClient) -> None:
    session = _create(client)
    assert session["status"] == "ready"
    document_type = session["document_type"]
    assert isinstance(document_type, dict)
    assert document_type["id"] == "act"
    corrections = session["corrections"]
    assert isinstance(corrections, list)
    assert [c["accepted"] for c in corrections] == [True, True]
    sid = session["id"]

    decided = client.post(
        f"/api/sessions/{sid}/decisions",
        json={
            "mode": "standard",
            "rejected_corrections": ["c2"],
            "accepted_suggestions": [1],
            "block_edits": {"b5": "Комиссия в составе следующих лиц:"},
        },
    ).json()

    assert decided["mode"] == "standard"
    assert [c["accepted"] for c in decided["corrections"]] == [True, False]
    texts = {b["id"]: b["text"] for b in decided["blocks"]}
    assert texts["b3"] == "г.Тестовый\t«12» Января 2026г"
    assert texts["b5"] == "Комиссия в составе следующих лиц:"
    assert texts["b8"].endswith("Тестоград»")
    download = client.get(f"/api/sessions/{sid}/document.docx")
    assert download.headers["content-type"].startswith("application/vnd.openxmlformats")
    paragraphs = [p.text for p in docx.Document(BytesIO(download.content)).paragraphs]
    assert "Комиссия в составе следующих лиц:" in paragraphs
    assert client.get(f"/api/sessions/{sid}/preview/1").content.startswith(b"\xff\xd8")
    assert client.get(f"/api/sessions/{sid}/photo/1").status_code == 200
    assert client.get(f"/api/sessions/{sid}/photo/2").status_code == 404


def test_changing_type_changes_profile(client: TestClient) -> None:
    sid = _create(client)["id"]

    decided = client.post(f"/api/sessions/{sid}/decisions", json={"document_type_id": "contract"})

    assert decided.json()["profile"] == "generic"
    assert decided.json()["chosen_type"] == "contract"


def test_source_edits_survive_repeated_saves_with_year_shift(client: TestClient) -> None:
    sid = _create(client)["id"]
    changes = {"block_edits": {"b5": "Состав комиссии 2026"}, "shift_years": 1}
    first = client.post(f"/api/sessions/{sid}/decisions", json=changes).json()
    block = next(b for b in first["blocks"] if b["id"] == "b5")
    assert block["text"] == "Состав комиссии 2027"
    assert block["source_text"] == "Состав комиссии 2026"
    changes["block_edits"] = first["block_edits"]
    second = client.post(f"/api/sessions/{sid}/decisions", json=changes).json()
    assert next(b for b in second["blocks"] if b["id"] == "b5")["text"] == block["text"]


def test_invalid_edit_and_unknown_session(client: TestClient) -> None:
    sid = _create(client)["id"]

    bad_edit = client.post(f"/api/sessions/{sid}/decisions", json={"block_edits": {"b19": "x"}})

    assert bad_edit.status_code == 400
    assert client.get("/api/sessions/nope").status_code == 404
    assert client.delete(f"/api/sessions/{sid}").status_code == 204
    assert client.get(f"/api/sessions/{sid}").status_code == 404


def test_upload_validation(client: TestClient) -> None:
    too_many = client.post("/api/sessions", files=[("photos", _photo())] * 3)
    not_image = client.post("/api/sessions", files=[("photos", ("a.txt", b"hello", "text/plain"))])

    assert too_many.status_code == 400
    assert not_image.status_code == 400


def test_model_failure_is_reported(profiles: YamlProfileRepository) -> None:
    client = _client(None, profiles)

    session = _create(client)

    assert session["status"] == "failed"
    assert session["error"] == "модель недоступна"


def test_decisions_before_analysis_are_rejected(
    act: Document, profiles: YamlProfileRepository
) -> None:
    client = _client(act, profiles, run=lambda job: None)  # analysis never runs

    sid = _create(client)["id"]

    assert client.get(f"/api/sessions/{sid}").json()["status"] == "analyzing"
    assert client.post(f"/api/sessions/{sid}/decisions", json={}).status_code == 409
    assert client.get(f"/api/sessions/{sid}/document.docx").status_code == 409


def test_store_expires_sessions() -> None:
    now = [datetime(2026, 1, 1, tzinfo=UTC)]
    store = InMemorySessionStore(ttl=timedelta(hours=24), clock=lambda: now[0])
    store.save(Session(id="s1", pages=(), review=False, created_at=now[0]))

    assert store.get("s1") is not None
    now[0] += timedelta(hours=25)
    assert store.get("s1") is None
    assert len(store) == 0


def test_session_status_values() -> None:
    assert [s.value for s in SessionStatus] == ["analyzing", "ready", "failed"]
